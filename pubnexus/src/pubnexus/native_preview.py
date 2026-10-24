"""Bounded native source windows, separate from extraction and PDF rendering.

All strings are literal content. Callers must use textContent, never innerHTML.
Asset IDs are opaque logical addresses; package member paths never cross the API.
"""
from __future__ import annotations

import hashlib
import io
import math
import re
import threading
import time
from collections import OrderedDict
from pathlib import Path
import zipfile

from lxml import etree

from . import documents, utils
from .preview import PreviewUnavailable

MAX_SOURCE_BYTES = documents.MAX_INPUT_BYTES
MAX_WINDOW_BYTES = 2 * 1024 * 1024
MAX_ASSET_BYTES = 16 * 1024 * 1024
MAX_ASSET_CACHE_BYTES = 32 * 1024 * 1024
MAX_WINDOW_IMAGES = 128
MAX_SEARCH_CELLS = 250_000
_SUPPORTED = frozenset({'.docx', '.xlsx'})
_RASTER_MIME = {'PNG': 'image/png', 'JPEG': 'image/jpeg', 'GIF': 'image/gif',
                'WEBP': 'image/webp', 'BMP': 'image/bmp'}
_assets = OrderedDict()
_assets_lock = threading.Lock()


def _stat_key(stat):
    return stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns


def _remember_asset(read, identity, blob, mime):
    key = (str(read.path), read.revision, identity)
    with _assets_lock:
        _assets[key] = (_stat_key(read.before), blob, mime)
        _assets.move_to_end(key)
        while len(_assets) > 128 or sum(len(value[1]) for value in _assets.values()) > MAX_ASSET_CACHE_BYTES:
            _assets.popitem(last=False)


def supports(source: str | Path) -> bool:
    return Path(source).suffix.lower() in _SUPPORTED


class _Read:
    def __init__(self, source, revision, cancel, timeout):
        if not isinstance(timeout, (int, float)) or not math.isfinite(timeout) or timeout <= 0:
            raise PreviewUnavailable('미리보기 제한 시간이 올바르지 않습니다.')
        self.cancel = cancel
        self.deadline = time.monotonic() + timeout
        self.started = time.monotonic()
        self.bytes = 0
        self.warnings = []
        self.check()
        self.path = Path(source).resolve(strict=True)
        self.before = self.path.stat()
        if not self.path.is_file() or self.before.st_size > MAX_SOURCE_BYTES:
            raise PreviewUnavailable('원본 문서 크기 제한을 초과했습니다.')
        digest = hashlib.sha256()
        with open(utils.long_path(self.path), 'rb') as stream:
            total = 0
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                self.check()
                total += len(chunk)
                if total > MAX_SOURCE_BYTES:
                    raise PreviewUnavailable('원본 문서 크기 제한을 초과했습니다.')
                digest.update(chunk)
        self.revision = digest.hexdigest()
        if revision is not None and revision != self.revision:
            raise PreviewUnavailable('원본이 변경됐습니다. 다시 여세요.')

    def check(self):
        if self.cancel is not None and self.cancel():
            raise PreviewUnavailable('원본 미리보기가 취소됐습니다.')
        if time.monotonic() > self.deadline:
            raise PreviewUnavailable('원본 미리보기 제한 시간을 초과했습니다.')

    def text(self, value):
        self.check()
        text = '' if value is None else value.isoformat() if hasattr(value, 'isoformat') else str(value)
        self.bytes += len(text.encode('utf-8'))
        if self.bytes > MAX_WINDOW_BYTES:
            raise PreviewUnavailable('내용 창 크기 제한을 초과했습니다. 더 작은 범위를 여세요.')
        return text

    def warn(self, message):
        if message not in self.warnings:
            self.warnings.append(message)

    def finish(self, result):
        self.check()
        after = self.path.stat()
        if _stat_key(self.before) != _stat_key(after):
            raise PreviewUnavailable('원본이 변경됐습니다. 다시 여세요.')
        result.update(source_ready=True, source_revision=self.revision, revision=self.revision,
                      format=self.path.suffix[1:].lower(), pages=[], source_label='원본 내용',
                      coordinate_space='logical_document', load_seconds=time.monotonic() - self.started)
        result.update(warnings=self.warnings, partial=bool(self.warnings))
        return result


def _window_number(value, minimum, maximum):
    if type(value) is not int or not minimum <= value <= maximum:
        raise PreviewUnavailable('미리보기 범위가 올바르지 않습니다.')


def _image_block(identity, blob):
    block = {'kind': 'image', 'id': identity}
    if len(blob) > MAX_ASSET_BYTES:
        return dict(block, unsupported=True, text='그림 미리보기 크기 제한을 초과했습니다.')
    from PIL import Image
    try:
        with Image.open(io.BytesIO(blob)) as image:
            mime = _RASTER_MIME.get(image.format)
            size = image.size
            if not mime or size[0] * size[1] > 64_000_000:
                raise ValueError('unsupported raster')
    except (OSError, ValueError, Image.DecompressionBombError):
        return dict(block, unsupported=True, text='이 그림 형식은 내용 보기에서 표시하지 못합니다.')
    return dict(block, asset_id=identity, mime=mime, image_sha256=hashlib.sha256(blob).hexdigest(),
                width=size[0], height=size[1])


def _docx_images(block):
    element = block._element
    for number, blip in enumerate(element.xpath('.//a:blip'), 1):
        relationship = blip.get(documents._R + 'embed')
        part = block.part.related_parts.get(relationship) if relationship else None
        if part is not None:
            yield number, part.blob


def _docx(read, start, count):
    from docx import Document
    from docx.table import Table
    parts = documents._guard_zip(read.path)
    native = Document(utils.long_path(read.path))
    if any(re.match(r'word/(?:header|footer|footnotes|endnotes|comments)', part) for part in parts):
        read.warn('내용 보기는 머리말·꼬리말·각주·미주·주석을 포함하지 않습니다. 원본 모양 보기에서 확인하세요.')
    blocks, more, image_count = [], False, 0
    for index, block in enumerate(native.iter_inner_content()):
        read.check()
        if index < start:
            continue
        if index >= start + count:
            more = True
            break
        identity = f'docx-b{index + 1:06d}'
        if block._element.xpath('.//m:oMath | .//w:txbxContent'):
            read.warn('수식·글상자는 내용 보기에 포함되지 않을 수 있습니다. 원본 모양 보기에서 확인하세요.')
        if isinstance(block, Table):
            rows, truncated = [], False
            for number, row in enumerate(block.rows):
                read.check()
                if number >= 200:
                    truncated = True
                    break
                rows.append([read.text(cell.text) for cell in row.cells[:64]])
                truncated |= len(row.cells) > 64
            blocks.append({'id': identity, 'kind': 'table', 'rows': rows, 'truncated': truncated})
        else:
            style = block.style.name if block.style is not None else ''
            blocks.append({'id': identity, 'kind': 'paragraph', 'text': read.text(block.text),
                           'heading': int(style[-1]) if re.fullmatch(r'Heading [1-9]', style) else 0})
        for number, blob in _docx_images(block):
            read.check()
            if image_count >= MAX_WINDOW_IMAGES:
                read.warn('이 내용 창의 그림 개수 제한을 초과했습니다. 더 작은 범위를 여세요.')
                break
            image_count += 1
            metadata = _image_block(f'{identity}-image{number:04d}', blob)
            if not metadata.get('unsupported'):
                _remember_asset(read, metadata['id'], blob, metadata['mime'])
            else:
                read.warn(metadata['text'])
            blocks.append(metadata)
    return {'source_kind': 'content', 'blocks': blocks, 'start': start, 'count': count,
            'next_start': start + count if more else None, 'truncated': more,
            'source_unit_label': '내용', 'title': native.core_properties.title or read.path.stem}


def _xlsx(read, sheet_index, start_row, start_col, row_count, col_count):
    from openpyxl import load_workbook
    documents._guard_zip(read.path)
    native = load_workbook(utils.long_path(read.path), read_only=True, data_only=False, keep_links=False)
    iterator = None
    try:
        read.check()
        sheets = [{'index': index, 'name': sheet.title, 'state': sheet.sheet_state}
                  for index, sheet in enumerate(native.worksheets)]
        if not 0 <= sheet_index < len(sheets):
            raise PreviewUnavailable('시트를 선택하세요.')
        sheet = native.worksheets[sheet_index]
        iterator = sheet.iter_rows(min_row=start_row, max_row=start_row + row_count,
                                   min_col=start_col, max_col=start_col + col_count - 1, values_only=True)
        rows, more = [], False
        for index, row in enumerate(iterator):
            read.check()
            if index == row_count:
                more = True
                break
            rows.append([read.text(value) for value in row])
        return {'source_kind': 'sheet', 'sheets': sheets, 'sheet_index': sheet_index,
                'start_row': start_row, 'start_col': start_col, 'row_count': row_count, 'col_count': col_count,
                'rows': rows, 'max_row': sheet.max_row, 'max_col': sheet.max_column,
                'has_more': more,
                'source_unit_label': '셀', 'title': sheet.title}
    finally:
        if iterator is not None:
            iterator.close()
        native.close()


def view(source: str | Path, *, revision=None, start=0, count=80, sheet_index=0,
         start_row=1, start_col=1, row_count=60, col_count=24, cancel=None, timeout=10) -> dict | None:
    """Return a content/sheet window, or None so the existing renderer can handle it."""
    if not supports(source):
        return None
    for value, minimum, maximum in ((start, 0, 1_000_000), (count, 1, 200), (sheet_index, 0, 10000),
            (start_row, 1, 1_048_576), (start_col, 1, 16_384), (row_count, 1, 200), (col_count, 1, 64)):
        _window_number(value, minimum, maximum)
    if not isinstance(timeout, (int, float)) or not math.isfinite(timeout) or timeout <= 0:
        raise PreviewUnavailable('미리보기 제한 시간이 올바르지 않습니다.')
    try:
        read = _Read(source, revision, cancel, timeout)
        result = (_xlsx(read, sheet_index, start_row, start_col, row_count, col_count)
                  if read.path.suffix.lower() == '.xlsx' else _docx(read, start, count))
        return read.finish(result)
    except PreviewUnavailable:
        raise
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile,
            etree.LxmlError, documents.DocumentExtractionError) as error:
        raise PreviewUnavailable('원본 내용을 읽지 못했습니다: ' + str(error)) from error


def read_asset(source: str | Path, asset_id: str, revision: str, *, cancel=None, timeout=10) -> tuple[bytes, str]:
    """Read one original raster image using an opaque, revision-bound block ID."""
    match = re.fullmatch(r'docx-b([0-9]{6})-image([0-9]{4})', str(asset_id))
    if (not match or Path(source).suffix.lower() != '.docx' or
            not re.fullmatch(r'[0-9a-f]{64}', str(revision or ''))):
        raise PreviewUnavailable('그림 주소가 올바르지 않습니다.')
    try:
        if cancel is not None and cancel():
            raise PreviewUnavailable('원본 미리보기가 취소됐습니다.')
        source_path = Path(source).resolve(strict=True)
        with _assets_lock:
            key = (str(source_path), revision, asset_id)
            hit = _assets.get(key)
            if hit is not None and hit[0] == _stat_key(source_path.stat()):
                _assets.move_to_end(key)
                return hit[1], hit[2]
        read = _Read(source, revision, cancel, timeout)
        documents._guard_zip(read.path)
        from docx import Document
        native = Document(utils.long_path(read.path))
        target, image_index = int(match[1]) - 1, int(match[2])
        for index, block in enumerate(native.iter_inner_content()):
            read.check()
            if index == target:
                for number, blob in _docx_images(block):
                    if number == image_index:
                        metadata = _image_block(asset_id, blob)
                        if metadata.get('unsupported'):
                            raise PreviewUnavailable(metadata['text'])
                        read.finish({})
                        _remember_asset(read, asset_id, blob, metadata['mime'])
                        return blob, metadata['mime']
                break
        raise PreviewUnavailable('원본 그림을 찾지 못했습니다.')
    except PreviewUnavailable:
        raise
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, etree.LxmlError,
            documents.DocumentExtractionError) as error:
        raise PreviewUnavailable('원본 그림을 읽지 못했습니다: ' + str(error)) from error


def search(source: str | Path, query: str, revision=None, *, cancel=None, limit=200, timeout=10) -> dict | None:
    """Find literal text at native block/cell addresses without a display conversion."""
    if not supports(source):
        return None
    _window_number(limit, 1, 200)
    query = str(query or '').strip()[:500]
    try:
        read = _Read(source, revision, cancel, timeout)
        if not query:
            return read.finish({'ok': True, 'results': [], 'truncated': False})
        documents._guard_zip(read.path)
        results, scanned, truncated = [], 0, False
        pattern = re.compile(re.escape(query), re.IGNORECASE) if query else None

        def add(text, location):
            nonlocal scanned, truncated
            read.check()
            scanned += len(text.encode('utf-8'))
            if scanned > documents.MAX_TEXT_BYTES:
                truncated = True
                return False
            for match in pattern.finditer(text) if pattern else ():
                if len(results) >= limit:
                    truncated = True
                    return False
                loc = dict(location, offset=match.start(), length=match.end() - match.start())
                results.append({'side': 'source', 'text': text[max(0, match.start() - 120):match.end() + 240],
                                'locations': [loc], 'mapping': 'exact', 'block_id': location.get('block_id')})
            return True

        if read.path.suffix.lower() == '.docx':
            from docx import Document
            from docx.table import Table
            for index, block in enumerate(Document(utils.long_path(read.path)).iter_inner_content(), 1):
                read.check()
                location = {'kind': 'content', 'block_id': f'docx-b{index:06d}'}
                if isinstance(block, Table):
                    stop = False
                    for row_number, row in enumerate(block.rows, 1):
                        for col_number, cell in enumerate(row.cells, 1):
                            if not add(cell.text, dict(location, row=row_number, col=col_number)):
                                stop = True
                                break
                        if stop:
                            break
                    if stop:
                        break
                elif not add(block.text, location):
                    break
        else:
            from openpyxl import load_workbook
            from openpyxl.utils.cell import get_column_letter
            native = load_workbook(utils.long_path(read.path), read_only=True, data_only=False, keep_links=False)
            cells, iterator = 0, None
            try:
                stop = False
                for sheet_index, sheet in enumerate(native.worksheets):
                    iterator = sheet.iter_rows()
                    try:
                        for row_number, row in enumerate(iterator, 1):
                            read.check()
                            for col_number, cell in enumerate(row, 1):
                                cells += 1
                                if cells > MAX_SEARCH_CELLS:
                                    truncated = stop = True
                                    break
                                value = cell.value
                                text = '' if value is None else value.isoformat() if hasattr(value, 'isoformat') else str(value)
                                location = {'kind': 'sheet', 'sheet_index': sheet_index, 'sheet': sheet.title,
                                            'cell': f'{get_column_letter(col_number)}{row_number}',
                                            'row': row_number, 'col': col_number}
                                if not add(text, location):
                                    stop = True
                                    break
                            if stop:
                                break
                    finally:
                        iterator.close()
                    if stop:
                        break
            finally:
                native.close()
        return read.finish({'ok': True, 'results': results, 'truncated': truncated})
    except PreviewUnavailable:
        raise
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile,
            etree.LxmlError, documents.DocumentExtractionError) as error:
        raise PreviewUnavailable('원본 내용을 검색하지 못했습니다: ' + str(error)) from error
