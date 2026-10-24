"""로컬 native 문서 → 정본. LLM·원격 자원·매크로 실행을 사용하지 않는다.

페이지가 없는 문서는 논리 unit/블록 주소만 부여한다. 그림 경로는 기존 앱
저장소 또는 asset_dir 기준 상대경로이며 export 모듈에서 휴대용 복사본을 만든다.
"""
from __future__ import annotations

import csv
import base64
from email import policy as email_policy
from email.parser import BytesParser
import hashlib
import html
import io
import json
import os
import posixpath
import re
import shutil
import stat
import subprocess
import tempfile
import time
import zipfile
from urllib.parse import unquote, urlsplit

from lxml import etree
from pathlib import Path

from . import schema, store, utils

MAX_INPUT_BYTES = 256 * 1024 * 1024
MAX_TEXT_BYTES = 32 * 1024 * 1024
MAX_CELLS = 250_000
MAX_ZIP_BYTES = 512 * 1024 * 1024
MAX_PART_BYTES = 64 * 1024 * 1024
MAX_ZIP_MEMBERS = 20_000
MAX_MAIL_PARTS = 2_000
MAX_MAIL_MESSAGES = 10_000
PPTX_IMAGE_OCR_VERSION = 1
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_O = "{urn:oasis:names:tc:opendocument:xmlns:office:1.0}"
_T = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_TB = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
_D = "{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}"
_PR = "{urn:oasis:names:tc:opendocument:xmlns:presentation:1.0}"
_XL = "{http://www.w3.org/1999/xlink}"


class DocumentExtractionError(RuntimeError):
    def __init__(self, message: str, code: str = "malformed"):
        super().__init__(message)
        self.code = code


def current(doc: dict) -> bool:
    """PPTX cache freshness records the attempted image-text pipeline version."""
    return (doc.get('source') == 'document_pptx' and
            (doc.get('extraction_report') or {}).get('image_ocr_version') == PPTX_IMAGE_OCR_VERSION)


def decode_text(data: bytes) -> tuple[str, str]:
    """Decode supported local text without newline or replacement-character repair."""
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), 'utf-16'
    if b"\0" in data:
        raise DocumentExtractionError("텍스트 파일에 이진 데이터가 있습니다.")
    for encoding in ('utf-8-sig' if data.startswith(b'\xef\xbb\xbf') else 'utf-8', 'cp949'):
        try:
            return data.decode(encoding), encoding
        except UnicodeError:
            pass
    raise DocumentExtractionError("텍스트 인코딩을 읽을 수 없습니다(UTF-8/UTF-16/CP949 지원).")


def _decode(data: bytes) -> str:
    return decode_text(data)[0]


def _guard_zip(path: Path) -> list[str]:
    with open(utils.long_path(path), "rb") as stream:
        if stream.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise DocumentExtractionError("암호화된 Office 문서 또는 구형 바이너리 형식입니다.", "encrypted")
    with zipfile.ZipFile(utils.long_path(path)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ZIP_MEMBERS or sum(e.file_size for e in entries) > MAX_ZIP_BYTES:
            raise DocumentExtractionError("문서 압축 해제 예산을 초과했습니다.", "limit")
        names = set()
        for entry in entries:
            name = entry.filename
            if (name in names or name.startswith(("/", "\\")) or "\\" in name
                    or ".." in name.split("/") or urlsplit(name).scheme
                    or stat.S_ISLNK(entry.external_attr >> 16)):
                raise DocumentExtractionError("모호하거나 패키지 밖을 가리키는 ZIP 경로입니다.")
            names.add(name)
            if entry.flag_bits & 1:
                raise DocumentExtractionError("암호화된 문서입니다.", "encrypted")
            if entry.file_size > MAX_PART_BYTES:
                raise DocumentExtractionError("문서 내부 part 크기 제한을 초과했습니다.", "limit")
        return [entry.filename for entry in entries]


def _xml(data: bytes):
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False, huge_tree=False)
    root = etree.fromstring(data, parser=parser)
    if root.getroottree().docinfo.doctype:
        raise DocumentExtractionError("DTD/외부 엔터티가 있는 XML은 읽지 않습니다.")
    return root


def _zip_target(part: str, target: str) -> str:
    if urlsplit(target).scheme or "\\" in target:
        raise DocumentExtractionError("외부 문서 part를 읽지 않습니다.")
    target = unquote(target)
    result = posixpath.normpath(target.lstrip("/") if target.startswith("/") else
                                posixpath.join(posixpath.dirname(part), target))
    if result == ".." or result.startswith("../"):
        raise DocumentExtractionError("문서 part 경로가 패키지를 벗어납니다.")
    return result


def _relationships(archive, part: str) -> dict:
    name = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels")
    if name not in archive.namelist():
        return {}
    return {entry.get("Id"): _zip_target(part, entry.get("Target", ""))
            for entry in _xml(archive.read(name)) if entry.get("TargetMode") != "External"}


def _save_asset(builder: _Builder, data: bytes, extension: str) -> str:
    if len(data) > MAX_PART_BYTES:
        raise DocumentExtractionError("그림 크기 제한을 초과했습니다.", "limit")
    extension = extension.lower()
    if not re.fullmatch(r"\.[a-z0-9]{1,10}", extension):
        extension = ".bin"
    digest = hashlib.sha256(data).hexdigest()
    builder.assets.mkdir(parents=True, exist_ok=True)
    target = builder.assets / (digest + extension)
    if utils.path_exists(target):
        with open(utils.long_path(target), "rb") as stream:
            if hashlib.sha256(stream.read()).hexdigest() != digest:
                raise DocumentExtractionError("기존 추출 그림이 변경되어 덮어쓰지 않았습니다.")
        return target.name
    fd, temporary = tempfile.mkstemp(prefix=".native-", dir=utils.long_path(builder.assets))
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
            stream.flush()
            os.fsync(stream.fileno())
        if os.name == "nt":
            os.rename(utils.long_path(temporary), utils.long_path(target))
        else:
            os.link(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return target.name


class _Builder:
    def __init__(self, path: Path, asset_dir: str | Path | None, timeout: float, cancel=None,
                 progress=None, default_timeout=False):
        self.started = time.monotonic()
        self.deadline = time.monotonic() + timeout
        self.cancel = cancel
        self.progress = progress
        self.default_timeout = default_timeout
        self.check()
        self.sha1 = store.file_sha1(path)
        self.check()
        self.assets = Path(asset_dir).resolve() if asset_dir is not None else store.figs_dir(self.sha1)
        self.kind = path.suffix.lower().lstrip(".")
        self.doc = schema.Document(
            paper_id=f"sha1:{self.sha1}", source=f"document_{self.kind}",
            source_file=str(path), meta=schema.Meta(title=path.stem)).to_dict()
        self.doc.update(sha1=self.sha1, scope="general_document")
        self.doc["source_map"] = {
            "schema_version": "1.0", "coordinate_space": "logical_document",
            "engine": "native-adapters", "units": [], "blocks": [],
        }
        self.doc["extraction_report"] = {
            "version": 1, "status": "complete", "warnings": [], "pages": [],
            "asset_base": str(self.assets), "source_map_status": "logical_blocks",
        }
        self.block_count = 0

    def check(self):
        if self.cancel is not None and self.cancel():
            from .general_pdf import ExtractionCancelled
            raise ExtractionCancelled("문서 추출이 취소되었습니다.")
        if time.monotonic() > self.deadline:
            raise DocumentExtractionError("문서 추출 제한 시간을 초과했습니다.", "timeout")

    def warn(self, message: str):
        self.doc["extraction_report"]["status"] = "partial"
        if message not in self.doc["extraction_report"]["warnings"]:
            self.doc["extraction_report"]["warnings"].append(message)

    def unit(self, label: str, kind: str, **locator) -> dict:
        self.check()
        index = len(self.doc["source_map"]["units"]) + 1
        unit = {"index": index, "id": f"{self.kind}-u{index:06d}",
                "kind": kind, "label": label, **locator}
        self.doc["source_map"]["units"].append(unit)
        section = {"path": [label] if label else [], "section_type": "other",
                   "paragraphs": [], "source": {"unit": index, "unit_id": unit["id"], **locator}}
        self.doc["body_text"].append(section)
        return section

    def origin(self, section: dict, kind: str, **locator) -> dict:
        self.check()
        self.block_count += 1
        block_id = f"{self.kind}-u{section['source']['unit']:06d}-b{self.block_count:06d}"
        origin = {**section["source"], "block_id": block_id, "kind": kind, **locator}
        self.doc["source_map"]["blocks"].append({"id": block_id, **origin})
        return origin

    def paragraph(self, section: dict, text: str, **locator) -> dict:
        origin = self.origin(section, "paragraph", **locator)
        paragraph = schema.Paragraph(id=origin["block_id"], text=text).__dict__.copy()
        paragraph["source"] = origin
        section["paragraphs"].append(paragraph)
        return paragraph

    def table(self, section: dict, rows: int, cols: int, cells: list[dict],
              caption: str = "", **locator) -> dict:
        if rows * cols > MAX_CELLS:
            raise DocumentExtractionError("표가 허용 셀 수를 초과했습니다.", "limit")
        origin = self.origin(section, "table", **locator)
        table_id = origin["block_id"]
        matrix = [["" for _ in range(cols)] for _ in range(rows)]
        anchors = {(c["row"], c["col"]): c for c in cells}
        parts = ["<table>"]
        for row in range(rows):
            parts.append("<tr>")
            for col in range(cols):
                cell = anchors.get((row, col))
                if cell is None:
                    continue
                text = str(cell.get("text", ""))
                matrix[row][col] = text
                tag = "th" if cell.get("is_header") else "td"
                span = f' rowspan="{cell.get("row_span", 1)}" colspan="{cell.get("col_span", 1)}"'
                parts.append(f"<{tag}{span}>" + html.escape(text).replace("\n", "<br>") + f"</{tag}>")
            parts.append("</tr>")
        parts.append("</table>")
        table_html = "".join(parts)
        merged = any(c.get("row_span", 1) > 1 or c.get("col_span", 1) > 1 for c in cells)
        lines = ["| " + " | ".join(html.escape(x).replace("|", "\\|").replace("\n", "<br>")
                                   for x in row) + " |" for row in matrix]
        if lines:
            lines.insert(1, "| " + " | ".join("---" for _ in range(cols)) + " |")
        buffer = io.StringIO(newline="")
        csv.writer(buffer).writerows(matrix)
        table = {"id": table_id, "caption": caption, "rows": rows, "cols": cols,
                 "cells": cells, "html": table_html, "csv": buffer.getvalue(),
                 "markdown": table_html if merged else "\n".join(lines), "source": origin}
        self.doc["tables"].append(table)
        anchor = self.paragraph(section, "", table_id=table_id)
        anchor["refs_table"] = [table_id]
        return table

    def matrix(self, section: dict, matrix: list[list[str]], **locator):
        rows = len(matrix)
        cols = max((len(row) for row in matrix), default=0)
        if rows * max(1, cols) > MAX_CELLS:
            raise DocumentExtractionError("표가 허용 셀 수를 초과했습니다.", "limit")
        cells = [{"row": r, "col": c, "row_span": 1, "col_span": 1,
                  "is_header": r == 0, "text": str(row[c]) if c < len(row) else ""}
                 for r, row in enumerate(matrix) for c in range(cols)]
        return self.table(section, rows, cols, cells, **locator)


def _docx(builder: _Builder, path: Path):
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise DocumentExtractionError("DOCX 추출에는 python-docx 1.2.0 설치가 필요합니다.", "dependency") from exc
    parts = _guard_zip(path)
    native = Document(utils.long_path(path))
    if native.core_properties.title:
        builder.doc["meta"]["title"] = native.core_properties.title
    section = builder.unit("", "document")

    def images(paragraph, location: dict):
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
            if not relationship:
                builder.warn("외부 연결 그림은 네트워크에서 가져오지 않았습니다.")
                continue
            part = paragraph.part.related_parts.get(relationship)
            if part is None:
                raise DocumentExtractionError("DOCX 그림 part가 없습니다.")
            name = _save_asset(builder, part.blob, Path(str(part.partname)).suffix)
            origin = builder.origin(section, "image", **location)
            figure = {"id": origin["block_id"], "caption": "", "image": name, "source": origin}
            builder.doc["figures"].append(figure)
            anchor = builder.paragraph(section, "", **location)
            anchor["refs_figure"] = [figure["id"]]

    for block_index, block in enumerate(native.iter_inner_content(), 1):
        builder.check()
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ""
            if re.fullmatch(r"Heading [1-9]", style) and block.text.strip():
                section = builder.unit(block.text, "section", paragraph=block_index)
            else:
                builder.paragraph(section, block.text, paragraph=block_index)
            images(block, {"paragraph": block_index})
            if block._p.xpath(".//w:txbxContent | .//m:oMath | .//c:chart"):
                builder.warn("DOCX 글상자·수식·차트는 완전한 구조 추출을 지원하지 않습니다.")
        elif isinstance(block, Table):
            rows, cols = len(block.rows), len(block.columns)
            if rows * cols > MAX_CELLS:
                raise DocumentExtractionError("DOCX 표 셀 수 제한을 초과했습니다.", "limit")
            seen = {}
            for row_index, row in enumerate(block.rows):
                for cell_index, cell in enumerate(row.cells, row.grid_cols_before):
                    builder.check()
                    key = cell._tc
                    if key in seen:
                        anchor = seen[key]
                        anchor["row_span"] = max(anchor["row_span"], row_index - anchor["row"] + 1)
                        anchor["col_span"] = max(anchor["col_span"], cell_index - anchor["col"] + 1)
                        continue
                    text = "\n".join(p.text for p in cell.paragraphs)
                    if cell.tables:
                        text += "\n" + "\n".join(t.text or "" for t in cell._tc.xpath(".//w:tbl//w:t"))
                        builder.warn("DOCX 중첩 표는 셀 텍스트로 보존했습니다(중첩 격자 미지원).")
                    entry = {"row": row_index, "col": cell_index, "row_span": 1, "col_span": 1,
                             "is_header": bool(row._tr.xpath("./w:trPr/w:tblHeader")), "text": text}
                    seen[key] = entry
                    for paragraph in cell.paragraphs:
                        images(paragraph, {"table": block_index, "row": row_index, "col": cell_index})
            builder.table(section, rows, cols, list(seen.values()), table=block_index)
    if any(re.search(r"word/(?:header|footer|footnotes|endnotes|comments)", p) for p in parts):
        builder.warn("DOCX 머리말·꼬리말·각주·미주·주석은 본문 추출 범위 밖입니다.")
    if any("vbaProject" in p or "/embeddings/" in p for p in parts):
        builder.warn("내장 실행체·매크로·OLE 개체는 실행하거나 추출하지 않았습니다.")


def _pptx(builder: _Builder, path: Path):
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:
        raise DocumentExtractionError("PPTX 추출에는 python-pptx 1.0.2 설치가 필요합니다.", "dependency") from exc
    parts = _guard_zip(path)
    native = Presentation(utils.long_path(path))
    builder.doc["source_map"]["ordering"] = "slide_shape_tree"
    builder.doc['extraction_report']['image_ocr_version'] = PPTX_IMAGE_OCR_VERSION
    pictures, slide_sections = [], {}

    def transform(shape, shape_path):
        nodes = shape._element.xpath('./p:spPr/a:xfrm | ./p:grpSpPr/a:xfrm')
        node = nodes[0] if nodes else None
        return {'shape': shape_path, 'rotation_degrees': float(shape.rotation),
                'source_shape_id': str(shape.shape_id),
                'flip_h': node is not None and node.get('flipH') in ('1', 'true'),
                'flip_v': node is not None and node.get('flipV') in ('1', 'true'),
                'components': ({etree.QName(child).localname: dict(child.attrib) for child in node}
                               if node is not None else {})}

    for slide_number, slide in enumerate(native.slides, 1):
        section = builder.unit(f"슬라이드 {slide_number}", "slide", slide=slide_number,
                               width_emu=int(native.slide_width), height_emu=int(native.slide_height))
        slide_sections[slide_number] = section

        def shapes(items, prefix="", groups=()):
            for index, shape in enumerate(items, 1):
                builder.check()
                shape_path = f"{prefix}/{index}" if prefix else str(index)
                location = {"shape": shape_path, 'source_shape_id': str(shape.shape_id),
                            'group_shape_ids': [group['source_shape_id'] for group in groups],
                            "shape_bbox_emu": [int(shape.left), int(shape.top),
                                               int(shape.left + shape.width), int(shape.top + shape.height)]}
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, shape_path, (*groups, transform(shape, shape_path)))
                elif shape.has_table:
                    native_table = shape.table
                    rows, cols = len(native_table.rows), len(native_table.columns)
                    if rows * cols > MAX_CELLS:
                        raise DocumentExtractionError("PPTX 표 셀 수 제한을 초과했습니다.", "limit")
                    cells = []
                    for row in range(rows):
                        for col in range(cols):
                            cell = native_table.cell(row, col)
                            if not cell.is_spanned:
                                cells.append({"row": row, "col": col, "row_span": cell.span_height,
                                              "col_span": cell.span_width, "is_header": row == 0,
                                              "text": cell.text})
                    builder.table(section, rows, cols, cells, **location)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image = shape.image
                    relative = _save_asset(builder, image.blob, "." + image.ext)
                    geometry = {'image_crop_fraction': [shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom],
                                'image_sha256': hashlib.sha256(image.blob).hexdigest(),
                                'shape_transform': transform(shape, shape_path), 'group_transforms': list(groups),
                                'shape_bbox_space': 'shape_parent_emu'}
                    origin = builder.origin(section, "image", **location, **geometry)
                    figure = {"id": origin["block_id"], "caption": shape.name,
                              "image": relative, "source": origin}
                    builder.doc["figures"].append(figure)
                    anchor = builder.paragraph(section, "", **location)
                    anchor["refs_figure"] = [origin["block_id"]]
                    pictures.append({'blob': image.blob, 'section': section, 'figure': figure,
                                     'anchor': anchor, 'location': {**location, **geometry}})
                elif shape.has_text_frame:
                    for pindex, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                        builder.paragraph(section, paragraph.text, paragraph=pindex,
                                          level=paragraph.level, **location)
                else:
                    builder.warn("PPTX 차트·SmartArt·도형·미디어는 텍스트/그림 추출 범위 밖입니다.")
        shapes(slide.shapes)
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None and frame.text.strip():
                notes = builder.unit(f"슬라이드 {slide_number} 발표자 노트", "speaker_notes", slide=slide_number)
                for index, paragraph in enumerate(frame.paragraphs, 1):
                    builder.paragraph(notes, paragraph.text, paragraph=index)
        if builder.progress:
            builder.progress(slide_number, len(native.slides), 'document_native')
    _pptx_image_text(builder, pictures, slide_sections)
    if any("vbaProject" in p or "/embeddings/" in p for p in parts):
        builder.warn("내장 실행체·매크로·OLE 개체는 실행하거나 추출하지 않았습니다.")
    if any("/comments/" in p for p in parts):
        builder.warn("PPTX 검토 주석은 추출하지 않았습니다.")


def _pptx_image_text(builder, pictures, slide_sections):
    from . import image_ocr, local_ocr
    if not pictures:
        builder.doc['extraction_report']['image_ocr'] = {'status': 'no_images', 'images': 0}
        return
    report = {'status': 'processed', 'images': len(pictures), 'unique_images': 0, 'cache_reused': 0,
              'recognized_images': 0, 'empty_images': 0, 'failed_images': [],
              'native_duplicate_lines': 0, 'unlocated_lines': 0, 'text_characters': 0, 'records': []}
    builder.doc['extraction_report']['image_ocr'] = report
    if builder.default_timeout:
        count = len({(hashlib.sha256(item['blob']).hexdigest(),
                      tuple(item['location']['image_crop_fraction'])) for item in pictures})
        builder.deadline = max(builder.deadline, builder.started + min(1800, 30 + 10 * count))
    builder.check()
    provider = local_ocr.discover()
    builder.check()
    if provider is None:
        report.update(status='unavailable', unrecognized_images=len(pictures))
        builder.warn('로컬 OCR 모델이 없어 PPTX 그림 안 글자를 읽지 못했습니다. 원본 그림은 보존했습니다.')
        return
    report.update(engine=provider['engine'], language=provider['language'], engine_identity=provider['identity'])
    images, prepared_cache, slide_images = {}, {}, {number: set() for number in slide_sections}
    prepared_bytes = 0
    for item in pictures:
        builder.check()
        fingerprint = (hashlib.sha256(item['blob']).hexdigest(), tuple(item['location']['image_crop_fraction']))
        try:
            prepared = prepared_cache.get(fingerprint)
            if prepared is None:
                if len(images) >= local_ocr.MAX_BATCH_IMAGES:
                    raise local_ocr.LocalOCRError('OCR 그림 묶음 개수 예산을 초과했습니다.', 'limit')
                prepared = image_ocr.prepare_picture(item['blob'], fingerprint[1])
                if prepared['id'] not in images:
                    if prepared_bytes + len(prepared['png']) > local_ocr.MAX_BATCH_BYTES:
                        raise local_ocr.LocalOCRError('OCR 그림 묶음 PNG 크기 예산을 초과했습니다.', 'limit')
                    prepared_bytes += len(prepared['png'])
                prepared['png'] = images.get(prepared['id'], prepared['png'])
                prepared_cache[fingerprint] = prepared
            item['prepared'] = prepared
            images.setdefault(prepared['id'], prepared['png'])
            slide_images[item['section']['source']['slide']].add(prepared['id'])
        except local_ocr.LocalOCRError as error:
            item['error'] = {'code': error.code, 'error': str(error)}
    report['unique_images'] = len(images)
    report['cache_reused'] = sum('prepared' in item for item in pictures) - len(images)
    completed = set()

    def notify(key=None):
        if key is not None:
            completed.add(key)
        builder.check()
        if builder.progress:
            count = sum(keys.issubset(completed) for keys in slide_images.values())
            builder.progress(count, len(slide_sections), 'document_ocr')

    notify()
    try:
        results = (local_ocr.extract_images(images, provider, cancel=builder.cancel,
                   timeout=max(.001, builder.deadline - time.monotonic()), progress=notify) if images else {})
    except local_ocr.LocalOCRError as error:
        results = {key: {'ok': False, 'code': error.code, 'error': str(error)} for key in images}
    builder.check()
    native_lines = {number: {' '.join(line.split()).casefold() for paragraph in section['paragraphs']
                            for line in paragraph['text'].splitlines() if line.strip()}
                    for number, section in slide_sections.items()}
    text_bytes = sum(len(p['text'].encode('utf-8')) for section in builder.doc['body_text'] for p in section['paragraphs'])
    for item in pictures:
        builder.check()
        prepared = item.get('prepared')
        record = results.get(prepared['id']) if prepared else None
        diagnostic = {'slide': item['section']['source']['slide'], 'shape': item['location']['shape'],
                      'image_id': item['figure']['id'], 'image': item['figure']['image']}
        report['records'].append(diagnostic)
        if not record or not record.get('ok'):
            error = item.get('error') or record or {'code': 'ocr', 'error': 'OCR 결과가 없습니다.'}
            diagnostic.update(status='failed', code=error.get('code'), error=error.get('error'))
            report['failed_images'].append(dict(diagnostic))
            continue
        recovered = []
        for text, origin in image_ocr.lines(record):
            if ' '.join(text.split()).casefold() in native_lines[diagnostic['slide']]:
                report['native_duplicate_lines'] += 1
                continue
            text_bytes += len(text.encode('utf-8'))
            if text_bytes > MAX_TEXT_BYTES:
                raise DocumentExtractionError('PPTX 그림 OCR 본문 크기 예산을 초과했습니다.', 'limit')
            metadata = {key: value for key, value in prepared.items() if key not in ('id', 'png')}
            if origin['geometry_status'] == 'unlocated_ocr':
                report['unlocated_lines'] += 1
            paragraph = builder.paragraph(item['section'], text, image_id=item['figure']['id'],
                                          **{**item['location'], **metadata, **origin})
            item['section']['paragraphs'].pop()
            recovered.append(paragraph)
        paragraphs = item['section']['paragraphs']
        position = next(index for index, paragraph in enumerate(paragraphs) if paragraph is item['anchor']) + 1
        paragraphs[position:position] = recovered
        count = sum(len(p['text']) for p in recovered)
        diagnostic.update(status='recognized' if record['words'] else 'empty', text_characters=count,
                          image_sha256=prepared['image_sha256'])
        report['recognized_images' if record['words'] else 'empty_images'] += 1
        report['text_characters'] += count
    if report['text_characters']:
        builder.warn('PPTX 그림 안 글자를 로컬 OCR로 읽었습니다. 오인식·누락과 숫자는 원본 그림을 대조해 주세요.')
    if report['unlocated_lines']:
        builder.warn('일부 OCR 단어 좌표가 그림 밖으로 나와 해당 글자는 그림 단위 출처만 제공합니다. 원본 그림을 대조해 주세요.')
    if report['failed_images']:
        report['status'] = 'partial'
        builder.warn(f"PPTX 그림 {len(report['failed_images'])}개의 글자를 읽지 못했습니다. 원본 그림은 보존했습니다.")
    completed.update(images)
    notify()


def _xlsx(builder: _Builder, path: Path):
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
    except ImportError as exc:
        raise DocumentExtractionError("XLSX 추출에는 openpyxl 3.1.5 설치가 필요합니다.", "dependency") from exc
    parts = _guard_zip(path)
    native = load_workbook(utils.long_path(path), read_only=True, data_only=False, keep_links=False)
    cached = None
    iterators = []
    try:
        cached = load_workbook(utils.long_path(path), read_only=True, data_only=True, keep_links=False)
        with zipfile.ZipFile(utils.long_path(path)) as archive:
            workbook = _xml(archive.read("xl/workbook.xml"))
            links = _relationships(archive, "xl/workbook.xml")
            sheet_parts = {s.get("name"): links[s.get(_R + "id")]
                           for s in workbook.findall(_S + "sheets/" + _S + "sheet")}
            for sheet in native.worksheets:
                builder.check()
                section = builder.unit(sheet.title, "sheet", sheet=sheet.title)
                if sheet.sheet_state != "visible":
                    builder.warn(f"숨김 시트도 추출했습니다: {sheet.title}")
                part = sheet_parts[sheet.title]
                root = _xml(archive.read(part))
                actual_cells = root.findall(_S + "sheetData/" + _S + "row/" + _S + "c")
                formulas = {c.get("r"): c.find(_S + "f") for c in actual_cells if c.find(_S + "f") is not None}
                rows, cols = sheet.max_row or 0, sheet.max_column or 0
                for actual_cell in actual_cells:
                    row_number, col_number = coordinate_to_tuple(actual_cell.get("r"))
                    rows, cols = max(rows, row_number), max(cols, col_number)
                    if rows * max(1, cols) > MAX_CELLS:
                        raise DocumentExtractionError("XLSX 시트 셀 수 제한을 초과했습니다.", "limit")
                if rows * max(1, cols) > MAX_CELLS:
                    raise DocumentExtractionError("XLSX 시트 셀 수 제한을 초과했습니다.", "limit")
                covered, spans = set(), {}
                for merge in root.findall(_S + "mergeCells/" + _S + "mergeCell"):
                    c1, r1, c2, r2 = range_boundaries(merge.get("ref"))
                    if r2 * c2 > MAX_CELLS:
                        raise DocumentExtractionError("XLSX 병합 범위가 셀 수 제한을 초과했습니다.", "limit")
                    spans[(r1 - 1, c1 - 1)] = (r2 - r1 + 1, c2 - c1 + 1)
                    covered.update((r - 1, c - 1) for r in range(r1, r2 + 1)
                                   for c in range(c1, c2 + 1) if (r, c) != (r1, c1))
                cells, missing_cache = [], 0
                cached_rows = cached[sheet.title].iter_rows(max_row=rows, max_col=cols)
                native_rows = sheet.iter_rows(max_row=rows, max_col=cols)
                iterators.extend([cached_rows, native_rows])
                for rindex, (row, cached_row) in enumerate(zip(native_rows, cached_rows)):
                    builder.check()
                    for cindex, cell in enumerate(row):
                        if (rindex, cindex) in covered:
                            continue
                        value = cell.value
                        coordinate = f"{get_column_letter(cindex + 1)}{rindex + 1}"
                        formula_node = formulas.get(coordinate)
                        formula = None
                        if cell.data_type == "f":
                            formula = value if isinstance(value, str) else "=" + (formula_node.text or "")
                            value = formula
                        if value is None:
                            text = ""
                        elif hasattr(value, "isoformat"):
                            text = value.isoformat()
                        else:
                            text = str(value)
                        row_span, col_span = spans.get((rindex, cindex), (1, 1))
                        entry = {"row": rindex, "col": cindex, "row_span": row_span,
                                 "col_span": col_span, "is_header": rindex == 0, "text": text,
                                 "source": {"sheet": sheet.title, "cell": coordinate}}
                        if formula is not None:
                            cached_value = cached_row[cindex].value
                            if hasattr(cached_value, "isoformat"):
                                cached_value = cached_value.isoformat()
                            entry.update(formula=formula, cached_value=cached_value)
                            if formula_node is not None:
                                entry["formula_attributes"] = dict(formula_node.attrib)
                            missing_cache += cached_value is None
                        cells.append(entry)
                builder.table(section, rows, cols, cells, sheet=sheet.title)
                if missing_cache:
                    builder.warn(f"{sheet.title}: 수식 {missing_cache}개의 계산 캐시가 없습니다(수식 미실행).")
                rels = _relationships(archive, part)
                for drawing in root.findall(_S + "drawing"):
                    drawing_part = rels.get(drawing.get(_R + "id"))
                    if drawing_part is None:
                        builder.warn("외부/누락 drawing은 읽지 않았습니다.")
                        continue
                    drawing_rels = _relationships(archive, drawing_part)
                    drawing_xml = _xml(archive.read(drawing_part))
                    for anchor in drawing_xml:
                        for blip in anchor.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                            image_part = drawing_rels.get(blip.get(_R + "embed"))
                            if image_part is None:
                                builder.warn("외부 연결 그림은 가져오지 않았습니다.")
                                continue
                            relative = _save_asset(builder, archive.read(image_part), Path(image_part).suffix)
                            origin = builder.origin(section, "image", drawing=drawing_part)
                            builder.doc["figures"].append({"id": origin["block_id"], "caption": "",
                                                           "image": relative, "source": origin})
                            marker = builder.paragraph(section, "")
                            marker["refs_figure"] = [origin["block_id"]]
        if any(p.startswith("xl/charts/") or "comments" in p or "pivotTable" in p for p in parts):
            builder.warn("XLSX 차트·주석·피벗 구조는 추출하지 않았습니다.")
        if any("vbaProject" in p or "/embeddings/" in p for p in parts):
            builder.warn("매크로·OLE 개체는 실행하거나 추출하지 않았습니다.")
    finally:
        for iterator in iterators:
            iterator.close()
        native.close()
        if cached is not None:
            cached.close()


def _read_text(path: Path) -> str:
    if path.stat().st_size > MAX_TEXT_BYTES:
        raise DocumentExtractionError("텍스트 크기 제한을 초과했습니다.", "limit")
    with open(utils.long_path(path), "rb") as stream:
        return _decode(stream.read())


def _text(builder: _Builder, path: Path):
    text = _read_text(path)
    if not text.strip():
        raise DocumentExtractionError("추출할 텍스트가 없는 문서입니다.", "empty")
    section = builder.unit("", "document")
    builder.paragraph(section, text, line_start=1, line_end=text.count("\n") + 1)


def _markdown(builder: _Builder, path: Path):
    try:
        from markdown_it import MarkdownIt
    except ImportError as exc:
        raise DocumentExtractionError("Markdown 그림 추출에는 markdown-it-py 설치가 필요합니다.", "dependency") from exc
    text = _read_text(path)
    if not text.strip():
        raise DocumentExtractionError("추출할 텍스트가 없는 문서입니다.", "empty")
    builder.doc["native_markdown"] = text
    parser = MarkdownIt("commonmark", {"html": True}).enable("table")
    env = {}
    tokens = parser.parse(text, env)
    builder.check()
    lines, cursor, section = text.splitlines(keepends=True), 0, None

    def append_raw(start, end):
        nonlocal section
        raw = "".join(lines[start:end])
        if raw.strip():
            if section is None:
                section = builder.unit("", "markdown", line_start=start + 1)
            builder.paragraph(section, raw, line_start=start + 1, line_end=end)

    def load(src):
        relative = Path(unquote(urlsplit(src).path))
        if relative.is_absolute() or not str(relative) or str(relative).startswith("\\"):
            builder.warn("Markdown 절대경로 그림은 수집하지 않았습니다.")
            return None
        original = (path.parent / relative).resolve()
        if not original.is_relative_to(path.parent):
            builder.warn("Markdown 그림이 원본 폴더 밖에 있어 수집하지 않았습니다.")
            return None
        if not original.is_file():
            builder.warn("Markdown 로컬 그림 파일을 찾을 수 없습니다.")
            return None
        if original.stat().st_size > MAX_PART_BYTES:
            raise DocumentExtractionError("Markdown 그림 크기 제한을 초과했습니다.", "limit")
        with open(utils.long_path(original), "rb") as stream:
            return stream.read(MAX_PART_BYTES + 1), original.suffix

    def contains(token, kinds):
        return token.type in kinds or any(contains(child, kinds) for child in token.children or [])

    index = 0
    while index < len(tokens):
        builder.check()
        token = tokens[index]
        if token.level != 0 or token.map is None:
            index += 1
            continue
        start, end = token.map
        stop = index + 1
        if token.nesting == 1:
            while stop < len(tokens):
                if tokens[stop].level == 0 and tokens[stop].nesting == -1:
                    stop += 1
                    break
                stop += 1
        group = tokens[index:stop]
        append_raw(cursor, start)
        has_image = any(contains(item, {"image"}) for item in group)
        has_html = any(contains(item, {"html_block", "html_inline"}) for item in group)
        if has_image or has_html or token.type == "table_open":
            rendered = parser.renderer.render(group, parser.options, env)
            _html_content(builder, rendered, "", load, unit_kind="markdown_block", use_title=False,
                          line_start=start + 1, line_end=end)
            section = None
            if has_image or has_html:
                builder.warn("Markdown 그림/HTML 포함 블록은 구조로 변환했습니다. 일부 인라인 서식/위치는 달라질 수 있으며 원문은 native_markdown에 보존했습니다.")
        else:
            append_raw(start, end)
        cursor, index = end, stop
    append_raw(cursor, len(lines))


def _html_text(node, skip=()) -> str:
    text = node.text or ""
    for child in node:
        tag = child.tag if isinstance(child.tag, str) else ""
        block = tag in ("p", "div", "li", "pre", "blockquote", "tr")
        if tag not in skip:
            if block and text and not text.endswith("\n"):
                text += "\n"
            text += "\n" if tag == "br" else _html_text(child, skip)
            if block and text and not text.endswith("\n"):
                text += "\n"
        text += child.tail or ""
    return text


def _html_table(builder: _Builder, section: dict, element, **location):
    native_rows = element.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
    occupied, cells = set(), []
    rows, cols = len(native_rows), 0
    for rindex, row in enumerate(native_rows):
        builder.check()
        col = 0
        for cell in row.xpath("./th | ./td"):
            while (rindex, col) in occupied:
                col += 1
            try:
                row_span = int(cell.get("rowspan", "1"))
                col_span = int(cell.get("colspan", "1"))
            except ValueError as exc:
                raise DocumentExtractionError("HTML 표의 병합 범위가 올바르지 않습니다.") from exc
            row_span = len(native_rows) - rindex if row_span == 0 else row_span
            if row_span < 1 or col_span < 1 or (rindex + row_span) * (col + col_span) > MAX_CELLS:
                raise DocumentExtractionError("HTML 표 병합/셀 수 제한을 초과했습니다.", "limit")
            cells.append({"row": rindex, "col": col, "row_span": row_span,
                          "col_span": col_span, "is_header": cell.tag == "th",
                          "text": _html_text(cell).strip("\n")})
            occupied.update((r, c) for r in range(rindex, rindex + row_span)
                            for c in range(col, col + col_span))
            rows, cols = max(rows, rindex + row_span), max(cols, col + col_span)
            col += col_span
            if cell.xpath(".//table"):
                builder.warn("HTML 중첩 표는 셀 텍스트로 보존했습니다(중첩 격자 미지원).")
    if rows * max(1, cols) > MAX_CELLS:
        raise DocumentExtractionError("HTML 표 셀 수 제한을 초과했습니다.", "limit")
    for row in range(rows):
        for col in range(cols):
            if (row, col) not in occupied:
                cells.append({"row": row, "col": col, "row_span": 1, "col_span": 1,
                              "is_header": False, "text": ""})
    builder.table(section, rows, cols, cells, **location)


def _html_content(builder: _Builder, data: str, label: str, loader, *,
                  unit_kind="html_section", use_title=True, local_schemes=(), **unit_locator):
    from lxml import html as lhtml
    root = lhtml.fromstring(data, parser=lhtml.HTMLParser(no_network=True))
    for active in root.xpath("//script | //style | //iframe | //object | //embed | //template"):
        active.drop_tree()
    titles = root.xpath("//title/text()")
    if use_title and titles and len(builder.doc["source_map"]["units"]) == 0:
        builder.doc["meta"]["title"] = titles[0]
    body = root.find("body")
    if body is None:
        body = root
    section = builder.unit(label, unit_kind, **unit_locator)

    def image(node):
        src = node.get("src", "")
        if src.startswith("data:"):
            match = re.fullmatch(r"data:image/(png|jpeg|gif|webp);base64,(.*)", src, re.S | re.I)
            if not match or len(match.group(2)) > MAX_PART_BYTES * 4 // 3 + 4:
                builder.warn("지원 범위 밖의 인라인 그림은 추출하지 않았습니다.")
                return
            blob = base64.b64decode(match.group(2), validate=True)
            extension = "." + match.group(1).lower()
        elif (urlsplit(src).scheme and urlsplit(src).scheme not in local_schemes) or src.startswith("//"):
            builder.warn("외부 연결 그림은 네트워크에서 가져오지 않았습니다.")
            return
        else:
            loaded = loader(src)
            if loaded is None:
                builder.warn("지원 범위 밖의 그림 링크는 추출하지 않았습니다.")
                return
            blob, extension = loaded
        name = _save_asset(builder, blob, extension)
        origin = builder.origin(section, "image", dom_path=node.getroottree().getpath(node))
        builder.doc["figures"].append({"id": origin["block_id"], "caption": node.get("alt", ""),
                                       "image": name, "source": origin})
        marker = builder.paragraph(section, "")
        marker["refs_figure"] = [origin["block_id"]]

    def walk(node):
        nonlocal section
        if not isinstance(node.tag, str):
            return
        builder.check()
        tag = node.tag.lower()
        location = {"dom_path": node.getroottree().getpath(node)}
        if re.fullmatch(r"h[1-6]", tag):
            section = builder.unit("".join(node.itertext()), "heading", **unit_locator, **location)
        elif tag == "table":
            _html_table(builder, section, node, **location)
            for item in node.xpath(".//img"):
                image(item)
        elif tag == "img":
            image(node)
        elif tag in ("p", "pre", "li", "dt", "dd", "figcaption"):
            text = _html_text(node, skip=("table", "ul", "ol"))
            if text.strip():
                builder.paragraph(section, text, **location)
            for item in node.xpath(".//img"):
                image(item)
            for child in node:
                if isinstance(child.tag, str) and child.tag in ("table", "ul", "ol"):
                    walk(child)
        else:
            if tag in ("head", "script", "style"):
                return
            if node.text and node.text.strip():
                builder.paragraph(section, node.text, **location)
            for child in node:
                walk(child)
                if child.tail and child.tail.strip():
                    builder.paragraph(section, child.tail, **location)
    walk(body)


def _html(builder: _Builder, path: Path):
    def load(src):
        relative = unquote(urlsplit(src).path)
        original = (path.parent / relative).resolve(strict=True)
        if not original.is_relative_to(path.parent):
            raise DocumentExtractionError("HTML 그림이 원본 폴더 밖에 있습니다.")
        if original.stat().st_size > MAX_PART_BYTES:
            raise DocumentExtractionError("HTML 그림 크기 제한을 초과했습니다.", "limit")
        with open(utils.long_path(original), "rb") as stream:
            return stream.read(), original.suffix
    _html_content(builder, _read_text(path), "", load)


def _epub(builder: _Builder, path: Path):
    # EPUB OCF container + OPF spine, W3C EPUB 3.3 §4/§5.7.
    parts = _guard_zip(path)
    with zipfile.ZipFile(utils.long_path(path)) as archive:
        if "META-INF/rights.xml" in parts:
            raise DocumentExtractionError("권리 보호(DRM) EPUB은 지원하지 않습니다.", "encrypted")
        if "META-INF/encryption.xml" in parts:
            encryption = _xml(archive.read("META-INF/encryption.xml"))
            targets = [e.get("URI", "") for e in encryption.findall(".//{*}CipherReference")]
            if not targets or any(Path(urlsplit(t).path).suffix.lower() not in
                                  (".ttf", ".otf", ".woff", ".woff2") for t in targets):
                raise DocumentExtractionError("본문/그림이 암호화된 EPUB은 지원하지 않습니다.", "encrypted")
            builder.warn("EPUB 난독화/암호화 글꼴은 사용하지 않고 본문만 추출했습니다.")
        container = _xml(archive.read("META-INF/container.xml"))
        rootfile = container.find(".//{*}rootfile")
        if rootfile is None or not rootfile.get("full-path"):
            raise DocumentExtractionError("EPUB package 주소가 없습니다.")
        package_part = _zip_target("", rootfile.get("full-path"))
        package = _xml(archive.read(package_part))
        title = package.findtext(".//{http://purl.org/dc/elements/1.1/}title")
        if title:
            builder.doc["meta"]["title"] = title
        builder.doc["meta"]["authors"] = package.xpath("//*[local-name()='creator']/text()")
        manifest = {entry.get("id"): entry for entry in package.findall("{*}manifest/{*}item")}
        spine = package.findall("{*}spine/{*}itemref")
        if not spine:
            raise DocumentExtractionError("EPUB spine 읽기 순서가 없습니다.")
        for index, reference in enumerate(spine, 1):
            builder.check()
            item = manifest.get(reference.get("idref"))
            if item is None or not item.get("href"):
                raise DocumentExtractionError("EPUB spine의 본문 item이 없습니다.")
            if item.get("media-type") not in ("application/xhtml+xml", "text/html"):
                builder.warn("EPUB spine의 비HTML 콘텐츠는 본문 추출 범위 밖입니다.")
                continue
            chapter_part = _zip_target(package_part, urlsplit(item.get("href")).path)

            def load(src):
                image_part = _zip_target(chapter_part, urlsplit(src).path)
                return archive.read(image_part), Path(image_part).suffix

            _html_content(builder, _decode(archive.read(chapter_part)), f"Chapter {index}", load,
                          unit_kind="chapter", use_title=False, spine_index=index, entry=chapter_part)


def _odf_count(value: str | None) -> int:
    number = int(value or "1")
    if number < 1 or number > MAX_CELLS:
        raise DocumentExtractionError("ODF 반복/병합 수가 제한을 초과했습니다.", "limit")
    return number


def _odf_positive(value: str | None) -> int:
    number = int(value or "1")
    if number < 1:
        raise DocumentExtractionError("ODF 반복 수는 양수여야 합니다.")
    return number


def _odf_has_value(cell, builder: _Builder) -> bool:
    return (cell.tag == _TB + "covered-table-cell" or bool(_odf_text(cell, builder).strip())
            or any(cell.get(_O + key) is not None for key in
            ("value", "string-value", "date-value", "time-value", "boolean-value"))
            or bool(cell.get(_TB + "formula"))
            or cell.get(_TB + "number-rows-spanned", "1") != "1"
            or cell.get(_TB + "number-columns-spanned", "1") != "1"
            or any(cell.find(".//" + tag) is not None for tag in
                   (_D + "image", _T + "note", _O + "annotation", _TB + "table")))


def _odf_charge(builder: _Builder, amount: int):
    builder.check()
    processed = getattr(builder, "odf_text_bytes", 0) + amount
    if processed > MAX_TEXT_BYTES:
        raise DocumentExtractionError("ODF 누적 텍스트 처리 예산을 초과했습니다.", "limit")
    builder.odf_text_bytes = processed


def _odf_text(node, builder: _Builder) -> str:
    chunks = []

    def append(text):
        _odf_charge(builder, len(text.encode("utf-8")))
        chunks.append(text)

    def walk(current):
        builder.check()
        if current.tag in (_D + "frame", _D + "image", _O + "annotation", _O + "binary-data"):
            return
        if current.tag == _T + "note":
            append("[" + (current.findtext(_T + "note-citation") or "") + "]")
            return
        if current.tag == _T + "s":
            count = _odf_count(current.get(_T + "c"))
            _odf_charge(builder, count)
            chunks.append(" " * count)
            return
        if current.tag in (_T + "tab", _T + "line-break"):
            append("\t" if current.tag == _T + "tab" else "\n")
            return
        append(current.text or "")
        for child in current:
            walk(child)
            append(child.tail or "")
    walk(node)
    return "".join(chunks)


def _odf_row_nodes(node):
    for child in node:
        if child.tag == _TB + "table-row":
            yield child
        elif child.tag in (_TB + "table-header-rows", _TB + "table-rows", _TB + "table-row-group"):
            yield from _odf_row_nodes(child)


def _odf_table(builder: _Builder, section: dict, element):
    row_index, rows, cols = 0, 0, 0
    cells, occupied = [], set()
    native_rows = list(_odf_row_nodes(element))
    omitted_rows = omitted_columns = 0
    if builder.kind == "ods":
        last = -1
        for index, row in enumerate(native_rows):
            builder.check()
            if any(_odf_has_value(cell, builder) for cell in row
                   if cell.tag in (_TB + "table-cell", _TB + "covered-table-cell")):
                last = index
        omitted_rows = sum(_odf_positive(r.get(_TB + "number-rows-repeated")) for r in native_rows[last + 1:])
        native_rows = native_rows[:last + 1]
    for row in native_rows:
        repeats = _odf_count(row.get(_TB + "number-rows-repeated"))
        if (row_index + repeats) * max(1, cols) > MAX_CELLS:
            raise DocumentExtractionError("ODF 표 행 수 제한을 초과했습니다.", "limit")
        native_cells = [cell for cell in row if cell.tag in (_TB + "table-cell", _TB + "covered-table-cell")]
        if builder.kind == "ods":
            last = max((i for i, cell in enumerate(native_cells) if _odf_has_value(cell, builder)), default=-1)
            omitted_columns += repeats * sum(_odf_positive(c.get(_TB + "number-columns-repeated"))
                                             for c in native_cells[last + 1:])
            native_cells = native_cells[:last + 1]
        for _ in range(repeats):
            builder.check()
            col = 0
            for cell in native_cells:
                if cell.tag not in (_TB + "table-cell", _TB + "covered-table-cell"):
                    continue
                count = _odf_count(cell.get(_TB + "number-columns-repeated"))
                if (row_index + 1) * (col + count) > MAX_CELLS:
                    raise DocumentExtractionError("ODF 표 셀 수 제한을 초과했습니다.", "limit")
                if cell.tag == _TB + "covered-table-cell":
                    col += count
                    continue
                row_span = _odf_count(cell.get(_TB + "number-rows-spanned"))
                col_span = _odf_count(cell.get(_TB + "number-columns-spanned"))
                if (row_index + row_span) * (col + count + col_span - 1) > MAX_CELLS:
                    raise DocumentExtractionError("ODF 표 병합 범위 제한을 초과했습니다.", "limit")
                text = "\n".join(_odf_text(p, builder) for p in cell.findall(_T + "p"))
                if not text:
                    text = next((cell.get(_O + key) for key in
                                 ("string-value", "value", "date-value", "time-value", "boolean-value")
                                 if cell.get(_O + key) is not None), "")
                _odf_charge(builder, len(text.encode("utf-8")) * count)
                for _ in range(count):
                    if (row_index, col) in occupied:
                        raise DocumentExtractionError("ODF 병합 셀의 covered-cell 구조가 겹칩니다.")
                    entry = {"row": row_index, "col": col, "row_span": row_span, "col_span": col_span,
                             "is_header": any(a.tag == _TB + "table-header-rows" for a in cell.iterancestors()),
                             "text": text, "source": {"cell": f"R{row_index + 1}C{col + 1}"}}
                    if cell.get(_TB + "formula"):
                        entry.update(formula=cell.get(_TB + "formula"), cached_value=text,
                                     value_type=cell.get(_O + "value-type"))
                    cells.append(entry)
                    occupied.update((r, c) for r in range(row_index, row_index + row_span)
                                    for c in range(col, col + col_span))
                    rows, cols = max(rows, row_index + row_span), max(cols, col + col_span)
                    col += 1
                if cell.find(_TB + "table") is not None:
                    builder.warn("ODF 중첩 표는 전체 격자 추출 범위 밖입니다.")
            cols = max(cols, col)
            row_index += 1
            rows = max(rows, row_index)
    if rows * max(1, cols) > MAX_CELLS:
        raise DocumentExtractionError("ODF 표 셀 수 제한을 초과했습니다.", "limit")
    for row in range(rows):
        for col in range(cols):
            if (row, col) not in occupied:
                cells.append({"row": row, "col": col, "row_span": 1, "col_span": 1,
                              "is_header": False, "text": ""})
    builder.table(section, rows, cols, cells, caption=element.get(_TB + "name", ""),
                  xml_path=element.getroottree().getpath(element),
                  omitted_blank_capacity={"rows": omitted_rows, "columns": omitted_columns})


def _odf(builder: _Builder, path: Path):
    parts = _guard_zip(path)
    with zipfile.ZipFile(utils.long_path(path)) as archive:
        if "META-INF/manifest.xml" in parts:
            manifest = _xml(archive.read("META-INF/manifest.xml"))
            if manifest.find(".//{*}encryption-data") is not None:
                raise DocumentExtractionError("암호화된 ODF 문서는 지원하지 않습니다.", "encrypted")
        root = _xml(archive.read("content.xml"))
        body_kind = {".odt": "text", ".ods": "spreadsheet", ".odp": "presentation"}[path.suffix.lower()]
        body = root.find(_O + "body/" + _O + body_kind)
        if body is None:
            raise DocumentExtractionError("ODF 확장자와 본문 형식이 일치하지 않습니다.")
        section = builder.unit("", "document")
        slide_number = 0
        if "meta.xml" in parts:
            title = _xml(archive.read("meta.xml")).findtext(".//{http://purl.org/dc/elements/1.1/}title")
            if title:
                builder.doc["meta"]["title"] = title

        def picture(node):
            href = node.get(_XL + "href")
            if href:
                if urlsplit(href).scheme or href.startswith("//"):
                    builder.warn("ODF 외부 연결 그림은 가져오지 않았습니다.")
                    return
                part = _zip_target("content.xml", urlsplit(href).path)
                data, extension = archive.read(part), Path(part).suffix
            else:
                encoded = node.findtext(_O + "binary-data") or ""
                if len(encoded) > MAX_PART_BYTES * 4 // 3 + 4096:
                    raise DocumentExtractionError("ODF 인라인 그림 크기 제한을 초과했습니다.", "limit")
                data = base64.b64decode(re.sub(r"\s+", "", encoded), validate=True)
                extension = "." + node.get(_D + "mime-type", "image/bin").split("/")[-1]
            relative = _save_asset(builder, data, extension)
            origin = builder.origin(section, "image", xml_path=node.getroottree().getpath(node))
            builder.doc["figures"].append({"id": origin["block_id"], "caption": "",
                                           "image": relative, "source": origin})
            marker = builder.paragraph(section, "")
            marker["refs_figure"] = [origin["block_id"]]

        def inline_resources(node):
            for child in node:
                if child.tag in (_D + "frame", _D + "image", _T + "note", _O + "annotation"):
                    walk(child)
                else:
                    inline_resources(child)

        def walk(node):
            nonlocal section, slide_number
            builder.check()
            location = {"xml_path": node.getroottree().getpath(node)}
            if node.tag == _D + "page" and builder.kind == "odp":
                slide_number += 1
                section = builder.unit(node.get(_D + "name") or f"Slide {slide_number}", "slide",
                                       slide=slide_number, **location)
                for child in node:
                    walk(child)
            elif node.tag == _PR + "notes":
                saved = section
                section = builder.unit("발표자 노트", "speaker_notes", slide=slide_number, **location)
                for child in node:
                    walk(child)
                section = saved
            elif node.tag == _T + "h":
                section = builder.unit(_odf_text(node, builder), "heading", **location)
            elif node.tag == _T + "p":
                text = _odf_text(node, builder)
                if text.strip():
                    builder.paragraph(section, text, **location)
                inline_resources(node)
            elif node.tag == _TB + "table":
                if builder.kind == "ods" and node.getparent() is body:
                    name = node.get(_TB + "name", "Sheet")
                    section = builder.unit(name, "sheet", sheet=name, **location)
                _odf_table(builder, section, node)
                inline_resources(node)
            elif node.tag == _D + "image":
                picture(node)
            elif node.tag in (_T + "note", _O + "annotation"):
                saved = section
                citation = node.findtext(_T + "note-citation") or ""
                kind = "footnote" if node.tag == _T + "note" else "comment"
                section = builder.unit(("각주 " + citation).strip() if kind == "footnote" else "주석", kind,
                                       note_id=node.get(_T + "id"), **location)
                target = node.find(_T + "note-body") if kind == "footnote" else node
                if target is not None:
                    for child in target:
                        walk(child)
                section = saved
            elif node.tag in (_O + "scripts", _T + "tracked-changes", _D + "object", _D + "object-ole", _D + "plugin"):
                builder.warn("ODF 스크립트·변경이력·내장 개체는 실행하거나 본문에 합치지 않았습니다.")
            else:
                for child in node:
                    walk(child)
        walk(body)


def _mail_body_text(part) -> str:
    data = part.get_payload(decode=True) or b""
    if len(data) > MAX_TEXT_BYTES:
        raise DocumentExtractionError("메일 본문 크기 제한을 초과했습니다.", "limit")
    charset = part.get_content_charset()
    if charset:
        try:
            return data.decode(charset)
        except (LookupError, UnicodeError) as exc:
            raise DocumentExtractionError(f"메일 본문 인코딩을 읽을 수 없습니다: {charset}") from exc
    return _decode(data)


def _mail_message(builder: _Builder, message, index: int):
    builder.check()
    if index > MAX_MAIL_MESSAGES:
        raise DocumentExtractionError("메일 개수 제한을 초과했습니다.", "limit")
    parts, stack, attachments, embedded = [], [(message, 0)], [], {}
    while stack:
        builder.check()
        part, depth = stack.pop()
        if depth > 30 or len(parts) >= MAX_MAIL_PARTS:
            raise DocumentExtractionError("메일 MIME 깊이/파트 수 제한을 초과했습니다.", "limit")
        parts.append(part)
        content_type = part.get_content_type()
        if (content_type == "multipart/encrypted" or content_type == "application/pgp-encrypted"
                or content_type in ("application/pkcs7-mime", "application/x-pkcs7-mime")
                and part.get_param("smime-type") != "signed-data"):
            raise DocumentExtractionError("암호화된 메일은 복호화하지 않습니다.", "encrypted")
        disposition = part.get_content_disposition()
        filename = part.get_filename()
        if filename or disposition == "attachment":
            attachments.append({"filename": filename or "", "content_type": content_type,
                                "disposition": disposition or "inline", "extracted": False})
        if disposition == "attachment":
            continue
        if part.get_content_maintype() == "image":
            for key in (str(part.get("Content-ID", "")).strip("<>"), str(part.get("Content-Location", ""))):
                if key:
                    embedded[key] = part
        if part.is_multipart() and content_type != "message/rfc822":
            stack.extend((child, depth + 1) for child in reversed(list(part.iter_parts())))
    metadata = {"index": index, "message_id": str(message.get("Message-ID", "")),
                "subject": str(message.get("Subject", "")), "from": str(message.get("From", "")),
                "to": str(message.get("To", "")), "cc": str(message.get("Cc", "")),
                "date": str(message.get("Date", "")), "attachments": attachments}
    builder.doc.setdefault("mail_messages", []).append(metadata)
    if builder.kind == "eml" and metadata["subject"]:
        builder.doc["meta"]["title"] = metadata["subject"]
    body = message.get_body(preferencelist=("html", "plain"))
    locator = {"message_index": index, "message_id": metadata["message_id"]}
    label = metadata["subject"] or f"Message {index}"
    header_unit = builder.unit(label, "email_headers", **locator)
    header_text = "\n".join(f"{key.title()}: {metadata[key]}" for key in ("from", "to", "cc", "date") if metadata[key])
    if header_text:
        builder.paragraph(header_unit, header_text)
    if body is not None and body.get_content_type() == "text/html":
        def load(src):
            key = unquote(src[4:]) if src.startswith("cid:") else src
            resource = embedded.get(key.strip("<>"))
            if resource is None:
                return None
            data = resource.get_payload(decode=True) or b""
            if len(data) > MAX_PART_BYTES:
                raise DocumentExtractionError("메일 그림 크기 제한을 초과했습니다.", "limit")
            for attachment in attachments:
                if attachment["filename"] == (resource.get_filename() or ""):
                    attachment["extracted"] = True
            extension = "." + resource.get_content_subtype().split("+")[0]
            return data, extension
        _html_content(builder, _mail_body_text(body), label, load, use_title=False,
                      unit_kind="email_message", local_schemes=("cid",), **locator)
    elif body is not None and body.get_content_type() == "text/plain":
        section = builder.unit(label, "email_message", **locator)
        builder.paragraph(section, _mail_body_text(body))
    else:
        builder.warn(f"{index}번째 메일에 지원하는 본문이 없습니다.")
    if any(not attachment["extracted"] for attachment in attachments):
        builder.warn("메일 첨부는 목록으로 보존했으며 첨부 문서를 실행하거나 재귀 추출하지 않았습니다.")
    if any(part.defects for part in parts):
        builder.warn("메일 MIME/전송 인코딩 진단이 있어 원본 확인이 필요합니다.")


def _parse_mail(builder: _Builder, data: bytes):
    if len(data) > MAX_TEXT_BYTES:
        raise DocumentExtractionError("단일 메일 크기 제한을 초과했습니다.", "limit")
    count = 0
    factory = email_policy.default.message_factory

    def create_message(*args, **kwargs):
        nonlocal count
        builder.check()
        count += 1
        if count > MAX_MAIL_PARTS:
            raise DocumentExtractionError("메일 MIME 파트 수 제한을 초과했습니다.", "limit")
        return factory(*args, **kwargs)

    parser = BytesParser(policy=email_policy.default.clone(message_factory=create_message))
    message = parser.parsebytes(data)
    builder.check()
    return message


def _eml(builder: _Builder, path: Path):
    if path.stat().st_size > MAX_TEXT_BYTES:
        raise DocumentExtractionError("단일 메일 크기 제한을 초과했습니다.", "limit")
    with open(utils.long_path(path), "rb") as stream:
        message = _parse_mail(builder, stream.read(MAX_TEXT_BYTES + 1))
    _mail_message(builder, message, 1)


def _mbox(builder: _Builder, path: Path):
    # Read-only mboxrd; Content-Length based mboxcl/cl2 cannot be inferred safely.
    builder.doc["source_map"]["mbox_variant"] = "mboxrd"
    builder.warn("MBOX는 mboxrd 구분/본문 escape 규칙으로 읽었습니다. Content-Length 변종은 지원하지 않습니다.")
    current, seen_envelope, in_body, count = bytearray(), False, False, 0

    def emit():
        nonlocal count
        if not current:
            return
        count += 1
        if count > MAX_MAIL_MESSAGES:
            raise DocumentExtractionError("메일 개수 제한을 초과했습니다.", "limit")
        builder.check()
        message = _parse_mail(builder, bytes(current))
        if message.get("Content-Length") is not None:
            raise DocumentExtractionError("Content-Length 기반 MBOX 변종은 지원하지 않습니다.", "unsupported")
        _mail_message(builder, message, count)

    with open(utils.long_path(path), "rb") as stream:
        while True:
            builder.check()
            line = stream.readline(MAX_TEXT_BYTES + 1)
            if not line:
                break
            if len(line) > MAX_TEXT_BYTES:
                raise DocumentExtractionError("MBOX 행 크기 제한을 초과했습니다.", "limit")
            if line.startswith(b"From "):
                emit()
                current.clear()
                seen_envelope, in_body = True, False
                continue
            if not seen_envelope:
                if line.strip():
                    raise DocumentExtractionError("MBOX의 From 구분 행이 없습니다.")
                continue
            if in_body and re.match(br"^>+From ", line):
                line = line[1:]
            if len(current) + len(line) > MAX_TEXT_BYTES:
                raise DocumentExtractionError("MBOX 단일 메일 크기 제한을 초과했습니다.", "limit")
            current.extend(line)
            if line in (b"\n", b"\r\n"):
                in_body = True
    emit()


def _run_rhwp(builder: _Builder, args: list[str], cwd: Path) -> bytes:
    builder.check()
    flags = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0
    with tempfile.TemporaryFile() as stdout, tempfile.TemporaryFile() as stderr:
        try:
            options = dict(stdin=subprocess.DEVNULL, stdout=stdout, stderr=stderr,
                           shell=False, cwd=str(cwd), creationflags=flags)
            if builder.cancel is None:
                result = subprocess.run(args, timeout=max(0.001, builder.deadline - time.monotonic()), **options)
                returncode = result.returncode
            else:
                process = subprocess.Popen(args, **options)
                try:
                    while process.poll() is None:
                        builder.check()
                        time.sleep(0.05)
                    returncode = process.returncode
                finally:
                    if process.poll() is None:
                        process.kill()
                        process.wait(timeout=5)
        except subprocess.TimeoutExpired as exc:
            raise DocumentExtractionError("rHWP 추출 제한 시간을 초과했습니다.", "timeout") from exc
        builder.check()
        stdout.seek(0)
        stderr.seek(0)
        output, diagnostic = stdout.read(65537), stderr.read(65537)
    if len(output) > 65536:
        raise DocumentExtractionError("rHWP 응답 크기 제한을 초과했습니다.", "limit")
    message = diagnostic.decode("utf-8", "replace").strip()
    if returncode:
        protected = re.search(r"암호|password|encrypt|DRM|Fasoo|SoftCamp", message, re.I)
        raise DocumentExtractionError(f"rHWP 추출 실패: {message[:4000] or returncode}",
                                      "encrypted" if protected else "unsupported")
    if message:
        builder.warn(f"rHWP 진단: {message[:4000]}")
    return output


def _doclang_text(node) -> str:
    if node.tag in ("location", "custom", "hwp_prop", "src"):
        return ""
    text = node.text or ""
    for child in node:
        text += _doclang_text(child) + (child.tail or "")
    return "$" + text + "$" if node.tag == "formula" else text


def _doclang_grid(builder: _Builder, element) -> dict:
    anchors, grid, current = [], {}, None
    row, col, max_col = 0, 0, 0
    for child in element:
        builder.check()
        tag = child.tag
        if tag in ("fcel", "ecel", "ched"):
            current = {"row": row, "col": col, "row_span": 1, "col_span": 1,
                       "is_header": tag == "ched", "text": ""}
            anchors.append(current)
            grid[(row, col)] = current
            col += 1
        elif tag in ("lcel", "ucel", "xcel"):
            key = (row, col - 1) if tag == "lcel" else (row - 1, col)
            anchor = grid.get(key)
            if anchor is None:
                raise DocumentExtractionError("rHWP 표의 병합 연결이 올바르지 않습니다.")
            anchor["row_span"] = max(anchor["row_span"], row - anchor["row"] + 1)
            anchor["col_span"] = max(anchor["col_span"], col - anchor["col"] + 1)
            grid[(row, col)] = anchor
            current = None
            col += 1
        elif tag == "nl":
            max_col = max(max_col, col)
            row, col, current = row + 1, 0, None
        elif current is not None and tag == "table":
            nested = _doclang_grid(builder, child)
            current.setdefault("nested", []).append(nested)
            current["text"] += "\n".join(c["text"] for c in nested["cells"])
            builder.warn("한글 중첩 표는 JSON 격자를 보존하고 Markdown 셀에는 텍스트를 표시했습니다.")
        elif current is not None and tag in ("text", "heading", "formula", "list"):
            value = _doclang_text(child)
            current["text"] += ("\n" if current["text"] else "") + value
        max_col = max(max_col, col)
        if (row + 1) * max(1, max_col) > MAX_CELLS:
            raise DocumentExtractionError("rHWP 표 셀 수 제한을 초과했습니다.", "limit")
    return {"rows": row + bool(col), "cols": max_col, "cells": anchors}


def _doclang(builder: _Builder, root, assets: Path):
    section = builder.unit("", "doclang_section")

    def picture(node, **locator):
        src = node.find("src")
        if src is None or not src.get("uri"):
            raise DocumentExtractionError("rHWP 그림 자원 주소가 없습니다.")
        original = Path(src.get("uri")).resolve(strict=True)
        if not original.is_relative_to(assets):
            raise DocumentExtractionError("rHWP 그림 자원이 임시 자원 폴더 밖에 있습니다.")
        if original.stat().st_size > MAX_PART_BYTES:
            raise DocumentExtractionError("rHWP 그림 크기 제한을 초과했습니다.", "limit")
        relative = _save_asset(builder, original.read_bytes(), original.suffix)
        origin = builder.origin(section, "image", doclang_path=node.getroottree().getpath(node), **locator)
        builder.doc["figures"].append({"id": origin["block_id"], "caption": "",
                                       "image": relative, "source": origin})
        marker = builder.paragraph(section, "")
        marker["refs_figure"] = [origin["block_id"]]

    def walk(node):
        nonlocal section
        builder.check()
        location = {"doclang_path": node.getroottree().getpath(node)}
        if node.tag == "heading":
            section = builder.unit(_doclang_text(node), "heading", **location)
        elif node.tag in ("text", "formula"):
            builder.paragraph(section, _doclang_text(node), **location)
        elif node.tag == "table":
            model = _doclang_grid(builder, node)
            caption = node.find("caption")
            builder.table(section, model["rows"], model["cols"], model["cells"],
                          caption=_doclang_text(caption) if caption is not None else "", **location)
            for image in node.findall(".//picture"):
                picture(image, table_path=location["doclang_path"])
        elif node.tag == "picture":
            picture(node)
        elif node.tag in ("footnote", "page_header", "page_footer"):
            saved = section
            labels = {"footnote": "각주", "page_header": "머리말", "page_footer": "꼬리말"}
            section = builder.unit(labels[node.tag], node.tag, **location)
            for child in node:
                walk(child)
            section = saved
        elif node.tag == "page_break":
            builder.origin(section, "page_break", **location)
        elif node.tag in ("doclang", "list", "group"):
            for child in node:
                walk(child)
    walk(root)


def _rhwp(builder: _Builder, path: Path):
    configured = os.environ.get("PUBNEXUS_RHWP")
    if configured and not Path(configured).is_file():
        raise DocumentExtractionError("rHWP 실행 파일이 없습니다. PUBNEXUS_RHWP 경로를 확인해 주세요.", "dependency")
    try:
        from .preview import find_engine
        engine = find_engine("rhwp")
    except ImportError:
        engine = configured or shutil.which("rhwp")
    if engine is None:
        raise DocumentExtractionError("HWP/HWPX 추출용 rHWP가 없습니다. PUBNEXUS_RHWP를 설정해 주세요.", "dependency")
    if path.suffix.lower() == ".hwpx":
        _guard_zip(path)
    with tempfile.TemporaryDirectory(prefix="pubnexus-rhwp-") as folder:
        work = Path(folder)
        output, assets = work / "document.xml", work / "assets"
        version = _run_rhwp(builder, [str(engine), "--version"], work).decode("utf-8", "replace").strip()
        envelope = json.loads(_run_rhwp(builder, [str(engine), "export-doclang", str(path),
                                               "-o", str(output), "--assets-dir", str(assets), "--json"], work))
        if (not isinstance(envelope, dict) or envelope.get("format") != "doclang"
                or envelope.get("doclangVersion") != "0.6"):
            raise DocumentExtractionError("지원하지 않는 rHWP DocLang 응답입니다.", "unsupported")
        if not output.is_file() or output.stat().st_size > MAX_PART_BYTES:
            raise DocumentExtractionError("rHWP 의미 XML이 없거나 크기 제한을 초과했습니다.", "limit")
        asset_files = list(assets.rglob("*")) if assets.exists() else []
        asset_files = [p for p in asset_files if p.is_file()]
        if len(asset_files) != envelope.get("assetCount") or sum(p.stat().st_size for p in asset_files) > MAX_ZIP_BYTES:
            raise DocumentExtractionError("rHWP 그림 자원 수 또는 크기가 응답과 다릅니다.", "limit")
        root = _xml(output.read_bytes())
        if root.tag != "doclang" or root.get("version") != "0.6":
            raise DocumentExtractionError("지원하지 않는 rHWP 의미 XML 버전입니다.", "unsupported")
        builder.doc["source_map"].update(engine="rHWP DocLang 0.6", engine_version=version,
                                          logical_origin="doclang_xpath", page_accuracy="none")
        loss = envelope.get("lossCount")
        if not isinstance(loss, int) or loss < 0:
            raise DocumentExtractionError("rHWP 손실 진단이 올바르지 않습니다.")
        builder.doc["extraction_report"]["loss_count"] = loss
        if loss:
            builder.warn(f"rHWP 의미 변환 손실 보고 {loss}건 — 원본 대조가 필요합니다.")
        _doclang(builder, root, assets.resolve())


def _finish(builder: _Builder) -> dict:
    builder.check()
    doc = builder.doc
    has_content = (any(p["text"].strip() for s in doc["body_text"] for p in s["paragraphs"])
                   or doc["figures"] or any(t["rows"] and t["cols"] for t in doc["tables"])
                   or any(u["kind"] == "heading" and u["label"].strip() for u in doc["source_map"]["units"]))
    if not has_content:
        raise DocumentExtractionError("추출할 본문·표·그림이 없는 문서입니다.", "empty")
    return doc


def extract_document(path: str | Path, *, asset_dir: str | Path | None = None,
                     timeout: float | None = None, cancel=None, progress=None) -> dict:
    """형식에 맞는 native parser로 추출한다. 의존성·암호·제한 오류는 명시한다."""
    if timeout is not None and timeout <= 0:
        raise ValueError("timeout은 양수여야 합니다.")
    try:
        return _extract_document(path, asset_dir=asset_dir, timeout=120 if timeout is None else timeout,
                                 cancel=cancel, progress=progress, default_timeout=timeout is None)
    except DocumentExtractionError:
        raise
    except RecursionError as exc:
        raise DocumentExtractionError("문서/MIME 중첩 깊이 제한을 초과했습니다.", "limit") from exc
    except OSError as exc:
        raise DocumentExtractionError(f"문서 또는 그림 읽기/쓰기 실패: {exc}", "io") from exc
    except (zipfile.BadZipFile, etree.LxmlError, ValueError, KeyError, TypeError) as exc:
        raise DocumentExtractionError(f"문서 구조를 읽을 수 없습니다: {exc}", "malformed") from exc


def _extract_document(path: str | Path, *, asset_dir: str | Path | None,
                      timeout: float, cancel, progress=None, default_timeout=False) -> dict:
    source = Path(path).resolve(strict=True)
    if not source.is_file() or source.stat().st_size > MAX_INPUT_BYTES:
        raise DocumentExtractionError("파일이 아니거나 문서 크기 제한을 초과했습니다.", "limit")
    builder = _Builder(source, asset_dir, timeout, cancel, progress, default_timeout)
    adapters = {".docx": _docx, ".pptx": _pptx, ".xlsx": _xlsx, ".txt": _text, ".md": _markdown,
                ".html": _html, ".htm": _html, ".epub": _epub, ".odt": _odf, ".ods": _odf, ".odp": _odf,
                ".eml": _eml, ".mbox": _mbox}
    adapters.update({".hwp": _rhwp, ".hwpx": _rhwp})
    if source.suffix.lower() in adapters:
        adapters[source.suffix.lower()](builder, source)
        return _finish(builder)
    if source.suffix.lower() not in (".csv", ".tsv"):
        raise DocumentExtractionError(f"아직 지원하지 않는 문서 형식입니다: {source.suffix}", "unsupported")
    if source.stat().st_size > MAX_TEXT_BYTES:
        raise DocumentExtractionError("텍스트 크기 제한을 초과했습니다.", "limit")
    with open(utils.long_path(source), "rb") as stream:
        text = _decode(stream.read())
    matrix, columns = [], 0
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter="\t" if builder.kind == "tsv" else ","):
            builder.check()
            columns = max(columns, len(row))
            if (len(matrix) + 1) * max(1, columns) > MAX_CELLS:
                raise DocumentExtractionError("표가 허용 셀 수를 초과했습니다.", "limit")
            matrix.append(row)
    except csv.Error as exc:
        raise DocumentExtractionError(f"CSV 구문 또는 필드 크기 제한 오류: {exc}", "limit") from exc
    builder.matrix(builder.unit(source.stem, "table"), matrix)
    return _finish(builder)
