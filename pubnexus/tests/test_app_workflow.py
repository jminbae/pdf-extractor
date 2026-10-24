import hashlib
import os
import threading
import time
from pathlib import Path
import unittest
from unittest.mock import Mock, patch

import test_app_export as support
from pubnexus import store


class WorkflowTests(unittest.TestCase):
    setUp = support.AppExportTests.setUp
    source = support.AppExportTests.source
    extracted = support.AppExportTests.extracted

    def test_start_extracts_in_a_child_without_running_gui_process_pipeline(self):
        source = self.root / 'background.csv'
        source.write_text('Name,Value\nBackground,713\n', encoding='utf-8')
        self.app.pdfs = [source]
        with patch.object(self.app, '_work', side_effect=AssertionError('GUI must not execute the extraction pipeline')):
            self.app.start([source])
            self.app.worker.join(timeout=20)
        self.assertFalse(self.app.worker.is_alive())
        self.assertTrue((self.root / 'Markdown/background.md').is_file())
        started = [payload for kind, payload in self.events if kind == 'worker_started']
        self.assertTrue(started)
        self.assertNotEqual(started[0]['pid'], os.getpid())

    def test_each_document_uses_a_fresh_child_process(self):
        sources = [self.root / name for name in ('one.csv', 'two.csv')]
        for source in sources:
            source.write_text('Name,Value\n' + source.stem + ',713\n', encoding='utf-8')
        self.app.pdfs = sources
        self.app.start(sources)
        self.app.worker.join(timeout=20)
        self.assertFalse(self.app.worker.is_alive())
        started = [payload for kind, payload in self.events if kind == 'worker_started']
        self.assertEqual(len(started), 2, 'Each document needs its own process boundary')
        self.assertEqual(len({item['pid'] for item in started}), 2)
        self.assertTrue(all(row['done'] for row in self.app.file_rows()))

    def test_enqueue_during_run_deduplicates_and_a_crashed_document_cannot_stop_next(self):
        from pubnexus import extraction_process
        first = self.root / 'crashing.csv'
        second = self.root / 'next.csv'
        for source in (first, second):
            source.write_text('Name,Value\n' + source.stem + ',713\n', encoding='utf-8')
        self.app.pdfs = [first, second]
        entered, release = threading.Event(), threading.Event()
        self.addCleanup(release.set)
        real_run = extraction_process.run
        calls = []

        def run(paths, output_dir, **kwargs):
            calls.append(list(paths))
            if paths == [first]:
                entered.set()
                release.wait(5)
                raise RuntimeError('Synthetic worker crash')
            return real_run(paths, output_dir, **kwargs)

        with patch.object(extraction_process, 'run', side_effect=run):
            self.app.start([first])
            self.assertTrue(entered.wait(3))
            api = support.app_module.Api(self.app)
            added = api.extract_selected([1], True)
            self.assertTrue(added['ok'], added)
            self.assertTrue(added['queued'])
            self.assertEqual(added['n'], 1)
            duplicate = api.extract_selected([0, 1], True)
            self.assertTrue(duplicate['ok'], duplicate)
            self.assertEqual(duplicate['n'], 0)
            release.set()
            self.app.worker.join(timeout=20)
        self.assertFalse(self.app.worker.is_alive())
        self.assertEqual(calls, [[first], [second]])
        rows = self.app.file_rows()
        self.assertEqual(rows[0]['status'], 'failed')
        self.assertTrue(rows[1]['done'])
        final = [payload for kind, payload in self.events if kind == 'run'][-1]
        self.assertEqual((final['done'], final['total'], final['nfail']), (1, 2, 1))

    def test_cancel_stops_queued_documents_before_any_new_process_starts(self):
        from pubnexus import extraction_process
        from pubnexus.general_pdf import ExtractionCancelled
        sources = [self.root / name for name in ('working.csv', 'queued.csv')]
        for source in sources:
            source.write_text('Name,Value\n' + source.stem + ',713\n', encoding='utf-8')
        self.app.pdfs = sources
        entered = threading.Event()
        calls = []

        def run(paths, output_dir, **kwargs):
            calls.append(list(paths))
            entered.set()
            deadline = time.monotonic() + 5
            while not kwargs['cancel']() and time.monotonic() < deadline:
                time.sleep(.01)
            raise ExtractionCancelled('Stopped')

        with patch.object(extraction_process, 'run', side_effect=run):
            self.app.start(sources)
            self.assertTrue(entered.wait(3))
            support.app_module.Api(self.app).cancel()
            self.app.worker.join(timeout=3)
        self.assertFalse(self.app.worker.is_alive())
        self.assertEqual(calls, [[sources[0]]])
        self.assertTrue(all(row['status'] == 'paused' for row in self.app.file_rows()))

    def test_committed_document_is_reserved_until_child_exit_and_survives_stop(self):
        from pubnexus import extraction_process
        source = self.root / 'saved.csv'
        source.write_text('Name,Value\nSaved,713\n', encoding='utf-8')
        self.app.pdfs = [source]
        self.app._work([source])
        saved = dict(self.app.exported[support.app_module._norm(source)])
        accepted_after_commit = []

        def run(paths, output_dir, **kwargs):
            kwargs['on_event']('mark', dict(saved, path=str(source), name=source.name, md_ready=True))
            accepted_after_commit.extend(self.app.start([source], True))
            self.app.cancel.set()  # Stop requested while this committed child is exiting.
            kwargs['on_event']('run', {'on': False, 'stopped': True, 'nfail': 0})

        with patch.object(extraction_process, 'run', side_effect=run):
            self.app.start([source], True)
            self.app.worker.join(timeout=3)
        self.assertFalse(self.app.worker.is_alive())
        self.assertEqual(accepted_after_commit, [], 'Do not accept a duplicate until the old child exits')
        self.assertEqual(self.app.file_rows()[0]['status'], 'complete')
        self.assertTrue(self.app.file_rows()[0]['done'])

    def test_saved_markdown_remains_readable_while_same_file_is_reextracted(self):
        source = self.source()
        doc = self.extracted(source)
        store.save(doc['sha1'], doc, source)
        self.app.pdfs = [source]
        self.app._work([source])
        self.app._active_path = source
        self.app._pending_paths = {support.app_module._norm(source)}
        self.app._file_progress.clear()
        self.app.prog(1, 1, source.name, '문서 구조 추출', 0, path=source, file_pct=25)
        row = self.app.file_rows()[0]
        self.assertEqual(row['status'], 'running')
        self.assertTrue(row['md_ready'], 'The already saved Markdown must stay readable')
        self.assertFalse(row['done'], 'An active extraction is not finished')
        self.assertEqual(row['progress_pct'], 25)
        self.assertTrue(self.app.doc_view(source)['extracted'])

    def test_file_addition_is_available_during_extraction_without_changing_targets(self):
        first = self.source()
        second = self.root / 'notes.md'
        second.write_text('A separate readable document.', encoding='utf-8')
        self.app.pdfs = [first]
        release = threading.Event()
        self.addCleanup(release.set)
        self.app.worker = threading.Thread(target=release.wait, daemon=True)
        self.app.worker.start()
        self.app._pending_paths = {support.app_module._norm(first)}
        self.app.add_files([second])
        self.assertEqual(self.app.pdfs, [first, second])
        self.assertEqual(self.app._pending_paths, {support.app_module._norm(first)})
        self.assertIn('key', self.app.source_view(0))
        self.addCleanup(self.app.renderer.close_all)

    def test_open_original_before_extraction_never_reads_or_queues_markdown(self):
        source = self.source()
        self.app.pdfs = [source]
        self.addCleanup(self.app.renderer.close_all)
        with patch.object(self.app, 'doc_view', side_effect=AssertionError('No Markdown on original open')), \
             patch.object(self.app, 'sha_for', side_effect=AssertionError('No extraction cache on original open')), \
             patch.object(self.app, 'bump', side_effect=AssertionError('Reading never changes queue')):
            result = self.app.source_view(0)
        self.assertEqual(len(result['pages']), 1)
        self.assertIn('key', result)
        self.assertNotIn('doc', result)
        self.assertFalse(self.app.done)
        self.assertFalse(self.app.exported)
        self.assertFalse(self.app.file_rows()[0]['md_ready'])

    def test_result_only_http_open_does_not_wait_for_office_preview_or_change_queue(self):
        from pptx import Presentation
        source = self.root / 'slides.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = 'Saved slide content 8426'
        deck.save(source)
        self.app.pdfs = [source]
        self.app._work([source])
        handler = object.__new__(support.app_module.Handler)
        handler.app = self.app
        handler.path = '/open?p=0&mode=result'
        responses = []
        handler._json = responses.append
        with patch('pubnexus.preview.preview_path', side_effect=AssertionError('Result reading must not render Office')) as preview, \
             patch.object(self.app, 'bump', side_effect=AssertionError('Reading must not change extraction queue')):
            handler.do_GET()
        preview.assert_not_called()
        self.assertEqual(len(responses), 1)
        result = responses[0]
        self.assertTrue(result['doc']['extracted'])
        self.assertIn('Saved slide content 8426', result['doc']['md_text'])
        self.assertFalse(result['source_ready'])
        self.assertEqual(result['pages'], [])
        self.assertEqual(result['revision'], result['doc']['revision'])

    def test_source_navigation_labels_use_document_units(self):
        from pubnexus.preview import Preview
        display = self.source()
        self.addCleanup(self.app.renderer.close_all)
        for suffix, label, unit in [('.pptx', '원본 슬라이드', '슬라이드'),
                                    ('.odp', '원본 슬라이드', '슬라이드'),
                                    ('.xlsx', '원본 시트', '쪽'),
                                    ('.docx', '원본 문서', '쪽'),
                                    ('.hwp', '원본 문서', '쪽'),
                                    ('.pdf', '원본 PDF', '쪽')]:
            with self.subTest(suffix=suffix):
                source = self.root / ('original' + suffix)
                source.write_bytes(b'Synthetic source for label only')
                self.app.pdfs = [source]
                with patch('pubnexus.preview.preview_path', return_value=Preview(display, suffix != '.pdf', 'synthetic', 'sha')), \
                     patch.object(self.app.renderer, 'open', return_value={'key': 1, 'pages': [[612, 792]]}):
                    result = self.app.source_view(0, renderer='pdf')
                self.assertEqual(result.get('source_label'), label)
                self.assertEqual(result.get('source_unit_label'), unit)
                self.assertTrue(result.get('source_ready'))

    def test_fast_pptx_source_reads_header_without_running_full_pdf_renderer(self):
        import sys
        from types import SimpleNamespace
        source = self.root / 'fast.pptx'
        source.write_bytes(b'Synthetic opaque package for dispatch only')
        self.app.pdfs = [source]
        header = {'source_kind': 'presentation', 'source_ready': True,
                  'source_revision': 'current-source', 'pages': [[960, 540]] * 71}
        engine = SimpleNamespace(available=lambda: True, describe=Mock(return_value=header))
        import pubnexus
        with patch.dict(sys.modules, {'pubnexus.presentation_preview': engine}), \
             patch.object(pubnexus, 'presentation_preview', engine, create=True), \
             patch('pubnexus.preview.preview_path', side_effect=AssertionError('No full Office conversion for fast preview')) as slow:
            result = self.app.source_view(0)
        slow.assert_not_called()
        self.assertEqual(result['source_kind'], 'presentation')
        self.assertEqual(result['pages'], [[960, 540]] * 71)
        engine.describe.assert_called_once_with(source, 0)

    def test_fast_pptx_map_and_result_search_do_not_prepare_full_pdf(self):
        from pptx import Presentation
        from pubnexus import presentation_preview
        source = self.root / 'searchable.pptx'
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = 'Searchable presentation 8426'
        deck.save(source)
        self.app.pdfs = [source]
        self.app._work([source])
        view = self.app.result_view(0)
        with patch.object(presentation_preview, 'available', return_value=True), \
             patch('pubnexus.preview.preview_path', side_effect=AssertionError('No full conversion for fast map/search')):
            mapping = self.app.document_map(0, renderer='presentation')
            result = support.app_module.Api(self.app).search(0, '8426', view['revision'], 'presentation')
        self.assertEqual(mapping['source_kind'], 'presentation')
        self.assertTrue(mapping['source_revision'])
        self.assertTrue(any(block.get('slide') == 0 for block in mapping['blocks']))
        self.assertTrue(all(not block['locations'] for block in mapping['blocks']))
        self.assertTrue(result['ok'], result)
        self.assertTrue(result['results'])

    def test_presentation_raw_data_has_no_cors_and_rejects_stale_or_invalid_selection(self):
        from pptx import Presentation
        from pubnexus import presentation_preview
        source = self.root / 'raw.pptx'
        deck = Presentation(); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(source)
        self.app.pdfs = [source]
        revision = presentation_preview.describe(source, 0)['source_revision']
        handler = object.__new__(support.app_module.Handler)
        handler.app = self.app
        handler._send = Mock()
        handler.send_error = Mock()
        handler.path = '/presentation-data?p=0&r=' + revision
        handler.do_GET()
        self.assertEqual(handler._send.call_args.args[0], source.read_bytes())
        self.assertNotIn('headers', handler._send.call_args.kwargs)
        handler._send.reset_mock()
        handler.path = '/presentation-data?p=0&r=' + '0' * 64
        handler.do_GET()
        handler.send_error.assert_called_with(409)
        handler._send.assert_not_called()
        handler.path = '/presentation-data?p=-1&r=' + revision
        handler.do_GET()
        handler.send_error.assert_called_with(404)

    def test_presentation_static_modules_allow_opaque_frame_but_data_does_not(self):
        from pubnexus import presentation_preview
        handler = object.__new__(support.app_module.Handler)
        handler.app = self.app
        handler._send = Mock()
        with patch.object(presentation_preview, 'asset', return_value=b'const fixture = true;'):
            for route in ('/presentation-frame.mjs', '/presentation-vendor.js'):
                handler.path = route
                handler.do_GET()
                self.assertEqual(handler._send.call_args.kwargs['headers'], {'Access-Control-Allow-Origin': '*'})
            handler.path = '/presentation-frame'
            handler.do_GET()
        headers = handler._send.call_args.kwargs['headers']
        self.assertEqual(headers, {'Content-Security-Policy': presentation_preview.CSP})
        self.assertIn("connect-src 'none'", headers['Content-Security-Policy'])

    def test_original_search_works_before_markdown_exists(self):
        source = self.source()
        self.app.pdfs = [source]
        with patch.object(self.app, 'comparison_doc', side_effect=AssertionError('No extracted result required')):
            result = support.app_module.Api(self.app).search_source(0, 'offline')
        self.assertTrue(result['ok'], result)
        self.assertTrue(result['results'])
        self.assertTrue(all(hit['side'] == 'source' for hit in result['results']))
        self.assertEqual(result['results'][0]['page'], 0)
        self.assertFalse(self.app.exported)

    def test_reopened_original_uses_current_bytes_and_retires_previous_images(self):
        import fitz
        source = self.source()
        self.app.pdfs = [source]
        self.addCleanup(self.app.renderer.close_all)
        first = self.app.source_view(0)
        self.assertTrue(self.app.renderer.png(first['key'], 0, 300))
        with fitz.open() as replacement:
            replacement.new_page(width=500, height=600)
            replacement.new_page(width=500, height=600)
            source.write_bytes(replacement.tobytes())
        second = self.app.source_view(0)
        self.assertEqual(second['pages'], [[500.0, 600.0], [500.0, 600.0]])
        self.assertNotEqual(first['key'], second['key'])
        self.assertIsNone(self.app.renderer.png(first['key'], 0, 300))

    def test_row_progress_identifies_path_and_md_is_ready_only_after_save(self):
        source = self.source()
        other = self.root / 'other' / source.name
        other.parent.mkdir()
        other.write_bytes(source.read_bytes())
        self.app.pdfs = [source, other]
        self.app._active_path = source
        self.app.prog(1, 2, source.name, '원문 확인', 0, path=source,
                      file_pct=37.5, page=3, pages=8)
        rows = self.app.file_rows()
        self.assertEqual(rows[0]['status'], 'running')
        self.assertEqual(rows[0]['progress_pct'], 37.5)
        self.assertEqual(rows[0]['progress_page'], 3)
        self.assertEqual(rows[1]['status'], 'ready')
        self.assertFalse(rows[0]['md_ready'])
        event = [p for kind, p in self.events if kind == 'prog'][-1]
        self.assertEqual(event['path'], str(source))
        self.assertEqual(event['file_pct'], 37.5)
        self.app._active_path = None
        doc = self.extracted(source)
        store.save(doc['sha1'], doc, source)
        self.app._work([source])
        self.assertTrue(self.app.file_rows()[0]['md_ready'])
        self.assertEqual(self.app.file_rows()[0]['progress_pct'], 100)

    def test_previous_source_markdown_is_not_mixed_with_current_cached_html(self):
        source = self.source()
        old_doc = self.extracted(source)
        store.save(old_doc['sha1'], old_doc, source)
        self.app.pdfs = [source]
        self.app._work([source])
        markdown = Path(self.app.file_rows()[0]['markdown'])
        old_bytes = markdown.read_bytes()
        source.write_bytes(source.read_bytes() + b'\nChanged revision\n')
        new_doc = self.extracted(source)
        new_doc['body_text'][0]['paragraphs'][0]['text'] = 'Current canonical revision content.'
        store.save(new_doc['sha1'], new_doc, source)
        view = self.app.doc_view(source)
        self.assertIn('Current canonical revision content.', view['html'])
        self.assertIn('Current canonical revision content.', view['md_text'])
        self.assertEqual(markdown.read_bytes(), old_bytes)
        self.assertFalse(self.app.file_rows()[0]['md_ready'])

    def test_legacy_general_pdf_needs_refresh_without_deleting_previous_output(self):
        source = self.source()
        doc = self.extracted(source)
        store.save(doc['sha1'], doc, source)
        self.app.pdfs = [source]
        self.app._work([source])
        record = self.app.exported[support.app_module._norm(source)]
        markdown = Path(record['markdown'])
        before = markdown.read_bytes()
        from pubnexus import pdf_markdown
        record.pop('pdf_conversion_engine', None)
        record.pop('pdf_conversion_attempt', None)
        with patch.object(pdf_markdown, 'available', return_value=True):
            row = self.app.file_rows()[0]
        self.assertTrue(row['needs_refresh'])
        self.assertFalse(row['md_ready'])
        self.assertEqual(row['status'], 'ready')
        self.assertEqual(markdown.read_bytes(), before)
        self.app._work([source])
        self.assertTrue(self.app.file_rows()[0]['md_ready'])


if __name__ == '__main__':
    unittest.main()
