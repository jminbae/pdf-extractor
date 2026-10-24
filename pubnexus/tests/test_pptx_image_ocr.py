"""Embedded-picture OCR keeps original PPTX assets and logical provenance."""
from __future__ import annotations

import hashlib
import io
import os
from pathlib import Path
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))

from pubnexus import documents, local_ocr


class PPTXImageOCRTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix='pubnexus-pptx-ocr-test-')
        self.addCleanup(temporary.cleanup)
        self.base = Path(temporary.name)
        environment = patch.dict(os.environ, {'PUBNEXUS_STORE': str(self.base / 'store')})
        environment.start()
        self.addCleanup(environment.stop)

    def image(self, text='METHODS image body marker 7391'):
        image = Image.new('RGB', (1500, 300), 'white')
        draw = ImageDraw.Draw(image)
        draw.text((35, 100), text, font=ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 52), fill='black')
        output = io.BytesIO()
        image.save(output, format='PNG')
        return output.getvalue()

    def presentation(self, image, name='image-body.pptx'):
        from pptx import Presentation
        from pptx.util import Inches
        path = self.base / name
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(.3), Inches(6), Inches(.5)).text = 'NATIVE HEAD'
        slide.shapes.add_picture(io.BytesIO(image), Inches(1), Inches(1), width=Inches(8))
        slide.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(.5)).text = 'NATIVE END'
        presentation.save(path)
        return path

    def test_real_picture_body_is_markdown_text_and_keeps_image_origin_and_original_bytes(self):
        if local_ocr.discover() is None:
            self.skipTest('Pinned local OCR models are unavailable')
        image = self.image()
        path = self.presentation(image)
        before = hashlib.sha256(path.read_bytes()).hexdigest()

        doc = documents.extract_document(path)

        paragraphs = [p for section in doc['body_text'] for p in section['paragraphs']]
        recovered = [p for p in paragraphs if '7391' in p['text']]
        self.assertTrue(recovered, [p['text'] for p in paragraphs])
        figure = doc['figures'][0]
        for paragraph in recovered:
            origin = paragraph['source']
            self.assertEqual(origin['image_id'], figure['id'])
            self.assertEqual(origin['slide'], 1)
            self.assertEqual(origin['shape'], '2')
            self.assertEqual(origin['method'], 'tesseract')
            self.assertEqual(origin['coordinate_space'], 'ocr_image_pixels')
            self.assertNotIn('page', origin)
            self.assertNotIn('bbox', origin)
            self.assertTrue(origin['image_bbox_px'])
        self.assertLess(paragraphs.index(next(p for p in paragraphs if p['text'] == 'NATIVE HEAD')),
                        paragraphs.index(recovered[0]))
        self.assertLess(paragraphs.index(recovered[0]),
                        paragraphs.index(next(p for p in paragraphs if p['text'] == 'NATIVE END')))
        self.assertEqual((Path(doc['extraction_report']['asset_base']) / figure['image']).read_bytes(), image)
        from pubnexus.render import to_markdown
        markdown = to_markdown(doc)
        self.assertIn('7391', markdown)
        self.assertIn('![' , markdown)
        self.assertTrue(documents.current(doc))
        self.assertEqual(doc['extraction_report']['image_ocr_version'], documents.PPTX_IMAGE_OCR_VERSION)
        self.assertEqual(doc['extraction_report']['status'], 'partial')
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_repeated_images_share_one_worker_and_model_copy_but_keep_each_shape_and_skip_native_lines(self):
        if local_ocr.discover() is None:
            self.skipTest('Pinned local OCR models are unavailable')
        from pptx import Presentation
        from pptx.util import Inches
        image = Image.new('RGB', (1500, 300), 'white')
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 52)
        draw.text((35, 20), 'NATIVE HEAD', font=font, fill='black')
        draw.text((35, 135), 'METHODS image body marker 7391', font=font, fill='black')
        output = io.BytesIO()
        image.save(output, format='PNG')
        blob = output.getvalue()
        path = self.presentation(blob)
        presentation = Presentation(path)
        presentation.slides[0].shapes.add_picture(io.BytesIO(blob), Inches(1), Inches(5), width=Inches(8))
        presentation.save(path)
        progress = []
        with patch.object(local_ocr, '_run', wraps=local_ocr._run) as run, \
                patch.object(local_ocr.shutil, 'copyfile', wraps=local_ocr.shutil.copyfile) as copied:
            doc = documents.extract_document(path, progress=lambda *event: progress.append(event))
        self.assertEqual(run.call_count, 1)
        self.assertEqual(sum(str(call.args[0]).endswith('.traineddata') for call in copied.call_args_list), 2)
        report = doc['extraction_report']['image_ocr']
        self.assertEqual((report['images'], report['unique_images'], report['cache_reused']), (2, 1, 1))
        paragraphs = [p for section in doc['body_text'] for p in section['paragraphs']]
        self.assertEqual(sum(p['text'] == 'NATIVE HEAD' for p in paragraphs), 1)
        recovered = [p for p in paragraphs if '7391' in p['text']]
        self.assertEqual(len(recovered), 2)
        self.assertEqual(len({p['source']['image_id'] for p in recovered}), 2)
        self.assertEqual(report['native_duplicate_lines'], 2)
        self.assertEqual(progress[-1], (1, 1, 'document_ocr'))
        self.assertIn((1, 1, 'document_native'), progress)

    def test_visible_crop_does_not_reveal_hidden_text_and_keeps_group_rotation_as_logical_provenance(self):
        if local_ocr.discover() is None:
            self.skipTest('Pinned local OCR models are unavailable')
        from pptx import Presentation
        from pptx.util import Inches
        image = Image.new('RGB', (2000, 350), 'white')
        draw = ImageDraw.Draw(image)
        font = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 60)
        draw.text((35, 120), 'HIDDEN 8247', font=font, fill='black')
        draw.text((1050, 120), 'VISIBLE 7391', font=font, fill='black')
        output = io.BytesIO()
        image.save(output, format='PNG')
        blob = output.getvalue()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        picture = group.shapes.add_picture(io.BytesIO(blob), Inches(1), Inches(1), width=Inches(8))
        picture.crop_left = .5
        picture.rotation = 90
        group.rotation = 15
        path = self.base / 'cropped-group.pptx'
        presentation.save(path)

        doc = documents.extract_document(path)

        paragraphs = [p for section in doc['body_text'] for p in section['paragraphs'] if p['text']]
        self.assertIn('7391', ' '.join(p['text'] for p in paragraphs))
        self.assertNotIn('8247', ' '.join(p['text'] for p in paragraphs))
        origin = next(p['source'] for p in paragraphs if '7391' in p['text'])
        self.assertEqual(origin['shape'], '1/1')
        self.assertEqual(origin['shape_transform']['rotation_degrees'], 90)
        self.assertEqual(origin['group_transforms'][0]['rotation_degrees'], 15)
        self.assertEqual(origin['image_source_crop_px'][0], 1000)
        self.assertEqual(origin['ocr_orientation'], 'source_image_before_shape_transform')
        self.assertNotIn('bbox', origin)
        self.assertEqual((Path(doc['extraction_report']['asset_base']) / doc['figures'][0]['image']).read_bytes(), blob)

    def test_unavailable_or_failed_ocr_keeps_native_body_original_picture_and_attempt_version(self):
        blob = self.image()
        path = self.presentation(blob)
        for mode in ('unavailable', 'failed'):
            with self.subTest(mode=mode):
                def failed(images, provider, **kwargs):
                    return {key: {'ok': False, 'code': 'ocr', 'error': 'synthetic OCR failure'} for key in images}
                provider = {'engine': 'pymupdf_builtin_tesseract', 'language': 'eng', 'identity': 'test'}
                with patch.object(local_ocr, 'discover', return_value=None if mode == 'unavailable' else provider), \
                        patch.object(local_ocr, 'extract_images', side_effect=failed):
                    doc = documents.extract_document(path)
                self.assertTrue(documents.current(doc))
                self.assertEqual(doc['extraction_report']['status'], 'partial')
                paragraphs = [p['text'] for section in doc['body_text'] for p in section['paragraphs']]
                self.assertEqual(paragraphs, ['NATIVE HEAD', '', 'NATIVE END'])
                self.assertEqual(len(doc['figures']), 1)
                self.assertEqual((Path(doc['extraction_report']['asset_base']) / doc['figures'][0]['image']).read_bytes(), blob)
                report = doc['extraction_report']['image_ocr']
                self.assertEqual(report['status'], 'unavailable' if mode == 'unavailable' else 'partial')

    def test_default_ocr_budget_scales_with_unique_images_and_explicit_timeout_is_respected(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for number in range(35):
            output = io.BytesIO()
            Image.new('RGB', (4, 3), (number, 50, 100)).save(output, format='PNG')
            slide.shapes.add_picture(io.BytesIO(output.getvalue()), Inches(1), Inches(1))
        path = self.base / 'image-batch.pptx'
        presentation.save(path)
        observed = []
        def recognize(images, provider, **kwargs):
            observed.append(kwargs['timeout'])
            return {key: {'ok': True, 'words': []} for key in images}
        provider = {'engine': 'pymupdf_builtin_tesseract', 'language': 'eng', 'identity': 'test'}
        with patch.object(local_ocr, 'discover', return_value=provider), \
                patch.object(local_ocr, 'extract_images', side_effect=recognize):
            documents.extract_document(path)
            documents.extract_document(path, timeout=2)
        self.assertGreater(observed[0], 300)
        self.assertTrue(0 < observed[1] <= 2)

    def test_cancel_after_completed_image_reaps_batch_and_removes_private_models(self):
        provider = local_ocr.discover()
        if provider is None:
            self.skipTest('Pinned local OCR models are unavailable')
        from pubnexus.general_pdf import ExtractionCancelled
        from pubnexus.image_ocr import prepare_picture
        images = {item['id']: item['png'] for item in
                  (prepare_picture(self.image('FIRST OCR 7391')), prepare_picture(self.image('SECOND OCR 8247')))}
        paths, processes, stopped = [], [], []
        temporary_factory = tempfile.TemporaryDirectory
        popen = local_ocr.subprocess.Popen
        def temporary(*args, **kwargs):
            value = temporary_factory(*args, dir=self.base, **kwargs)
            paths.append(Path(value.name))
            return value
        def spawn(*args, **kwargs):
            process = popen(*args, **kwargs)
            processes.append(process)
            return process
        started = time.monotonic()
        with patch.object(local_ocr.tempfile, 'TemporaryDirectory', temporary), \
                patch.object(local_ocr.subprocess, 'Popen', spawn):
            with self.assertRaises(ExtractionCancelled):
                local_ocr.extract_images(images, provider, cancel=lambda: bool(stopped),
                                         progress=lambda key: stopped.append(key))
        self.assertLess(time.monotonic() - started, 10)
        self.assertTrue(processes)
        self.assertTrue(all(process.poll() is not None for process in processes))
        self.assertTrue(all(not path.exists() for path in paths))

    def test_subpixel_crop_is_a_per_image_failure_and_preserves_native_content(self):
        from pptx import Presentation
        from pubnexus.image_ocr import prepare_picture
        output = io.BytesIO()
        Image.new('RGB', (10, 10), 'white').save(output, format='PNG')
        blob = output.getvalue()
        try:
            prepare_picture(blob, (.49, 0, .49, 0))
        except Exception as error:
            self.assertIsInstance(error, local_ocr.LocalOCRError)
        else:
            self.fail('A rounded zero-width crop must not be sent to OCR')
        path = self.presentation(blob)
        presentation = Presentation(path)
        picture = presentation.slides[0].shapes[1]
        picture.crop_left = picture.crop_right = .49
        presentation.save(path)
        provider = {'engine': 'pymupdf_builtin_tesseract', 'language': 'eng', 'identity': 'test'}
        with patch.object(local_ocr, 'discover', return_value=provider):
            doc = documents.extract_document(path)
        self.assertEqual(doc['extraction_report']['status'], 'partial')
        self.assertEqual(len(doc['extraction_report']['image_ocr']['failed_images']), 1)
        self.assertIn('NATIVE HEAD', [p['text'] for s in doc['body_text'] for p in s['paragraphs']])
        self.assertEqual((Path(doc['extraction_report']['asset_base']) / doc['figures'][0]['image']).read_bytes(), blob)

    def test_prepared_image_cache_obeys_batch_count_and_bytes_before_retaining_more_pngs(self):
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        for number in range(3):
            output = io.BytesIO()
            Image.new('RGB', (4, 3), (number, 50, 100)).save(output, format='PNG')
            slide.shapes.add_picture(io.BytesIO(output.getvalue()), Inches(1), Inches(1))
        path = self.base / 'bounded-image-cache.pptx'
        presentation.save(path)
        provider = {'engine': 'pymupdf_builtin_tesseract', 'language': 'eng', 'identity': 'test'}
        for boundary, value in (('MAX_BATCH_IMAGES', 1), ('MAX_BATCH_BYTES', 100)):
            with self.subTest(boundary=boundary):
                received = []
                def recognize(images, provider, **kwargs):
                    received.append(images)
                    return {key: {'ok': True, 'words': []} for key in images}
                with patch.object(local_ocr, 'discover', return_value=provider), \
                        patch.object(local_ocr, boundary, value), \
                        patch.object(local_ocr, 'extract_images', side_effect=recognize):
                    doc = documents.extract_document(path)
                self.assertEqual(len(received[0]), 1)
                self.assertEqual(len(doc['extraction_report']['image_ocr']['failed_images']), 2)
                self.assertEqual(len(doc['figures']), 3)

    def test_invalid_coordinates_of_one_image_do_not_discard_other_completed_images(self):
        provider = local_ocr.discover()
        if provider is None:
            self.skipTest('Pinned local OCR models are unavailable')
        from pubnexus.image_ocr import prepare_picture
        prepared = [prepare_picture(self.image(text)) for text in ('GOOD 7391', 'BAD 8247')]
        images = {item['id']: item['png'] for item in prepared}
        records = {'ok': True, 'images': [
            {'id': prepared[0]['id'], 'result': {'ok': True, 'engine_version': 'test',
                                               'words': [[1, 1, 10, 10, 'GOOD 7391', 0, 0, 0]]}},
            {'id': prepared[1]['id'], 'result': {'ok': True, 'engine_version': 'test',
                                               'words': [[float('nan'), 1, 10, 10, 'BAD 8247', 0, 0, 0]]}}]}
        with patch.object(local_ocr, '_run', return_value=records):
            try:
                result = local_ocr.extract_images(images, provider)
            except local_ocr.LocalOCRError as error:
                self.fail(f'One image invalidated the entire completed batch: {error}')
        self.assertTrue(result[prepared[0]['id']]['ok'])
        self.assertFalse(result[prepared[1]['id']]['ok'])

    def test_ocr_font_metrics_outside_image_keep_text_with_only_logical_image_location(self):
        from pubnexus.image_ocr import lines
        record = {'width': 1600, 'height': 900, 'language': 'eng', 'engine': 'pymupdf_builtin_tesseract',
                  'engine_identity': 'test', 'engine_version': 'test',
                  'words': [[189., 872., 254.97496032714844, 901.0289916992188, 'phototherapy,', 19, 0, 0]]}
        text, origin = next(lines(record))
        self.assertEqual(text, 'phototherapy,')
        self.assertIsNone(origin['image_bbox_px'])
        self.assertIsNone(origin['words'][0]['bbox_px'])
        self.assertEqual(origin['raw_ocr_bbox_px'], record['words'][0][:4])
        self.assertEqual(origin['geometry_status'], 'unlocated_ocr')

    def test_different_encoded_images_with_identical_visible_pixels_share_one_bounded_png(self):
        from pptx import Presentation
        from pptx.util import Inches
        from PIL.PngImagePlugin import PngInfo
        from pubnexus import image_ocr
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        blobs = []
        for number in range(4):
            info = PngInfo()
            info.add_text('source', str(number))
            output = io.BytesIO()
            Image.new('RGB', (30, 20), 'white').save(output, format='PNG', pnginfo=info)
            blobs.append(output.getvalue())
            slide.shapes.add_picture(io.BytesIO(blobs[-1]), Inches(1), Inches(1))
        path = self.base / 'same-visible-pixels.pptx'
        presentation.save(path)
        prepared = []
        prepare = image_ocr.prepare_picture
        def capture(*args, **kwargs):
            value = prepare(*args, **kwargs)
            prepared.append(value)
            return value
        def recognize(images, provider, **kwargs):
            return {key: {'ok': True, 'words': []} for key in images}
        provider = {'engine': 'pymupdf_builtin_tesseract', 'language': 'eng', 'identity': 'test'}
        with patch.object(image_ocr, 'prepare_picture', side_effect=capture), \
                patch.object(local_ocr, 'MAX_BATCH_BYTES', 200), \
                patch.object(local_ocr, 'discover', return_value=provider), \
                patch.object(local_ocr, 'extract_images', side_effect=recognize):
            doc = documents.extract_document(path)
        self.assertEqual(len(prepared), 4)
        self.assertEqual(len({id(item['png']) for item in prepared}), 1,
                         'Cached equal PNGs must share storage to obey the decoded-byte budget')
        self.assertEqual({item['image_sha256'] for item in prepared},
                         {hashlib.sha256(blob).hexdigest() for blob in blobs})
        self.assertEqual(doc['extraction_report']['image_ocr']['unique_images'], 1)
        self.assertEqual(doc['extraction_report']['image_ocr']['failed_images'], [])


if __name__ == '__main__':
    unittest.main()
