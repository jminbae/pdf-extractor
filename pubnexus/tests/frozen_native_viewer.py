"""Six native formats in an owned frozen WebView2, with TEMP-only inputs/profile.

Reuses frozen_smoke's CDP transport, Windows Job ownership and four-asset hashing.
This is an executable test harness, never a production application entry point.
"""
from __future__ import annotations

import argparse
import ctypes
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import shutil
import socket
import stat
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile

sys.dont_write_bytecode = True
import frozen_smoke as smoke

UPSTREAM_COMMIT = "1a76570e833917d15817415a53c09ad61ab3203f"
UPSTREAM_PATH = "samples/hwpx/exam_social.hwpx"
UPSTREAM_SHA256 = "c11b117dfc55606bb3ea703f9ea5b0a14481f7e70e918bc09ce4e202a04c9fd0"
RHWP_SHA256 = "1e1b7ae6c9ffd895e1a330a191666a2629cc7eba2fc9e37442e9f1755741bb62"
FORMATS = ("docx", "xlsx", "txt", "md", "rtf", "hwpx")


def fingerprint(path):
    path = Path(path)
    value = path.stat()
    return {"path": str(path), "sha256": smoke._digest(path),
            "bytes": value.st_size, "mtime_ns": value.st_mtime_ns}


def fixtures(root, hwpx):
    from docx import Document
    from openpyxl import Workbook
    inputs, cache = root / "inputs", root / "cache"
    inputs.mkdir()
    cache.mkdir()
    specs = []
    def spec(fmt, name, kind, marker, query):
        specs.append({"format": fmt, "name": name, "source_kind": kind,
                      "marker": marker, "query": query})
        return inputs / name
    source = spec("docx", "01-source.docx", "content", "FROZEN_NATIVE_DOCX_127", "DOCX_TARGET_127")
    document = Document()
    document.add_heading(specs[-1]["marker"], 1)
    for i in range(65):
        document.add_paragraph("Original paragraph " + str(i) + ". Read the source without conversion.")
    document.add_paragraph(specs[-1]["query"])
    document.save(source)
    source = spec("xlsx", "02-source.xlsx", "sheet", "FROZEN_NATIVE_XLSX_127", "XLSX_TARGET_127")
    book = Workbook()
    sheet = book.active
    sheet.title = "Original"
    sheet.append([specs[-1]["marker"], "=21*2"])
    for i in range(2, 48):
        sheet.append(["Original row " + str(i), i])
    sheet.append([specs[-1]["query"], 48])
    book.save(source)
    book.close()
    # A real persisted OOXML calculation cache, not a newly calculated formula.
    with zipfile.ZipFile(source) as archive:
        members = {item.filename: archive.read(item) for item in archive.infolist()}
    previous = b"<f>21*2</f><v></v>"
    assert previous in members["xl/worksheets/sheet1.xml"]
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        previous, b"<f>21*2</f><v>42</v>", 1)
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, value in members.items():
            archive.writestr(name, value)
    source = spec("txt", "03-source.txt", "text", "FROZEN_NATIVE_TXT_127", "TXT_TARGET_127")
    raw = specs[-1]["marker"] + "\r\n😀 원문 <img src=x> 그대로\r\n" + "\r\n".join(
        "Original line " + str(i) for i in range(35)) + "\r\n" + specs[-1]["query"]
    source.write_bytes(raw.encode("utf-8"))
    specs[-1]["raw_text"] = raw
    source = spec("md", "04-source.md", "text", "FROZEN_NATIVE_MD_127", "MD_TARGET_127")
    raw = "# " + specs[-1]["marker"] + "\n\n**Markdown 원본**\n\n" + "\n\n".join(
        "Original paragraph " + str(i) for i in range(30)) + "\n\n" + specs[-1]["query"]
    source.write_bytes(raw.encode("utf-8"))
    specs[-1]["raw_text"] = raw
    source = spec("rtf", "05-source.rtf", "content", "FROZEN_NATIVE_RTF_127", "RTF_TARGET_127")
    korean = "".join("\\u" + str(ord(c) - 65536 if ord(c) > 32767 else ord(c)) + "?" for c in "원문 한글" if c != " ")
    rtf = r"{\rtf1\ansi\uc1 " + specs[-1]["marker"] + r"\par " + korean + r"\par "
    rtf += r"\par ".join("Original RTF paragraph " + str(i) for i in range(40))
    source.write_text(rtf + r"\par " + specs[-1]["query"] + "}", encoding="ascii")
    source = spec("hwpx", "06-upstream.hwpx", "hwp_native", "희소 자원의 차등 분배", "희소 자원의 차등 분배")
    shutil.copy2(hwpx, source)
    assert smoke._digest(source) == UPSTREAM_SHA256
    specs[-1]["expected_pages"] = 4
    specs[-1]["search_page"] = 3
    (cache / "ui.json").write_text(json.dumps({"folder": str(inputs)}), encoding="utf-8")
    (root / "config.yaml").write_text(json.dumps({
        "project": {"input_dir": str(inputs), "work_dir": str(root / "work")},
        "grobid": {"url": ""}, "identify": {"resolve_missing_doi": False},
        "metadata": {"request_delay_sec": 0}}), encoding="utf-8")
    sources = [inputs / item["name"] for item in specs]
    for source in sources:
        source.chmod(stat.S_IREAD)
    return {"inputs": str(inputs), "cache": str(cache), "output_dir": str(inputs / "Markdown"),
            "names": [item["name"] for item in specs], "specs": specs,
            "source_digests_before": {str(p): smoke._digest(p) for p in sources}}


_BODY = r'''
  await command('Runtime.enable');await command('Page.enable');await command('Network.enable');
  await until(()=>evaluate('!!window.pywebview?.api?.state'),'native bridge');
  const state=await until(async()=>{const s=await evaluate('window.pywebview.api.state()');return s.files?.length===cfg.specs.length?s:null;},'private six inputs');
  if(path.resolve(state.path).toLowerCase()!==path.resolve(cfg.inputs).toLowerCase()||path.resolve(state.output_dir).toLowerCase()!==path.resolve(cfg.output_dir).toLowerCase())throw Error('Foreign folder/store output');
  if(JSON.stringify(state.files.map(f=>f.name))!==JSON.stringify(cfg.names)||state.files.some(f=>f.md_ready))throw Error('Unexpected fixtures or saved result before explicit extraction');
  const settings=await evaluate('window.pywebview.api.settings()');
  if(settings.current_version!==cfg.expected_version)throw Error('Frozen version mismatch');
  if(focusEmulation.requested){
    focusEmulation.browser=await command('Browser.getVersion');focusEmulation.before=await focusState('before');
    await command('Emulation.setFocusEmulationEnabled',{enabled:true});focusEmulation.enabled=true;focusEmulation.after=await focusState('after');
    if(!focusEmulation.after.valid)throw Error('Focus emulation unavailable');
  }
  if(!(await focusState('native-input')).valid)throw Error('Native page must be focused and visible for trusted input');
  const requests=[];let phase='viewer';
  ws.addEventListener('message',event=>{const message=JSON.parse(String(event.data));if(message.method==='Network.requestWillBeSent')requests.push({phase,url:message.params.request.url});});
  async function click(selector){
    const point=await evaluate('(()=>{const n=document.querySelector('+JSON.stringify(selector)+');if(!n||n.disabled)throw Error("Missing enabled direct action");const r=n.getBoundingClientRect();if(r.width<=0||r.height<=0)throw Error("Hidden direct action");return{x:r.left+r.width/2,y:r.top+r.height/2};})()');
    await command('Input.dispatchMouseEvent',{type:'mouseMoved',...point,button:'none'});
    await command('Input.dispatchMouseEvent',{type:'mousePressed',...point,button:'left',clickCount:1});
    await command('Input.dispatchMouseEvent',{type:'mouseReleased',...point,button:'left',clickCount:1});
  }
  const formats=[];
  for(let i=0;i<cfg.specs.length;i++){
    const spec=cfg.specs[i];
    await evaluate('(()=>{const spec='+JSON.stringify(spec)+';window.__nativeReady=null;window.__nativeStarted=null;const tick=()=>{if(document.querySelector("#reader-name").textContent===spec.name&&!document.querySelector("#source-pane").hidden){const native=document.querySelector("#native-content,#native-sheet,#source-text,#source-markdown"),image=document.querySelector(".pdf-page img");const ready=spec.format==="hwpx"?image?.complete&&image.naturalWidth>0:native?.textContent.includes(spec.marker);if(ready){window.__nativeReady={first_display_ms:performance.now()-window.__nativeStarted,clock_boundary:"trusted-filename-click-to-dom-or-png-decode",trusted_click:true,marker_found:spec.format==="hwpx"?null:true,decoded:spec.format==="hwpx"?true:null};return;}}requestAnimationFrame(tick);};const node=[...document.querySelectorAll(".file-row")].find(n=>n.querySelector(".filename").textContent===spec.name).querySelector(".filename");node.addEventListener("click",event=>{if(!event.isTrusted)throw Error("Untrusted fixture activation");window.__nativeStarted=performance.now();requestAnimationFrame(tick);},{once:true,capture:true});return true;})()');
    await click('.file-row[data-index="'+i+'"] .filename');
    const display=await until(()=>evaluate('window.__nativeReady'),'real '+spec.format+' display');
    if(!Number.isFinite(display.first_display_ms)||display.first_display_ms<=0)throw Error('Invalid first display clock');
    display.target_ms=2000;display.performance_ok=display.first_display_ms<=2000;
    if(!display.performance_ok&&!cfg.diagnostic_continue_target_miss)throw Error('First visible source exceeded 2000ms: '+JSON.stringify({spec,display}));
    const sourceRequests=requests.filter(r=>r.url.includes('/open?')&&new URL(r.url).searchParams.get('p')===String(i)&&new URL(r.url).searchParams.get('mode')==='source');
    if(sourceRequests.length!==1)throw Error('Expected one original metadata request per document');
    const metadata=await evaluate('fetch('+JSON.stringify(sourceRequests[0].url)+').then(r=>r.json())');
    if(metadata.source_kind!==spec.source_kind||!metadata.source_revision)throw Error('Native source kind or revision mismatch');
    const initialPages=requests.filter(r=>r.url.includes('/hwp-page?')).map(r=>Number(new URL(r.url).searchParams.get('i')));
    const controls=await evaluate('({extract_disabled:document.querySelector("#extract").disabled,md_disabled:document.querySelector("#reader-md").disabled,visible_status:[...document.querySelectorAll(".status,.row-gauge")].some(n=>n.getClientRects().length),source_visible:!document.querySelector("#source-pane").hidden,primary_more:!!document.querySelector(".reader-tools details")})');
    if(controls.visible_status||!controls.source_visible||controls.primary_more||controls.extract_disabled!==['txt','md'].includes(spec.format))throw Error('Viewer actions/state noise mismatch: '+JSON.stringify(controls));
    if(spec.format==='xlsx'){
      const cell=await evaluate('(()=>{const n=document.querySelector("#native-sheet td[data-row=\\"1\\"][data-col=\\"2\\"]");return{text:n?.textContent,title:n?.title};})()');
      if(cell.text!=='42'||!cell.title.includes('21*2'))throw Error('Cached formula display/raw formula lost');
      display.cached_formula=cell;
    }
    if(spec.format==='rtf'&&!await evaluate('document.querySelector("#native-content").textContent.includes("한글")'))throw Error('RTF Korean original lost');
    await trustedKey('f',true);await click('#search');await command('Input.insertText',{text:spec.query});await click('#search-form button[type=submit]');
    await until(()=>evaluate('document.querySelectorAll(".search-hit").length>0'),'native original search results');
    await until(()=>evaluate('(()=>{const s='+JSON.stringify(spec)+';if(s.format==="hwpx")return document.querySelector("#source-position").textContent.startsWith("4 /");return !!document.querySelector("#source-scroll .native-location,#source-scroll .text-location");})()'),'source search navigation');
    const search=await evaluate('({hits:document.querySelectorAll(".search-hit").length,position:document.querySelector("#source-position").textContent,scroll:document.querySelector("#source-scroll").scrollTop,highlight:document.querySelector("#source-scroll .native-location,#source-scroll mark")?.textContent||"",snippet:document.querySelector(".search-hit")?.textContent||""})');
    if(!search.snippet.includes(spec.query)||search.scroll<=0||spec.format!=='hwpx'&&!search.highlight.includes(spec.query))throw Error('Search did not navigate to exact source content');
    if(spec.format==='hwpx'){
      if(metadata.pages.length!==4||metadata.pages.some(p=>!Array.isArray(p)||p.length!==2||!p.every(n=>Number.isFinite(n)&&n>0)))throw Error('Invalid exact HWP page sizes');
      if(!initialPages.includes(0)||new Set(initialPages).size>=metadata.pages.length)throw Error('HWP rasterized every page before navigation');
      await until(()=>evaluate('document.querySelector(".pdf-page:last-child img")?.complete&&document.querySelector(".pdf-page:last-child img").naturalWidth>0'),'actual last HWP PNG decode');
      const last=await evaluate('(()=>{const image=document.querySelector(".pdf-page:last-child img");return{url:image.src,natural_width:image.naturalWidth,natural_height:image.naturalHeight};})()');
      const u=new URL(last.url);
      if(u.searchParams.get('i')!=='3'||u.searchParams.get('r')!==metadata.source_revision||!u.searchParams.get('viewer')||!u.searchParams.get('generation'))throw Error('HWP last source identity mismatch');
      display.hwpx={pages_count:metadata.pages.length,pages_px:metadata.pages,initial_requested_pages:initialPages,last_page_decoded:true,last,search_page:3,source_revision:metadata.source_revision};
      display.marker_found=search.snippet.includes(spec.marker);
    }
    await trustedKey('Escape');
    const wheel=await evaluate('(()=>{const n=document.querySelector("#source-scroll"),r=n.getBoundingClientRect();window.__nativeWheel=null;n.addEventListener("wheel",e=>window.__nativeWheel={trusted:e.isTrusted,prevented:e.defaultPrevented},{once:true});return{x:r.left+Math.min(100,r.width/2),y:r.top+100};})()');
    await command('Input.dispatchMouseEvent',{type:'mouseWheel',...wheel,deltaX:0,deltaY:-120,modifiers:2});
    await until(()=>evaluate('document.querySelector("#source-zoom-reset").textContent==="110%"'),'trusted original zoom 110');
    const zoom=await evaluate('({...window.__nativeWheel,label:document.querySelector("#source-zoom-reset").textContent})');
    await click('#source-zoom-reset');await until(()=>evaluate('document.querySelector("#source-zoom-reset").textContent==="100%"'),'direct zoom reset');
    zoom.reset='100%';if(!zoom.trusted||!zoom.prevented)throw Error('Source zoom input was not trusted/prevented');
    if(spec.format==='md'){
      await click('#source-code-toggle');await until(()=>evaluate('!!document.querySelector("#source-text")'),'direct Markdown source code');
      const raw=await evaluate('document.querySelector("#source-text").textContent');
      if(raw!==spec.raw_text)throw Error('Raw Markdown changed in view');
      await click('#reader-md');await until(()=>evaluate('!!document.querySelector("#source-markdown h1")'),'direct Markdown reading view');
      display.markdown={raw_exact:true,rendered:true,direct_top_view:true};
    }
    formats.push({format:spec.format,name:spec.name,source_kind:metadata.source_kind,...display,search,zoom,extract_disabled:controls.extract_disabled});
    completedResult={ok:false,functional_ok:false,stage:'format-complete',formats,focus_emulation:focusEmulation};
    fs.writeFileSync(cfg.result,JSON.stringify(completedResult,null,2),'utf8');
    await optionalScreenshot(path.join(cfg.screenshots,spec.format+'.png'));
  }
  const viewerRequests=requests.filter(r=>r.phase==='viewer'),forbidden=viewerRequests.filter(r=>r.url.includes('/page?')||r.url.includes('/document-map?')||new URL(r.url).searchParams.get('renderer')==='pdf');
  if(forbidden.length)throw Error('Native viewer invoked full PDF/map: '+JSON.stringify(forbidden));
  const beforeExtract=await evaluate('window.pywebview.api.state()');
  if(beforeExtract.files.some(f=>f.md_ready)||beforeExtract.busy)throw Error('Viewing originals triggered automatic extraction');
  phase='explicit-md';await click('.file-row[data-index="0"] .filename');await until(()=>evaluate('!!document.querySelector("#native-content")'),'DOCX original for explicit MD');
  await click('#extract');await until(async()=>{const s=await evaluate('window.pywebview.api.state()');return s.files[0].md_ready?s:null;},'explicit top MD extraction saved');
  await click('#reader-md');await until(()=>evaluate('document.querySelector("#document").textContent.includes("FROZEN_NATIVE_DOCX_127")&&!document.querySelector("#result-pane").hidden'),'one-click top MD');
  const mdAction=await evaluate('({saved:document.querySelector("#document").textContent.includes("FROZEN_NATIVE_DOCX_127"),same_pane:!document.querySelector("#list-view").hidden,source_hidden:document.querySelector("#source-pane").hidden,button_visible:document.querySelector("#reader-md").getClientRects().length>0})');
  if(!mdAction.saved||!mdAction.same_pane||!mdAction.source_hidden||!mdAction.button_visible)throw Error('Direct MD same-pane action failed');
  await click('#source-tab');await until(()=>evaluate('!document.querySelector("#source-pane").hidden'),'direct original return');
  const direct=await evaluate('({open:!!document.querySelector("#open-files")?.getClientRects().length,extract:!!document.querySelector("#extract")?.getClientRects().length,md:!!document.querySelector("#reader-md")?.getClientRects().length,original:!!document.querySelector("#source-tab")?.getClientRects().length,zoom:!!document.querySelector("#source-zoom-reset")?.getClientRects().length})');
  if(Object.values(direct).some(v=>v!==true))throw Error('Primary action requires a hidden menu');
  const targetMisses=formats.filter(f=>!f.performance_ok).map(f=>({format:f.format,first_display_ms:f.first_display_ms,target_ms:f.target_ms}));
  completedResult={ok:targetMisses.length===0,functional_ok:true,performance_ok:targetMisses.length===0,target_misses:targetMisses,target_url:target.url,bridge:true,version:settings.current_version,formats,controls:direct,viewer_phase:{automatic_extract:false,full_pdf_requests:forbidden.length,requests:viewerRequests},explicit_md:mdAction,screenshots,focus_emulation:focusEmulation,hidden_decode_is_not_visible_os_paint:true};
  fs.writeFileSync(cfg.result,JSON.stringify(completedResult,null,2),'utf8');
  if(targetMisses.length)throw Error('Functional checks passed with native latency target misses: '+JSON.stringify(targetMisses));
'''


def cdp_script():
    marker = "  await command('Runtime.enable');await command('Page.enable');"
    tail = "} catch(error){\n  fs.writeFileSync(cfg.result,"
    assert smoke._CDP_SCRIPT.count(marker) == 1
    assert smoke._CDP_SCRIPT.count(tail) == 1
    return smoke._CDP_SCRIPT.split(marker)[0] + _BODY + tail + smoke._CDP_SCRIPT.split(tail)[1]


def runtime_assets(exe):
    result = smoke._runtime_assets(exe)
    engines = list(exe.resolve().parent.rglob("rhwp.exe"))
    if len(engines) != 1 or smoke._digest(engines[0]) != RHWP_SHA256:
        raise AssertionError("Expected one pinned shipped rHWP native renderer")
    engine = engines[0]
    result[str(engine.relative_to(exe.parent))] = {"bytes": engine.stat().st_size, "sha256": smoke._digest(engine)}
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exe", type=Path)
    parser.add_argument("--expected-exe-sha256")
    parser.add_argument("--out-dir", type=Path)
    parser.add_argument("--node", default=shutil.which("node"))
    parser.add_argument("--hwpx", type=Path, default=smoke._cache_root() / "native-viewer-upstream/exam_social.hwpx")
    parser.add_argument("--timeout", type=float, default=180)
    parser.add_argument("--focus-emulation-test", action="store_true")
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--diagnostic-continue-target-miss", action="store_true",
                        help="Complete functionality after target misses; receipt remains unsuccessful")
    args = parser.parse_args()
    if not args.node or args.timeout <= 0 or sys.platform != "win32":
        parser.error("Windows and Node.js WebSocket support are required")
    if not args.hwpx.is_file() or smoke._digest(args.hwpx) != UPSTREAM_SHA256:
        parser.error("Pass the exact pinned upstream HWPX fixture; no automatic download")
    if not args.prepare_only and (not args.exe or not args.exe.is_file() or
            not args.expected_exe_sha256 or smoke._digest(args.exe) != args.expected_exe_sha256.lower()):
        parser.error("The actual EXE must match the explicit expected SHA256 before execution")
    cache = smoke._cache_root().resolve()
    output = (args.out_dir or cache / "frozen-native-viewer" / (datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:8])).resolve()
    if not output.is_relative_to(cache):
        parser.error("Receipts must stay inside configured project DevCache")
    output.mkdir(parents=True, exist_ok=False)
    original = fingerprint(args.hwpx)
    receipt = {"schema_version": 1, "kind": "native-viewer-frozen", "ok": False,
               "mode": "prepare-only" if args.prepare_only else "frozen-exe",
               "started_at": datetime.now(timezone.utc).isoformat(), "output_dir": str(output),
               "hwpx_original": {"before": original, "provenance": {
                   "url": "https://raw.githubusercontent.com/edwardkim/rhwp/" + UPSTREAM_COMMIT + "/" + UPSTREAM_PATH,
                   "commit": UPSTREAM_COMMIT, "repository_path": UPSTREAM_PATH}},
               "harness_sha256": smoke._digest(Path(__file__)),
               "reused_harness_sha256": smoke._digest(Path(smoke.__file__))}
    process = job = root = None
    fixture = None
    try:
        temporary = tempfile.TemporaryDirectory(prefix="pnx-frozen-native-viewer-")
        root = Path(temporary.name)
        try:
            fixture = fixtures(root, args.hwpx)
            receipt.update(fixtures=fixture, isolated_root=str(root), source_readonly=True)
            from pubnexus.version import VERSION
            config = {**fixture, "timeout": args.timeout, "expected_version": VERSION,
                      "diagnostic_continue_target_miss": args.diagnostic_continue_target_miss,
                      "focus_emulation_test": args.focus_emulation_test, "result": str(root / "cdp-result.json"),
                      "screenshots": str(output), "probe": str(output / "cdp-probe.json")}
            script = root / "native-viewer.mjs"
            script.write_text(cdp_script(), encoding="utf-8")
            check = subprocess.run([args.node, "--check", str(script)], capture_output=True, timeout=10,
                                   creationflags=subprocess.CREATE_NO_WINDOW)
            if check.returncode:
                raise AssertionError(check.stderr.decode("utf-8", "replace"))
            if args.prepare_only:
                receipt.update(ok=True, exe_launched=False, fixtures_validated=True)
            else:
                receipt["runtime_assets_before"] = runtime_assets(args.exe)
                source_web = Path(__file__).resolve().parents[1] / "web"
                for name in ("presentation-frame.html", "presentation-frame.mjs"):
                    if receipt["runtime_assets_before"][name]["sha256"] != smoke._digest(source_web / name):
                        raise AssertionError("Frozen runtime frame differs from approved source")
                with socket.socket() as reserve:
                    reserve.bind(("127.0.0.1", 0))
                    config["port"] = reserve.getsockname()[1]
                config_path = root / "cdp-config.json"
                config_path.write_text(json.dumps(config), encoding="utf-8")
                environment = os.environ.copy()
                environment.update(PUBNEXUS_STORE=fixture["cache"], PUBNEXUS_ROOT=str(root),
                                   PUBNEXUS_OUTPUT=fixture["output_dir"],
                                   WEBVIEW2_ADDITIONAL_BROWSER_ARGUMENTS="--remote-debugging-port=" + str(config["port"]),
                                   WEBVIEW2_USER_DATA_FOLDER=str(root / "profile"), PYTHONDONTWRITEBYTECODE="1",
                                   PNX_TEST_DEBUG_PORT=str(config["port"]))
                startup = subprocess.STARTUPINFO()
                startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                startup.wShowWindow = subprocess.SW_HIDE
                with (output / "exe.log").open("wb") as log:
                    process = subprocess.Popen([str(args.exe.resolve()), "--hidden"], cwd=root, env=environment,
                                               stdout=log, stderr=log, stdin=subprocess.DEVNULL,
                                               creationflags=subprocess.CREATE_NO_WINDOW, startupinfo=startup)
                    job = smoke._PrivateJob(process)
                    receipt.update(exe=str(args.exe.resolve()), exe_sha256=smoke._digest(args.exe),
                                   expected_exe_sha256=args.expected_exe_sha256.lower(), pid=process.pid,
                                   job_bound=True, exe_launched=True)
                    (output / "running.json").write_text(json.dumps({"own_pid": process.pid, "config": config},
                                                                  ensure_ascii=False, indent=2), encoding="utf-8")
                    completed = subprocess.run([args.node, str(script), str(config_path)], capture_output=True,
                                               text=True, encoding="utf-8", errors="replace", timeout=args.timeout + 20,
                                               creationflags=subprocess.CREATE_NO_WINDOW)
                    (output / "cdp.log").write_text((completed.stdout or "") + (completed.stderr or ""), encoding="utf-8")
                    if Path(config["result"]).is_file():
                        receipt["webview"] = json.loads(Path(config["result"]).read_text(encoding="utf-8"))
                    if completed.returncode or not receipt.get("webview", {}).get("ok"):
                        raise AssertionError("Frozen native WebView failed: " + (completed.stderr or "")[-2500:])
                    receipt["ok"] = True
        finally:
            cleanup_errors = []
            def attempt(label, action):
                try:
                    action()
                except Exception as error:
                    cleanup_errors.append(label + ": " + type(error).__name__ + ": " + str(error))
                    receipt["ok"] = False
            def check_sources():
                if fixture:
                    after = {name: smoke._digest(Path(name)) for name in fixture["source_digests_before"]}
                    receipt.update(source_digests_after=after, sources_unchanged=after == fixture["source_digests_before"])
                    if not receipt["sources_unchanged"]:
                        raise AssertionError("A private source changed")
            def collect_pids():
                if job:
                    receipt["owned_job_pids_before_cleanup"] = job.pids()
            def close_job():
                if job:
                    job.close()
            def fallback_stop():
                if process and process.poll() is None:
                    subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], timeout=10,
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                   creationflags=subprocess.CREATE_NO_WINDOW)
            def check_process():
                if process:
                    process.wait(timeout=10)
                    receipt["exe_returncode_after_cleanup"] = process.returncode
                    import psutil
                    alive = [pid for pid in receipt.get("owned_job_pids_before_cleanup", []) if psutil.pid_exists(pid)]
                    receipt["owned_pids_terminated"] = not alive
                    if alive:
                        raise AssertionError("Owned process survived: " + str(alive))
            attempt("source integrity", check_sources)
            attempt("owned PID inventory", collect_pids)
            attempt("owned Job close", close_job)
            attempt("owned process fallback", fallback_stop)
            attempt("owned process completion", check_process)
            if fixture:
                for source in fixture["source_digests_before"]:
                    attempt("private source writable", lambda source=source: Path(source).chmod(stat.S_IREAD | stat.S_IWRITE))
            attempt("private TEMP/profile cleanup", temporary.cleanup)
            receipt["cleanup_errors"] = cleanup_errors
    except Exception as error:
        receipt.update(ok=False, error=type(error).__name__ + ": " + str(error))
    finally:
        try:
            after = fingerprint(args.hwpx)
            receipt["hwpx_original"]["after"] = after
            receipt["hwpx_original"]["unchanged"] = after == original
            if after != original:
                receipt.update(ok=False, original_error="Upstream original changed")
            if "runtime_assets_before" in receipt:
                receipt["runtime_assets_after"] = runtime_assets(args.exe)
                receipt["runtime_assets_unchanged"] = receipt["runtime_assets_before"] == receipt["runtime_assets_after"]
                if not receipt["runtime_assets_unchanged"]:
                    receipt.update(ok=False, runtime_assets_error="Frozen runtime assets changed")
            receipt["temporary_inputs_removed"] = root is not None and not root.exists()
            if not receipt["temporary_inputs_removed"]:
                receipt.update(ok=False, cleanup_error="Private TEMP/profile not removed")
        except Exception as error:
            receipt.update(ok=False, final_integrity_error=type(error).__name__ + ": " + str(error))
        receipt["finished_at"] = datetime.now(timezone.utc).isoformat()
        (output / "receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": receipt["ok"], "receipt": str(output / "receipt.json"), "error": receipt.get("error")},
                     ensure_ascii=False), flush=True)
    return 0 if receipt["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
