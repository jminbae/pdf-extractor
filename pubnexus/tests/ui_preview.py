"""Offline actual-app browser harness. Run manually, close with Ctrl+C.

All documents, exports and cache live in TemporaryDirectory. The HTTP RPC
bridge exists only here; the production app keeps its pywebview bridge.
"""
import importlib.util
import json
import os
import queue
import sys
import tempfile
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path

sys.dont_write_bytecode = True
spec = importlib.util.spec_from_file_location('preview_app', Path(__file__).resolve().parents[1] / 'app.py')
appmod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(appmod)
from pubnexus import general_pdf, store
import fitz


BRIDGE = '''<script>
window.pywebview={api:new Proxy({}, {get:(_,name)=>(...args)=>fetch('/preview-api',{
method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({name,args})
}).then(r=>r.json())})};
window.addEventListener('load',()=>{window.dispatchEvent(new Event('pywebviewready'));
let alive=true;window.addEventListener('pagehide',()=>alive=false);
async function poll(){if(!alive)return;try{for(const m of await fetch('/preview-events').then(r=>r.json()))window.pnx?.on(m);}finally{if(alive)setTimeout(poll,120);}}poll();});
</script>'''


def main():
    with tempfile.TemporaryDirectory(prefix='document-ui-') as folder:
        root = Path(folder)
        inputs = root / 'inputs'
        inputs.mkdir()
        os.environ['PUBNEXUS_STORE'] = str(root / 'cache')
        os.environ['PUBNEXUS_OUTPUT'] = str(inputs / 'Markdown')
        appmod.utils.load_config = lambda: {}
        events = queue.Queue()
        app = appmod.App()
        app.grobid.state = 'off'
        app.push = lambda kind, payload: events.put({'kind': kind, **payload})
        sources = []
        for name, texts in [('Book.pdf', ['Chapter one. Source navigation example.',
                                         'Chapter two. Unique search needle 739.',
                                         'Chapter three. Last source location.']),
                            ('Another document.pdf', ['A separate document with preserved native text.'])]:
            path = inputs / name
            with fitz.open() as pdf:
                for text in texts:
                    page = pdf.new_page()
                    page.insert_text((72, 72), text)
                pdf.save(path)
            sha = store.file_sha1(path)
            doc, _ = general_pdf.ensure_fulltext({'sha1': sha, 'meta': {'title': path.stem}}, path)
            store.save(sha, doc, path)
            sources.append(path)
        from docx import Document
        word = Document()
        word.add_heading('Word source comparison', 0)
        word.add_paragraph('Unique Word source location 8461.')
        table = word.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [('Name', 'Value'), ('Sample', '23.45')]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        path = inputs / 'Word document.docx'
        word.save(path)
        sources.append(path)
        from pptx import Presentation
        slides = Presentation()
        slide = slides.slides.add_slide(slides.slide_layouts[1])
        slide.shapes.title.text = 'Presentation source'
        slide.placeholders[1].text = 'Unique slide content 5682.'
        path = inputs / 'Presentation.pptx'
        slides.save(path)
        sources.append(path)
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(['Source', 'Value'])
        workbook.active.append(['Sample', 8461])
        path = inputs / 'Workbook.xlsx'
        workbook.save(path)
        workbook.close()
        sources.append(path)
        path = inputs / 'Table.csv'
        path.write_text('Name,Value\nPreview boundary,42\n', encoding='utf-8')
        sources.append(path)
        api = appmod.Api(app)

        # This browser fixture never checks a release or launches an installer.
        # Settings persistence and update events stay inside its private TEMP.
        settings_path = root / 'ui-settings.json'
        fixture_settings = {'ok': True, 'output_dir': str(app.output_dir),
                            'auto_update': True, 'current_version': 'UI fixture',
                            'update': {'status': 'idle', 'message': '로컬 화면 시험 · 실제 업데이트는 실행하지 않습니다.'}}

        def settings():
            return json.loads(json.dumps(fixture_settings))

        def save_settings(**changes):
            fixture_settings.update(changes)
            app.output_dir = Path(fixture_settings['output_dir'])
            settings_path.write_text(json.dumps(fixture_settings, ensure_ascii=False), encoding='utf-8')
            events.put({'kind': 'settings', **settings()})
            return settings()

        def check_update():
            update = {'status': 'ready', 'version': 'UI fixture next',
                      'message': '로컬 화면 시험 · 실제 업데이트는 실행하지 않습니다.'}
            fixture_settings['update'] = update
            events.put({'kind': 'update', **update})
            return settings()

        def add_fixtures():
            app.add_files(sources)
            return {'ok': True, 'files': app.file_rows(), 'path': str(inputs),
                    'output_dir': str(app.output_dir)}

        methods = {name: getattr(api, name) for name in ('state', 'extract_selected', 'cancel', 'search', 'search_source', 'location')}
        methods.update(pick_files=add_fixtures, pick_folder=add_fixtures,
                       reveal=lambda path: {'ok': True, 'path': path},
                       reveal_output=lambda: {'ok': True, 'path': str(app.output_dir)},
                       open_file=lambda index, kind: {'ok': True, 'index': index, 'kind': kind})
        methods.update(settings=settings,
                       choose_output_folder=lambda: save_settings(output_dir=str(inputs / '새 저장 폴더')),
                       reset_output_folder=lambda: save_settings(output_dir=str(inputs / 'Markdown')),
                       set_auto_update=lambda enabled: save_settings(auto_update=bool(enabled)),
                       check_update=check_update,
                       install_update=lambda: {'ok': False, 'why': '화면 시험에서는 설치 프로그램을 실행하지 않습니다.'})

        class Handler(appmod.Handler):
            def do_GET(self):
                if self.path == '/preview-events':
                    batch = []
                    while True:
                        try:
                            batch.append(events.get_nowait())
                        except queue.Empty:
                            break
                    self._json(batch)
                elif self.path == '/':
                    self._send(appmod.INDEX_HTML.replace('</head>', BRIDGE + '</head>').encode(), 'text/html; charset=utf-8')
                else:
                    super().do_GET()

            def do_POST(self):
                if self.path != '/preview-api':
                    self.send_error(404)
                    return
                try:
                    length = int(self.headers.get('Content-Length', '0'))
                    if not 0 < length <= 65536:
                        raise ValueError('Invalid bridge request')
                    data = json.loads(self.rfile.read(length))
                    method = methods[data['name']]
                    self._json(method(*data.get('args', [])))
                except Exception as error:
                    self._json({'ok': False, 'why': str(error)})

        Handler.app = app
        server = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        def shutdown():
            threading.Thread(target=server.shutdown, daemon=True).start()
            return {'ok': True}
        methods['shutdown'] = shutdown
        print(json.dumps({'url': f'http://127.0.0.1:{server.server_port}/', 'temp': folder}), flush=True)
        try:
            server.serve_forever()
        except KeyboardInterrupt:
            pass
        finally:
            app.cancel.set()
            if app.worker:
                app.worker.join(timeout=30)
            app.renderer.close_all()
            server.server_close()


if __name__ == '__main__':
    main()
