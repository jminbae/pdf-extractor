"""Hidden frozen EXE smoke through its real WebView2 JS bridge.

Run with the development Python, never from unittest discovery. Fixtures and
the app store/profile are TEMP-only; receipts and screenshots use DevCache.
"""
from __future__ import annotations

import argparse
import ctypes
from ctypes import wintypes
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import shutil
import socket
import stat
import subprocess
import sys
import tempfile
import time
import uuid

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))


_PRESENTATION_VISUAL_EXPRESSION = "(()=>{const viewport=document.querySelector('#viewport'),view=viewport.getBoundingClientRect(),slots=[...document.querySelectorAll('.slide-slot')],overlays=[...document.querySelectorAll('[data-pptx-search-highlight]')];const area=rect=>Math.max(0,Math.min(rect.right,view.right)-Math.max(rect.left,view.left))*Math.max(0,Math.min(rect.bottom,view.bottom)-Math.max(rect.top,view.top));const visible=slots.map(slot=>({slot,area:area(slot.getBoundingClientRect())})).filter(item=>item.area>0).sort((a,b)=>b.area-a.area);const highlighted=overlays.length===1?overlays[0].closest('.slide-slot'):null,slot=highlighted||visible[0]?.slot,images=slot?[...slot.querySelectorAll('img')]:[];return {page:slot?Number(slot.dataset.slideIndex)+1:null,highlighted_page:highlighted?Number(highlighted.dataset.slideIndex)+1:null,visible_pages:visible.map(item=>Number(item.slot.dataset.slideIndex)+1),highlights:overlays.length,highlight_intersects_viewport:overlays.length===1&&area(overlays[0].getBoundingClientRect())>0,width:slot?.getBoundingClientRect().width,paper:slot?getComputedStyle(slot).backgroundColor:null,filters:[...document.images].map(image=>getComputedStyle(image).filter),image_count:images.length,images_decoded:images.length>0&&images.every(image=>image.complete&&image.naturalWidth>0),slot_count:slots.length,mounted_count:slots.filter(slot=>slot.childElementCount>0).length,text:document.body.innerText};})()"

_CDP_SCRIPT = r'''
import fs from 'node:fs';
import path from 'node:path';
const cfg=JSON.parse(fs.readFileSync(process.argv[2],'utf8'));
const deadline=Date.now()+cfg.timeout*1000;
const PRESENTATION_VISUAL=__PNX_FRAME_VISUAL__;
const pause=ms=>new Promise(r=>setTimeout(r,ms));
async function until(fn,label){while(Date.now()<deadline){const value=await fn();if(value)return value;await pause(100);}throw Error('Timed out: '+label);}
let ws;
let command;
let focusEmulation={requested:!!cfg.focus_emulation_test,production:false};
let background=null;
let completedResult=null;
let previousProbe='';
function probe(value){const serialized=JSON.stringify(value,null,2);if(serialized!==previousProbe){fs.writeFileSync(cfg.probe,serialized,'utf8');previousProbe=serialized;}}
try {
  const target=await until(async()=>{try{const response=await fetch(`http://127.0.0.1:${cfg.port}/json/list`,{signal:AbortSignal.timeout(1500)});const pages=await response.json();probe({status:response.status,targets:pages.map(p=>({id:p.id,type:p.type,url:p.url,title:p.title}))});return pages.find(p=>p.type==='page'&&/^http:\/\/127\.0\.0\.1:\d+\/$/.test(p.url));}catch(error){probe({error:String(error),code:error.cause?.code||''});return null;}},'private WebView2 CDP target');
  ws=new WebSocket(target.webSocketDebuggerUrl);
  await new Promise((resolve,reject)=>{const timer=setTimeout(()=>reject(Error('CDP connect timeout')),5000);ws.addEventListener('open',()=>{clearTimeout(timer);resolve();},{once:true});ws.addEventListener('error',()=>{clearTimeout(timer);reject(Error('CDP connection failed'));},{once:true});});
  let next=0;const pending=new Map();
  ws.addEventListener('message',event=>{const message=JSON.parse(String(event.data));const item=pending.get(message.id);if(item){pending.delete(message.id);clearTimeout(item.timer);message.error?item.reject(Error(JSON.stringify(message.error))):item.resolve(message.result);}});
  command=(method,params={},timeoutMs=15000,sessionId=null)=>new Promise((resolve,reject)=>{const id=++next;const timer=setTimeout(()=>{pending.delete(id);reject(Error('CDP command timeout: '+method));},timeoutMs);pending.set(id,{resolve,reject,timer});ws.send(JSON.stringify({id,method,params,...(sessionId?{sessionId}:{})}));});
  async function evaluate(expression){const result=await command('Runtime.evaluate',{expression,awaitPromise:true,returnByValue:true});if(result.exceptionDetails)throw Error(JSON.stringify(result.exceptionDetails));return result.result?.value;}
  async function trustedKey(key,ctrl=false){
    const code=key==='Escape'?'Escape':'KeyF',keyCode=key==='Escape'?27:70;
    if(ctrl)await command('Input.dispatchKeyEvent',{type:'keyDown',key:'Control',code:'ControlLeft',windowsVirtualKeyCode:17,modifiers:2});
    try{await command('Input.dispatchKeyEvent',{type:'keyDown',key,code,windowsVirtualKeyCode:keyCode,modifiers:ctrl?2:0});await command('Input.dispatchKeyEvent',{type:'keyUp',key,code,windowsVirtualKeyCode:keyCode,modifiers:ctrl?2:0});}
    finally{if(ctrl)await command('Input.dispatchKeyEvent',{type:'keyUp',key:'Control',code:'ControlLeft',windowsVirtualKeyCode:17,modifiers:0});}
  }
  async function focusPresentation(){
    const point=await evaluate("(()=>{const box=document.querySelector('#presentation-frame').getBoundingClientRect();return{x:box.left+8,y:box.top+8};})()");
    await command('Input.dispatchMouseEvent',{type:'mousePressed',...point,button:'left',clickCount:1});await command('Input.dispatchMouseEvent',{type:'mouseReleased',...point,button:'left',clickCount:1});
    if(!await evaluate("document.activeElement?.id==='presentation-frame'"))throw Error('Trusted input did not focus the owned presentation iframe');
  }
  async function focusState(stage){
    const state=await evaluate('({has_focus:document.hasFocus(),visibility:document.visibilityState,hidden:document.hidden})');
    const sample={stage,...state,valid:state?.has_focus===true&&state?.visibility==='visible'&&state?.hidden===false};
    if(focusEmulation.requested)(focusEmulation.samples??=[]).push(sample);
    return sample;
  }
  const screenshots=[];
  async function optionalScreenshot(destination){
    try{const shot=await command('Page.captureScreenshot',{format:'png',captureBeyondViewport:false},3000);fs.writeFileSync(destination,Buffer.from(shot.data,'base64'));screenshots.push({path:destination,saved:true});}
    catch(error){screenshots.push({path:destination,saved:false,error:String(error)});}
  }
  const executionContexts=new Map();
  ws.addEventListener('message',event=>{const message=JSON.parse(String(event.data));if(message.method==='Runtime.executionContextCreated')executionContexts.set(message.params.context.id,message.params.context);else if(message.method==='Runtime.executionContextDestroyed')executionContexts.delete(message.params.executionContextId);else if(message.method==='Runtime.executionContextsCleared')executionContexts.clear();});
  const frameContextProbe={};
  async function frameEvaluate(expression){
    const expected=await evaluate("document.querySelector('#presentation-frame')?.src");if(!expected)throw Error('No owned presentation iframe');
    const tree=await command('Page.getFrameTree');frameContextProbe.tree=tree;
    const visit=node=>node.frame.url.split('#')[0]===expected.split('#')[0]?node.frame:(node.childFrames||[]).map(visit).find(Boolean);const frame=visit(tree.frameTree);
    const context=frame&&[...executionContexts.values()].find(c=>c.auxData?.isDefault&&c.auxData?.frameId===frame.id);
    let sessionId=null;
    try{
      if(!context){const targets=await command('Target.getTargets');frameContextProbe.targets=targets.targetInfos.map(t=>({id:t.targetId,type:t.type,url:t.url,openerFrameId:t.openerFrameId}));const candidates=targets.targetInfos.filter(t=>t.type==='iframe'&&t.url.split('#')[0]===expected.split('#')[0]);if(candidates.length!==1)throw Error('Expected exactly one owned opaque iframe target: '+JSON.stringify(frameContextProbe));sessionId=(await command('Target.attachToTarget',{targetId:candidates[0].targetId,flatten:true})).sessionId;}
      const checked=`(()=>{if(location.href!==${JSON.stringify(expected)})throw Error('Refusing a stale or foreign frame');return (${expression});})()`;
      const result=await command('Runtime.evaluate',{expression:checked,...(context?{contextId:context.id}:{}),awaitPromise:true,returnByValue:true},15000,sessionId);
      if(result.exceptionDetails)throw Error(JSON.stringify(result.exceptionDetails));return {...result.result?.value,context_transport:sessionId?'owned-opaque-iframe-target':'owned-default-frame-context'};
    }finally{if(sessionId)await command('Target.detachFromTarget',{sessionId});}
  }
  await command('Runtime.enable');await command('Page.enable');
  await until(()=>evaluate('!!(window.pywebview&&window.pywebview.api&&window.pywebview.api.state)'), 'pywebview bridge');
  const state=await until(async()=>{const s=await evaluate('window.pywebview.api.state()');return s.files?.length===cfg.names.length?s:null;},'isolated inputs');
  if(path.resolve(state.path).toLowerCase()!==path.resolve(cfg.inputs).toLowerCase())throw Error('Refusing to operate outside private fixture folder');
  if(path.resolve(state.output_dir||'').toLowerCase()!==path.resolve(cfg.output_dir).toLowerCase())throw Error('Refusing to save outside the explicit private output folder');
  if(JSON.stringify(state.files.map(f=>f.name))!==JSON.stringify(cfg.names))throw Error('Unexpected input list');
  if(state.files.some(f=>f.md_ready))throw Error('A private fixture was presented as saved Markdown before export');
  if(focusEmulation.requested){
    focusEmulation.browser=await command('Browser.getVersion');
    focusEmulation.before=await focusState('before');
    focusEmulation.protocol_result=await command('Emulation.setFocusEmulationEnabled',{enabled:true});
    focusEmulation.enabled=true;
    focusEmulation.after=await focusState('after');
    if(!focusEmulation.after.valid)throw Error('Focus emulation did not establish a focused visible page: '+JSON.stringify(focusEmulation.after));
  }
  await until(()=>evaluate(`document.querySelectorAll('.file-row .row-actions [data-action=read]').length===${cfg.names.length}`),'always-present Markdown buttons');
  const initialMdButtons=await evaluate("Array.from(document.querySelectorAll('.file-row .row-actions [data-action=read]')).map(b=>({label:b.textContent,disabled:b.disabled}))");
  if(initialMdButtons.some(b=>b.label!=='MD 보기'||!b.disabled))throw Error('Cached extraction enabled an unsaved Markdown button');
  const settings=await evaluate('window.pywebview.api.settings()');
  if(!settings?.ok||settings.current_version!==cfg.expected_version||typeof settings.auto_update!=='boolean'||path.resolve(settings.output_dir||'').toLowerCase()!==path.resolve(cfg.output_dir).toLowerCase())throw Error('Frozen settings API mismatch: '+JSON.stringify(settings));
  await evaluate("document.querySelector('#settings-open').click();true");
  await until(()=>evaluate(`document.querySelector('#settings-dialog')?.open&&document.querySelector('#settings-version').textContent===${JSON.stringify(cfg.expected_version)}&&document.querySelector('#settings-output').value===${JSON.stringify(cfg.output_dir)}`),'real settings dialog');
  const settingsUi=await evaluate("({output:document.querySelector('#settings-output').value,readonly:document.querySelector('#settings-output').readOnly,version:document.querySelector('#settings-version').textContent,auto:document.querySelector('#settings-auto').checked})");
  if(!settingsUi.readonly||settingsUi.auto!==settings.auto_update)throw Error('Frozen settings dialog is not an accurate read-only settings view');
  let manualUpdate=null;
  if(cfg.check_update){
    const check=await evaluate('window.pywebview.api.check_update()');if(!check?.ok)throw Error('Manual update check did not start: '+JSON.stringify(check));
    manualUpdate=await until(async()=>{const current=await evaluate('window.pywebview.api.settings()');const update=current.update||{};if(['idle','checking','downloading'].includes(update.status))return null;if(update.status!=='up_to_date')throw Error('Unexpected manual update state: '+JSON.stringify(update));return update;},'explicit manual update check');
  }
  await evaluate("document.querySelector('#settings-close').click();true");
  if(cfg.office_mode){
    const officeProof={};function officeCheckpoint(stage,data={}){Object.assign(officeProof,data,{stage});completedResult={ok:false,bridge:true,office:officeProof,focus_emulation:focusEmulation};fs.writeFileSync(cfg.result,JSON.stringify(completedResult,null,2),'utf8');}
    officeCheckpoint('starting');
    await command('Network.enable');const officeRequests=[];
    ws.addEventListener('message',event=>{const message=JSON.parse(String(event.data));if(message.method==='Network.requestWillBeSent')officeRequests.push(message.params.request.url);});
    const previewBegan=Date.now();
    await evaluate("document.querySelector('.file-row[data-index=\"0\"] .filename').dispatchEvent(new MouseEvent('dblclick',{bubbles:true}));true");
    const preparing=await evaluate("({reader:!document.querySelector('#reader').hidden,back:document.querySelector('#reader-back').textContent,extractEnabled:!document.querySelector('#reader-extract').disabled,globalVisible:!document.querySelector('#extract').hidden,globalLabel:document.querySelector('#extract').textContent})");
    if(!preparing.reader||!preparing.back.includes('목록')||!preparing.extractEnabled||!preparing.globalVisible||preparing.globalLabel!=='MD 추출')throw Error('Office preparation blocked reader actions: '+JSON.stringify(preparing));
    const slideCount=cfg.office_pages||3;
    await until(()=>evaluate(`(()=>{const frame=document.querySelector('#presentation-frame'),metrics=frame?JSON.parse(frame.dataset.metrics||'{}'):null;return frame?.dataset.ready==='true'&&Number.isFinite(metrics.render_width)&&metrics.render_width>=frame.getBoundingClientRect().width*.8;})()||document.querySelectorAll('#source-scroll .pdf-page').length===${slideCount}`),'actual PPTX source ready');
    const fastPresentation=await evaluate("!!document.querySelector('#presentation-frame')");
    if(cfg.require_presentation&&!fastPresentation)throw Error('Fast PPTX unexpectedly fell back to a full PDF');
    if(!fastPresentation)await until(()=>evaluate("Array.from(document.querySelectorAll('#source-scroll img')).some(i=>{i.loading='eager';return i.complete&&i.naturalWidth>0&&!i.hidden;})"),'real PPTX preview image');
    const previewDecodedMs=Date.now()-previewBegan;
    const original=await evaluate("(()=>{const frame=document.querySelector('#presentation-frame'),metrics=frame?JSON.parse(frame.dataset.metrics||'{}'):null;return{presentation:!!frame,metrics,frame_width:frame?.getBoundingClientRect().width,width:frame?metrics.render_width:document.querySelector('.pdf-page').getBoundingClientRect().width,body:document.body.getBoundingClientRect().width,pageLabel:document.querySelector('#page-input').getAttribute('aria-label'),total:document.querySelector('#page-total').textContent,label:document.querySelector('#source-label').textContent,editable:document.querySelectorAll('#reader [contenteditable=true],#reader textarea').length};})()");
    if(!Number.isFinite(original.width)||original.width<=0||original.presentation&&(!Number.isFinite(original.frame_width)||original.frame_width<200||original.width<original.frame_width*.8||original.width>original.frame_width*1.1)||original.editable||!original.pageLabel.includes('슬라이드')||original.total.trim()!=='/ '+slideCount||original.label!=='원본')throw Error('PPTX slide position label is wrong: '+JSON.stringify(original));
    if(fastPresentation&&(!original.metrics?.isolation?.opaque_parent||!original.metrics?.isolation?.native_bridge_absent||!Number.isInteger(original.metrics.resources?.images)||original.metrics.resources.images<1||original.metrics.resources?.fonts!=='loaded'))throw Error('Fast frame isolation/image decode/font ready failed: '+JSON.stringify(original));
    officeCheckpoint('source-ready-before-isolation',{original,preview_image_decoded_ms:previewDecodedMs,preview_target_ms:2000,preview_decode_within_target:previewDecodedMs<=2000,frame_context_probe:frameContextProbe});
    let frameSecurity=null;
    if(fastPresentation){frameSecurity=await frameEvaluate("(()=>{let parentBlocked=false;try{void parent.document.body;}catch{parentBlocked=true;}const host=window.chrome?.webview?.hostObjects;let hostKeys=[],hostKeysError=null;try{hostKeys=host?Object.keys(host):[];}catch(error){hostKeysError=String(error);}return {pywebview_api_present:!!window.pywebview?.api,parent_access_blocked:parentBlocked,chrome_webview_postMessage_type:typeof window.chrome?.webview?.postMessage,hostObjects_type:typeof host,hostObjects_keys:hostKeys,hostObjects_keys_error:hostKeysError,native_messages_sent:0};})()");if(frameSecurity.pywebview_api_present||!frameSecurity.parent_access_blocked)throw Error('Native presentation frame isolation failed: '+JSON.stringify(frameSecurity));}
    officeCheckpoint('source-ready',{frame_security:frameSecurity,original,preview_image_decoded_ms:previewDecodedMs,preview_target_ms:2000,preview_decode_within_target:previewDecodedMs<=2000,hidden_decode_is_not_visible_paint:true});
    if(cfg.require_presentation&&previewDecodedMs>2000)throw Error('First decoded PPTX preview exceeded 2000ms: '+previewDecodedMs);
    const zoomTarget=await evaluate("(()=>{const node=document.querySelector('#source-scroll'),box=node.getBoundingClientRect();window.__officeWheel=null;node.addEventListener('wheel',event=>window.__officeWheel={trusted:event.isTrusted,prevented:event.defaultPrevented},{once:true});return{x:box.left+Math.min(150,box.width/2),y:box.top+120};})()");
    await command('Input.dispatchMouseEvent',{type:'mouseMoved',x:zoomTarget.x,y:zoomTarget.y,button:'none'});
    await command('Input.dispatchMouseEvent',{type:'mouseWheel',x:zoomTarget.x,y:zoomTarget.y,deltaX:0,deltaY:-120,modifiers:2});
    await until(()=>evaluate("document.querySelector('#source-zoom-reset').textContent.trim()==='110%'&&(!document.querySelector('#presentation-frame')||JSON.parse(document.querySelector('#presentation-frame').dataset.metrics||'{}').zoom===110)"),'PPTX trusted original Ctrl-wheel');
    const sourceZoom=await evaluate("(()=>{const frame=document.querySelector('#presentation-frame'),metrics=frame?JSON.parse(frame.dataset.metrics||'{}'):null,frameEvent=frame?JSON.parse(frame.dataset.wheel||'null'):null,parentEvent=window.__officeWheel;return{width:frame?metrics.render_width:document.querySelector('.pdf-page').getBoundingClientRect().width,body:document.body.getBoundingClientRect().width,event:frameEvent||parentEvent,event_origin:frameEvent?'opaque-frame':'source-scroll',frame_event:frameEvent,parent_event:parentEvent,md:document.querySelector('#md-zoom-reset').textContent,filters:Array.from(document.querySelectorAll('#source-scroll img')).map(i=>getComputedStyle(i).filter)};})()");
    if(fastPresentation){const visual=await frameEvaluate(PRESENTATION_VISUAL);sourceZoom.filters=visual.filters;sourceZoom.frame_visual=visual;if(!visual.filters.length||visual.paper!=='rgb(255, 255, 255)')throw Error('Fast source visual proof missing');}
    if(!Number.isFinite(sourceZoom.width)||sourceZoom.width<=0||Math.abs(sourceZoom.width/original.width-1.1)>.02||sourceZoom.body!==original.body||!sourceZoom.event?.trusted||!sourceZoom.event?.prevented||sourceZoom.md!=='100%'||sourceZoom.filters.some(f=>f!=='none'))throw Error('Original Ctrl-wheel changed the app or source colors: '+JSON.stringify(sourceZoom));
    officeCheckpoint('source-zoom-verified',{source_zoom:sourceZoom});
    let lastSlide=null;
    if(fastPresentation){const began=Date.now();await evaluate(`document.querySelector('#page-input').value='${slideCount}';document.querySelector('#page-input').dispatchEvent(new Event('change',{bubbles:true}));true`);await until(()=>evaluate(`JSON.parse(document.querySelector('#presentation-frame').dataset.metrics||'{}').page===${slideCount}`),'on-demand final slide with decoded image');lastSlide={ms:Date.now()-began,metrics:await evaluate("JSON.parse(document.querySelector('#presentation-frame').dataset.metrics)")};if(!Number.isInteger(lastSlide.metrics.resources?.images)||lastSlide.metrics.resources.images<1)throw Error('Last slide JPEG was not decoded');lastSlide.visual=await frameEvaluate(PRESENTATION_VISUAL);if(lastSlide.visual.page!==slideCount||!lastSlide.visual.visible_pages.includes(slideCount)||!lastSlide.visual.images_decoded||lastSlide.visual.slot_count!==slideCount)throw Error('Last slide visible/image/continuous-slot proof failed: '+JSON.stringify(lastSlide.visual));}
    let sourceSearch=null,keyboard=null;
    if(fastPresentation){
      const query=cfg.pdfMarker+' slide '+slideCount;
      await frameEvaluate("(()=>{window.__officeKeys=[];document.addEventListener('keydown',event=>{if(event.key==='Escape'||event.key.toLowerCase()==='f')window.__officeKeys.push({key:event.key,ctrl:event.ctrlKey,trusted:event.isTrusted,prevented:event.defaultPrevented});});return{armed:true};})()");
      await evaluate(`document.querySelector('#search').value=${JSON.stringify(query)};true`);await focusPresentation();await trustedKey('f',true);
      await until(()=>evaluate("document.activeElement?.id==='search'"),'opaque iframe Ctrl+F focuses app search');
      const find=await evaluate("({focused:document.activeElement.id==='search',query:document.querySelector('#search').value,start:document.querySelector('#search').selectionStart,end:document.querySelector('#search').selectionEnd})");
      find.event=(await frameEvaluate("({event:window.__officeKeys.find(event=>event.key.toLowerCase()==='f'&&event.ctrl)})")).event;
      if(!find.focused||find.query!==query||find.start!==0||find.end!==query.length||!find.event?.trusted||!find.event?.prevented)throw Error('Native iframe Ctrl+F did not select the existing query: '+JSON.stringify(find));
      keyboard={find};officeCheckpoint('keyboard-find-verified',{keyboard});
      await evaluate("document.querySelector('#search-form').requestSubmit();true");await until(()=>evaluate("document.querySelectorAll('#search-results [data-search-hit]').length===1"),'unique final slide search result');sourceSearch=await until(async()=>{const current=await frameEvaluate(PRESENTATION_VISUAL);const note=await evaluate("document.querySelector('#location-note').textContent");if(note)throw Error('Fast source search showed an error: '+note);return current.page===slideCount&&current.highlighted_page===slideCount&&current.highlights===1&&current.visible_pages.includes(slideCount)&&current.highlight_intersects_viewport?{...current,note}:null;},'official source search highlight after results height changes');
      await focusPresentation();await trustedKey('Escape');await until(()=>evaluate("document.querySelector('#search-results').hidden"),'native iframe Escape dismisses search results');
      await until(async()=>!(await frameEvaluate("({highlights:document.querySelectorAll('[data-pptx-search-highlight]').length})")).highlights,'native iframe Escape clears official search highlights');
      keyboard.search_escape=await evaluate("({reader_visible:!document.querySelector('#reader').hidden,results_hidden:document.querySelector('#search-results').hidden,query:document.querySelector('#search').value,note:document.querySelector('#location-note').textContent})");
      Object.assign(keyboard.search_escape,await frameEvaluate("({highlights:document.querySelectorAll('[data-pptx-search-highlight]').length,event:window.__officeKeys.find(event=>event.key==='Escape')})"));
      if(!keyboard.search_escape.reader_visible||!keyboard.search_escape.results_hidden||keyboard.search_escape.query!==query||keyboard.search_escape.note||keyboard.search_escape.highlights||!keyboard.search_escape.event?.trusted||!keyboard.search_escape.event?.prevented)throw Error('Native iframe Escape did not dismiss only search: '+JSON.stringify(keyboard));
      officeCheckpoint('keyboard-search-escape-verified',{keyboard});
    }


    await evaluate("document.querySelector('#source-zoom-reset').click();document.querySelector('#page-input').value='2';document.querySelector('#page-input').dispatchEvent(new Event('change',{bubbles:true}));true");
    await until(()=>evaluate("document.querySelector('#page-input').value==='2'&&document.querySelector('#source-zoom-reset').textContent==='100%'&&(!document.querySelector('#presentation-frame')||JSON.parse(document.querySelector('#presentation-frame').dataset.metrics||'{}').page===2)"),'PPTX page navigation/reset');
    await evaluate("window.__officeWorkers=[];const on=window.pnx.on;window.pnx.on=function(m){if(m.kind==='worker_started')window.__officeWorkers.push({pid:m.pid,path:m.path});return on(m);};document.querySelector('#reader-extract').click();true");
    const afterClick=await evaluate("({disabled:document.querySelector('#reader-extract').disabled,text:document.querySelector('#reader-extract').textContent})");
    if(!afterClick.disabled)throw Error('Current-document extraction did not prevent an immediate duplicate');
    if(fastPresentation){
      await until(()=>evaluate("!document.querySelector('#stop').disabled&&window.__officeWorkers.some(worker=>worker.path.endsWith('01-slides.pptx'))"),'actual current-document worker starts before Escape');
      const before=await evaluate("({stop_enabled:!document.querySelector('#stop').disabled,reader_visible:!document.querySelector('#reader').hidden})");
      await focusPresentation();await trustedKey('Escape');await until(()=>evaluate("!document.querySelector('#list-view').hidden"),'second native iframe Escape returns to the file list');
      keyboard.close_while_running={before,...await evaluate("({list_visible:!document.querySelector('#list-view').hidden,reader_hidden:document.querySelector('#reader').hidden,frame_removed:!document.querySelector('#presentation-frame'),stop_enabled:!document.querySelector('#stop').disabled})")};
      if(!before.stop_enabled||!before.reader_visible||!keyboard.close_while_running.list_visible||!keyboard.close_while_running.reader_hidden||!keyboard.close_while_running.frame_removed||!keyboard.close_while_running.stop_enabled)throw Error('Native iframe Escape stopped extraction or failed to return to the list: '+JSON.stringify(keyboard));
      officeCheckpoint('keyboard-close-while-running-verified',{keyboard});
    }
    const pptxSaved=await until(async()=>{const value=await evaluate('window.pywebview.api.state()');if(value.files[0].error)throw Error(value.files[0].error);return value.files[0].md_ready?value:null;},'PPTX current-document extraction saves actual Markdown');
    if(pptxSaved.files[1].md_ready)throw Error('Current-document extraction used list selection');
    if(keyboard)keyboard.extraction_completed_after_close=true;
    await until(()=>evaluate("document.querySelector('#stop').disabled"),'office extraction coordinator finishes');
    officeCheckpoint('pptx-saved',{after_click:afterClick,pptx:pptxSaved.files[0]});
    const csvStarted=await evaluate('window.pywebview.api.extract_selected([1])');if(!csvStarted?.ok||csvStarted.n!==1)throw Error('Private CSV extraction failed');
    const finished=await until(async()=>{const value=await evaluate('window.pywebview.api.state()');if(value.files.some(f=>f.error))throw Error(JSON.stringify(value.files));return value.files.every(f=>f.md_ready)?value:null;},'private PPTX and CSV exports');
    await evaluate("document.querySelector('#reader-back').click();true");
    const requestStart=officeRequests.length;await evaluate("document.querySelector('.file-row[data-index=\"0\"] [data-action=read]').click();true");
    await until(()=>evaluate(`document.querySelector('#document').innerText.includes(${JSON.stringify(cfg.pdfMarker)})`),'result-only PPTX Markdown reader');
    const resultRequests=officeRequests.slice(requestStart).filter(value=>/\/open\?|\/document-map\?/.test(value));
    if(!resultRequests.some(value=>new URL(value).searchParams.get('mode')==='result')||resultRequests.some(value=>new URL(value).pathname==='/document-map'||new URL(value).pathname==='/open'&&new URL(value).searchParams.get('mode')!=='result'))throw Error('MD reader eagerly requested original preview: '+JSON.stringify(resultRequests));
    officeCheckpoint('result-only-reader',{result_requests:resultRequests});
    const mdBefore=await evaluate("document.querySelector('#document').innerText");
    const mdTarget=await evaluate("(()=>{const node=document.querySelector('#result-scroll'),box=node.getBoundingClientRect();window.__officeMdWheel=null;node.addEventListener('wheel',event=>window.__officeMdWheel={trusted:event.isTrusted,prevented:event.defaultPrevented},{once:true});return{x:box.left+80,y:box.top+80};})()");
    await command('Input.dispatchMouseEvent',{type:'mouseWheel',x:mdTarget.x,y:mdTarget.y,deltaX:0,deltaY:-120,modifiers:2});
    await until(()=>evaluate("document.querySelector('#md-zoom-reset').textContent==='110%'"),'PPTX Markdown trusted Ctrl-wheel');
    const mdZoom=await evaluate("({event:window.__officeMdWheel,text:document.querySelector('#document').innerText,zoom:Number(getComputedStyle(document.querySelector('#document')).zoom)})");
    if(mdZoom.text!==mdBefore||mdZoom.zoom!==1.1||!mdZoom.event?.trusted||!mdZoom.event?.prevented)throw Error('PPTX Markdown zoom changed content');
    await evaluate("document.querySelector('#code-tab').click();true");
    const code=await evaluate("({text:document.querySelector('#md-source').textContent,zoom:Number(getComputedStyle(document.querySelector('#md-source')).zoom),editable:document.querySelector('#md-source').isContentEditable})");
    if(code.zoom!==1.1||code.editable||!code.text.includes(cfg.pdfMarker))throw Error('PPTX raw Markdown zoom/read-only failed');
    await evaluate("document.querySelector('#compare-tab').click();true");
    await until(()=>evaluate(`document.querySelector('#reader').classList.contains('compare')&&(document.querySelector('#presentation-frame')?.dataset.ready==='true'||document.querySelectorAll('#source-scroll .pdf-page').length===${slideCount})`),'PPTX compare preview');
    const compare=await evaluate("({source:document.querySelector('#source-zoom-reset').textContent,md:document.querySelector('#md-zoom-reset').textContent,editable:document.querySelectorAll('#reader [contenteditable=true],#reader textarea').length})");
    if(compare.source!=='100%'||compare.md!=='110%'||compare.editable)throw Error('Compare pane zoom labels are not independent');
    let nativeMapping=null;
    if(fastPresentation){const clicked=await evaluate(`(()=>{const node=[...document.querySelectorAll('#document [data-block-id]')].find(n=>n.textContent.trim()===${JSON.stringify(cfg.pdfMarker+' slide 1')});if(!node)return false;node.click();return true;})()`);if(!clicked)throw Error('Saved native PPTX paragraph missing from comparison DOM');nativeMapping=await until(async()=>{const state=await frameEvaluate(PRESENTATION_VISUAL);const note=await evaluate("document.querySelector('#location-note').textContent");if(note)throw Error('Native PPTX paragraph mapping failed: '+note);return state.page===1&&state.highlighted_page===1&&state.highlights===1&&state.visible_pages.includes(1)&&state.highlight_intersects_viewport?{...state,note}:null;},'native canonical paragraph to official presentation search bounds');}

    await evaluate("document.querySelector('#theme-toggle').click();true");await until(()=>evaluate("document.documentElement.dataset.theme==='dark'"),'office dark theme');
    if(fastPresentation)await until(()=>evaluate("JSON.parse(document.querySelector('#presentation-frame')?.dataset.theme||'{}').theme==='dark'"),'sandbox presentation dark theme');
    const theme=await evaluate("(()=>{const frame=document.querySelector('#presentation-frame');if(frame)return JSON.parse(frame.dataset.theme||'{}');return{paper:getComputedStyle(document.querySelector('.pdf-page')).backgroundColor,filters:Array.from(document.querySelectorAll('#source-scroll img')).map(i=>getComputedStyle(i).filter)};})()");
    if(fastPresentation)Object.assign(theme,await frameEvaluate(PRESENTATION_VISUAL));
    if(theme.paper!=='rgb(255, 255, 255)'||theme.filters.some(f=>f!=='none'))throw Error('Office dark theme changed source colors');
    theme.ok=true;await optionalScreenshot(cfg.compareShot);await evaluate("document.querySelector('#theme-toggle').click();true");
    let pdfFallback=null;
    if(fastPresentation){await evaluate("document.querySelector('#source-pdf-fallback').click();true");await until(()=>evaluate(`document.querySelectorAll('#source-scroll .pdf-page').length===${slideCount}&&!document.querySelector('#presentation-frame')`),'explicit standard PDF fallback pages');await until(()=>evaluate("(()=>{const image=document.querySelector('#source-scroll img');if(!image)return false;image.loading='eager';return image.complete&&image.naturalWidth>0&&!image.hidden;})()"),'actual fallback PNG decode');pdfFallback=await evaluate("({pages:document.querySelectorAll('.pdf-page').length,images:document.querySelectorAll('.pdf-page img').length,frame:!!document.querySelector('#presentation-frame'),note:document.querySelector('#page-note').textContent,md:document.querySelector('#md-zoom-reset').textContent})");pdfFallback.body_unchanged=(await evaluate("document.querySelector('#document').innerText"))===mdBefore;if(!pdfFallback.body_unchanged||pdfFallback.note||pdfFallback.frame||pdfFallback.md!=='110%')throw Error('PDF fallback replaced Markdown or left an error: '+JSON.stringify(pdfFallback));}
    await evaluate("document.querySelector('#reader-back').click();true");
    const back=await evaluate("({list:!document.querySelector('#list-view').hidden,globalVisible:!document.querySelector('#extract').hidden})");
    if(!back.list||!back.globalVisible)throw Error('Explicit list navigation failed');
    await optionalScreenshot(cfg.listShot);
    const officeWorkers=await evaluate('window.__officeWorkers');if(officeWorkers.length!==2||new Set(officeWorkers.map(w=>w.pid)).size!==2||cfg.names.some(name=>!officeWorkers.some(w=>path.basename(w.path)===name)))throw Error('Office/CSV must use separate owned document workers: '+JSON.stringify(officeWorkers));
    completedResult={ok:true,bridge:true,office:{ok:true,keyboard,frame_security:frameSecurity,last_slide:lastSlide,source_search:sourceSearch,native_mapping:nativeMapping,pdf_fallback:pdfFallback,fast_presentation:fastPresentation,slide_count:slideCount,preparing,original,preview_image_decoded_ms:previewDecodedMs,preview_target_ms:2000,preview_decode_within_target:previewDecodedMs<=2000,hidden_decode_is_not_visible_paint:true,source_zoom:sourceZoom,after_click:afterClick,result_requests:resultRequests,md_zoom:mdZoom,raw_markdown:code,compare,back,workers:officeWorkers},state:finished,settings,settings_ui:settingsUi,theme,screenshots,focus_emulation:focusEmulation};
    fs.writeFileSync(cfg.result,JSON.stringify(completedResult,null,2),'utf8');
  }else{
  await evaluate("document.querySelector('.file-row[data-index=\"0\"] .filename').dispatchEvent(new MouseEvent('dblclick',{bubbles:true}));true");
  await until(()=>evaluate("document.querySelector('#reader').classList.contains('source-only')&&document.querySelectorAll('#source-scroll img').length===1"),'original double-click before Markdown exists');
  await until(()=>evaluate("Array.from(document.querySelectorAll('#source-scroll img')).some(i=>{i.loading='eager';return !i.hidden&&i.complete&&i.naturalWidth>0;})"),'original PNG before extraction');
  const original=await evaluate("({source_visible:!document.querySelector('#source-pane').hidden,result_hidden:document.querySelector('#result-pane').hidden,md_disabled:document.querySelector('#content-tab').disabled,editable:document.querySelectorAll('#reader textarea,#reader [contenteditable=true]').length,error_note:document.querySelector('#page-note').textContent})");
  if(!original.source_visible||!original.result_hidden||!original.md_disabled||original.editable||original.error_note)throw Error('Original-only view before export is invalid: '+JSON.stringify(original));
  const originalSearch=await evaluate(`window.pywebview.api.search_source(0,${JSON.stringify(cfg.pdfMarker)})`);
  if(originalSearch?.ok===false||!Array.isArray(originalSearch.results)||(!cfg.scan_mode&&!originalSearch.results.some(hit=>hit.side==='source'&&hit.locations?.length)))throw Error('Original search before export failed: '+JSON.stringify(originalSearch));
  original.search_hits=originalSearch.results.length;original.hidden_test_eager_load=true;
  if((await evaluate('window.pywebview.api.state()')).files.some(f=>f.md_ready))throw Error('Original view/search unexpectedly created saved Markdown');
  const theme={ok:false,initial:await evaluate("({theme:document.documentElement.dataset.theme,background:getComputedStyle(document.body).backgroundColor})")};
  await evaluate("document.querySelector('#theme-toggle').click();true");
  await until(()=>evaluate("document.documentElement.dataset.theme==='dark'&&!document.querySelector('#theme-toggle').disabled"),'native dark theme switch');
  theme.dark=await evaluate("({theme:document.documentElement.dataset.theme,background:getComputedStyle(document.body).backgroundColor,color:getComputedStyle(document.body).color,paper:getComputedStyle(document.querySelector('.pdf-page')).backgroundColor,imageFilters:Array.from(document.querySelectorAll('#source-scroll img')).map(i=>getComputedStyle(i).filter)})");
  theme.native_dark=await evaluate('window.pywebview.api.settings()');
  if(theme.native_dark.theme!=='dark'||theme.dark.background===theme.initial.background||theme.dark.paper!=='rgb(255, 255, 255)'||!theme.dark.imageFilters.length||theme.dark.imageFilters.some(v=>v!=='none'))throw Error('Dark theme changed original pixels or failed to persist natively');
  await command('Page.reload');
  await until(()=>evaluate("document.documentElement.dataset.theme==='dark'&&!!window.pywebview?.api?.settings&&document.querySelectorAll('.file-row').length>0"),'saved dark theme restored after a real WebView reload');
  if(focusEmulation.requested&&!((await focusState('after-theme-reload')).valid))throw Error('Focus emulation was lost during the theme reload');
  theme.restored=await evaluate('window.pywebview.api.settings()');
  if(theme.restored.theme!=='dark')throw Error('Native theme did not survive the UI reload');
  await evaluate("document.querySelector('#theme-toggle').click();true");
  await until(()=>evaluate("document.documentElement.dataset.theme==='light'&&!document.querySelector('#theme-toggle').disabled"),'light theme restored');
  theme.ok=true;
  await evaluate("document.querySelector('#reader-back').click();true");
  await evaluate("window.__documentWorkers=[];const initialOn=window.pnx.on;window.pnx.on=function(m){if(m.kind==='worker_started')window.__documentWorkers.push({pid:m.pid,path:m.path});return initialOn(m);};true");
  const started=await evaluate('window.pywebview.api.extract_selected([0,1])');
  if(!started?.ok||started.n!==2)throw Error('Native extraction did not start: '+JSON.stringify(started));
  const finished=await until(async()=>{const s=await evaluate('window.pywebview.api.state()');const failed=s.files?.find(f=>f.error||f.status==='failed');if(failed)throw Error('Native extraction failed: '+JSON.stringify(failed));return s.files?.length===cfg.names.length&&s.files.slice(0,2).every(f=>f.done)?s:null;},'both portable Markdown files saved');
  if(!finished.files.slice(0,2).every(f=>f.md_ready===true))throw Error('Saved Markdown readiness was not verified by the backend');
  const documentWorkers=await evaluate('window.__documentWorkers');
  if(documentWorkers.length!==2||new Set(documentWorkers.map(item=>item.pid)).size!==2)throw Error('Each document did not receive a separate extraction process: '+JSON.stringify(documentWorkers));
  if(cfg.background_book&&finished.files[2].md_ready)throw Error('The unprocessed book must not have a Markdown result');
  if(cfg.scan_mode&&(finished.files[0].status!=='partial'||finished.files[1].status!=='complete'||!finished.files[0].warnings.some(w=>/OCR|오인식|인식/.test(String(w)))))throw Error('Scan must preserve its partial status and OCR warning');
  const labels=['저장 완료','저장 완료'];
  await until(()=>evaluate(`(()=>{const nodes=Array.from(document.querySelectorAll('.file-row .status'));const expected=${JSON.stringify(labels)};return nodes.length===${cfg.names.length}&&nodes.slice(0,2).every((node,i)=>node.textContent.includes(expected[i]));})()`),'rendered saved rows');
  await optionalScreenshot(cfg.listShot);
  await evaluate("document.querySelector('[data-action=read][data-index=\"1\"]').click();true");
  await until(()=>evaluate(`document.querySelector('#document')?.innerText.includes(${JSON.stringify(cfg.csvMarker)})&&document.querySelectorAll('#document table tr').length>=2`),'CSV reader table');
  const csvUi=await evaluate("({tables:document.querySelectorAll('#document table').length,editable:document.querySelectorAll('#reader [contenteditable=true],#reader textarea').length,readonly:document.querySelector('#reader .readonly')?.textContent,content:document.querySelector('#document').innerText})");
  if(csvUi.editable)throw Error('Result reader exposes editable controls');
  const zoomTarget=await evaluate("(()=>{const node=document.querySelector('#result-scroll');const box=node.getBoundingClientRect();window.__nativeZoomProbe=null;node.addEventListener('wheel',event=>window.__nativeZoomProbe={trusted:event.isTrusted,prevented:event.defaultPrevented},{once:true});return {x:box.left+Math.min(80,box.width/2),y:box.top+Math.min(80,box.height/2),bodyWidth:document.body.getBoundingClientRect().width};})()");
  await command('Input.dispatchMouseEvent',{type:'mouseWheel',x:zoomTarget.x,y:zoomTarget.y,deltaX:0,deltaY:-120,modifiers:2});
  await until(()=>evaluate("document.querySelector('#md-zoom-reset')?.textContent.trim()==='110%'"),'trusted Ctrl+wheel magnifies Markdown to 110 percent');
  const mdZoom=await evaluate("({label:document.querySelector('#md-zoom-reset').textContent.trim(),zoom:Number(getComputedStyle(document.querySelector('#document')).zoom),bodyWidth:document.body.getBoundingClientRect().width,content:document.querySelector('#document').innerText,event:window.__nativeZoomProbe})");
  if(mdZoom.zoom<=1||mdZoom.content!==csvUi.content||mdZoom.bodyWidth!==zoomTarget.bodyWidth||!mdZoom.event?.trusted||!mdZoom.event?.prevented)throw Error('Markdown zoom did not remain local and read-only: '+JSON.stringify(mdZoom));
  await evaluate("document.querySelector('#md-zoom-reset').click();true");
  await until(()=>evaluate("document.querySelector('#md-zoom-reset').textContent.trim()==='100%'&&Number(getComputedStyle(document.querySelector('#document')).zoom)===1"),'Markdown zoom reset');
  mdZoom.reset=true;
  await evaluate("document.querySelector('#code-tab').click();true");
  const code=await evaluate("({tag:document.querySelector('#md-source').tagName,text:document.querySelector('#md-source').textContent,editable:document.querySelector('#md-source').isContentEditable})");
  if(code.tag!=='PRE'||code.editable||!code.text.includes(cfg.csvMarker))throw Error('Raw Markdown read-only view failed');
  await evaluate("document.querySelector('#reader-back').click();document.querySelector('[data-action=read][data-index=\"0\"]').click();true");
  await until(()=>evaluate(`document.querySelector('#document')?.innerText.includes(${JSON.stringify(cfg.pdfMarker)})`),'PDF Markdown reader');
  await evaluate("document.querySelector('#compare-tab').click();true");
  // A deliberately hidden WebView need not activate browser lazy-loading.
  // Eager loading is confined to this test and still decodes the real PNG URL.
  await until(()=>evaluate("Array.from(document.querySelectorAll('#source-scroll img')).some(i=>{i.loading='eager';return !i.hidden&&i.complete&&i.naturalWidth>0;})"),'real PDF source rendering');
  const compare=await evaluate("({visible:!document.querySelector('#source-pane').hidden,images:document.querySelectorAll('#source-scroll img').length,pages:document.querySelector('#page-total').textContent,editable:document.querySelectorAll('#reader [contenteditable=true],#reader textarea').length,error_note:document.querySelector('#page-note').textContent,page_errors:document.querySelectorAll('#source-scroll .page-error').length})");
  if(!compare.visible||compare.editable||compare.error_note||compare.page_errors)throw Error('PDF comparison has an error or is not a read-only view: '+JSON.stringify(compare));
  compare.hidden_test_eager_load=true;
  await optionalScreenshot(cfg.compareShot);
  await evaluate("document.querySelector('#reader-close').click();true");
  await until(()=>evaluate("document.querySelector('#reader').hidden&&!document.querySelector('#list-view').hidden"),'viewer close X returns to the file list');
  const minimalChrome=await evaluate("({close_returns_list:document.querySelector('#reader').hidden,savebar_visible:!!document.querySelector('.savebar')?.getClientRects().length,footer_visible:!!document.querySelector('.footer')?.getClientRects().length})");
  if(minimalChrome.savebar_visible||minimalChrome.footer_visible)throw Error('Redundant persistent status/path chrome remains visible');
  await evaluate("document.querySelector('[data-action=read][data-index=\"0\"]').click();true");
  await until(()=>evaluate(`document.querySelector('#document')?.innerText.includes(${JSON.stringify(cfg.pdfMarker)})`),'saved Markdown reopens after X');
  if(cfg.background_book){
    await evaluate("window.__backgroundWorkers=[];const on=window.pnx.on;window.pnx.on=function(m){if(m.kind==='worker_started')window.__backgroundWorkers.push(m.pid);return on(m);};true");
    const start=await evaluate('window.pywebview.api.extract_selected([2],true)');
    if(!start?.ok)throw Error('Book background extraction did not start');
    const running=await until(async()=>{const s=await evaluate('window.pywebview.api.state()');const f=s.files?.[2];if(f?.error)throw Error(f.error);return f?.status==='running'&&f.progress_pct>=5&&f.progress_pages>1?s:null;},'actual multi-page book SDK progress');
    const bookBegan=Date.now();
    fs.writeFileSync(cfg.bookActive,JSON.stringify({started_at:new Date().toISOString(),book:running.files[2],resources:running.resources}),'utf8');
    const measurements=[];
    background={ok:false,measurements,latency_failures:[],stage:'measuring-book',
      latency_gate:{limit_ms:2000,targets:['bridge_inside_browser','node_http_files','browser_http_files'],cdp_roundtrip:'diagnostic_only'}};
    for(let index=0;index<4;index++){
      const pingBegan=Date.now();
      const browser=await focusState('book-round-'+index);
      const pingMs=Date.now()-pingBegan;
      const began=Date.now();
      const timed=await evaluate('(async()=>{const began=performance.now();const state=await window.pywebview.api.state();return {state,bridge_ms:performance.now()-began};})()');
      const s=timed.state;
      const elapsed=Date.now()-began;
      const httpBegan=Date.now();
      const backendResponse=await fetch(new URL('files',target.url),{cache:'no-store',signal:AbortSignal.timeout(15000)});
      const backend=await backendResponse.json();
      const backendMs=Date.now()-httpBegan;
      const browserHttp=await evaluate('(async()=>{const began=performance.now();const response=await fetch("/files",{cache:"no-store"});const state=await response.json();return {http_ms:performance.now()-began,status:response.status,files:state.files};})()');
      const failures=[];
      if(timed.bridge_ms>2000)failures.push('state_api_over_2000ms');
      if(backendMs>2000)failures.push('http_backend_over_2000ms');
      if(browserHttp.http_ms>2000)failures.push('browser_http_over_2000ms');
      if(focusEmulation.requested&&!browser.valid)failures.push('focus_emulation_state_invalid');
      if(s.files[2].status!=='running')failures.push('book_not_running');
      if(!s.files[0].md_ready)failures.push('saved_markdown_unavailable');
      measurements.push({roundtrip_ms:elapsed,bridge_ms:timed.bridge_ms,cdp_ping_ms:pingMs,
        node_http_files_ms:backendMs,browser_http_files_ms:browserHttp.http_ms,
        http_status:backendResponse.status,browser_http_status:browserHttp.status,
        has_focus:browser.has_focus,visibility:browser.visibility,hidden:browser.hidden,failures,
        percent:s.files[2].progress_pct,stage:s.files[2].progress_stage,page:s.files[2].progress_page,
        book:s.files[2],saved:s.files[0],resources:s.resources,backend_files:backend.files});
      background.latency_failures.push(...failures.map(reason=>({round:index,reason})));
      probe({stage:'book-latency',measurements,latency_failures:background.latency_failures});
      fs.writeFileSync(cfg.result,JSON.stringify({ok:false,background,focus_emulation:focusEmulation}),'utf8');
      await pause(1000);
    }
    // This private diagnostic job stays alive for a 15-second own-PID profiler.
    // The latency gate remains 2000ms and is never relaxed.
    await pause(Math.max(0,20000-(Date.now()-bookBegan)));
    const searchFocus=await focusState('source-search');
    if(focusEmulation.requested&&!searchFocus.valid)background.latency_failures.push({reason:'focus_emulation_state_invalid_before_search'});
    const begun=Date.now();
    const searched=await evaluate(`(async()=>{const began=performance.now();const hits=await window.pywebview.api.search_source(0,${JSON.stringify(cfg.pdfMarker)});return {hits,bridge_ms:performance.now()-began};})()`);
    const hits=searched.hits;
    const searchMs=Date.now()-begun;
    background.source_search={roundtrip_ms:searchMs,bridge_ms:searched.bridge_ms,ok:hits?.ok,hits:hits?.results?.length,error:hits?.why};
    if(searched.bridge_ms>2000)background.latency_failures.push({reason:'source_search_api_over_2000ms'});
    fs.writeFileSync(cfg.result,JSON.stringify({ok:false,background,focus_emulation:focusEmulation}),'utf8');
    if(!hits?.ok||!hits.results?.length)throw Error('Original search failed during actual book extraction: '+JSON.stringify(background.source_search));
    const controls=await evaluate("({add_files:!document.querySelector('[data-action=add-files]').disabled,add_folder:!document.querySelector('[data-action=add-folder]').disabled,saved_md:!document.querySelector('#content-tab').disabled,reader_visible:!document.querySelector('#reader').hidden})");
    if(!Object.values(controls).every(Boolean))throw Error('Reading/file controls were disabled during extraction: '+JSON.stringify(controls));
    await evaluate(`document.querySelector('#search').value=${JSON.stringify(cfg.pdfMarker)};document.querySelector('#search-form').dispatchEvent(new Event('submit',{bubbles:true,cancelable:true}));true`);
    await until(()=>evaluate("document.querySelectorAll('.search-hit').length>0"),'MD/source search through actual busy reader');
    const queryBefore=await evaluate("document.querySelector('#search').value");
    const enqueued=await evaluate('window.pywebview.api.extract_selected([1],true)');
    if(!enqueued?.ok||!enqueued.queued||enqueued.n!==1)throw Error('Cannot enqueue another document during book extraction: '+JSON.stringify(enqueued));
    const duplicate=await evaluate('window.pywebview.api.extract_selected([2],true)');
    if(!duplicate?.ok||duplicate.n!==0)throw Error('Queued/running document was duplicated: '+JSON.stringify(duplicate));
    const queueState=await evaluate('window.pywebview.api.state()');
    if(!['queued','running','complete'].includes(queueState.files[1].status)||!queueState.files[1].md_ready)throw Error('Another job made the existing Markdown unreadable');
    const parallelFinished=await until(async()=>{const s=await evaluate('window.pywebview.api.state()');if(s.files[1].error)throw Error(s.files[1].error);return s.files[1].done&&s.files[2].status==='running'?s:null;},'CSV completes in its own process while the book keeps extracting');
    const cancelFocus=await focusState('cancel');
    if(focusEmulation.requested&&!cancelFocus.valid)background.latency_failures.push({reason:'focus_emulation_state_invalid_before_cancel'});
    const cancelBegan=Date.now();
    const cancelled=await evaluate('window.pywebview.api.cancel()');
    if(!cancelled?.ok)throw Error('Cancel was not accepted');
    const stopped=await until(async()=>{const s=await evaluate('window.pywebview.api.state()');return s.files?.[2]?.status==='paused'?s:null;},'owned background process cancelled');
    const cancelMs=Date.now()-cancelBegan;
    if(cancelMs>4500)throw Error('Cancellation did not stop the native child promptly: '+cancelMs);
    const queryAfter=await evaluate("document.querySelector('#search').value");
    if(queryBefore!==queryAfter||!stopped.files[0].md_ready||stopped.files[2].md_ready)throw Error('Cancellation reset the reader or promoted an unfinished book');
    const workers=await evaluate('window.__backgroundWorkers');
    if(!workers.length)throw Error('No actual extraction child PID was observed');
    if(!stopped.files[1].done||!stopped.files[1].md_ready||workers.length!==2||new Set(workers).size!==2)throw Error('The independent completed CSV or two worker PIDs were not preserved');
    background={...background,ok:background.latency_failures.length===0,stage:'cancelled-and-verified',measurements,source_search_ms:searchMs,controls,cancel_ms:cancelMs,worker_pids:workers,reader_query_preserved:true,enqueue:enqueued,duplicate,parallel_completed:parallelFinished.files[1],resources:parallelFinished.resources,book_final:stopped.files[2]};
  }
  completedResult={ok:background?.ok!==false,target_url:target.url,bridge:true,settings,settings_ui:settingsUi,manual_update:manualUpdate,initial_md_buttons:initialMdButtons,original_before_export:original,start_response:started,document_workers:documentWorkers,state:finished,csv_ui:csvUi,theme,md_zoom:mdZoom,minimal_chrome:minimalChrome,raw_markdown:code,pdf_comparison:compare,background,screenshots,focus_emulation:focusEmulation};
  fs.writeFileSync(cfg.result,JSON.stringify(completedResult,null,2),'utf8');
  if(!completedResult.ok)throw Error('Book latency failed the unchanged 2000ms gate: '+JSON.stringify(background.latency_failures));
  }
} catch(error){
  fs.writeFileSync(cfg.result,JSON.stringify({...completedResult,ok:false,error:String(error),background,focus_emulation:focusEmulation},null,2),'utf8');
  throw error;
} finally {
  try{
    if(focusEmulation.enabled){
      try{focusEmulation.reset_protocol_result=await command('Emulation.setFocusEmulationEnabled',{enabled:false});focusEmulation.reset=true;}
      catch(error){focusEmulation.reset=false;focusEmulation.reset_error=String(error);}
      const saved=fs.existsSync(cfg.result)?JSON.parse(fs.readFileSync(cfg.result,'utf8')):{ok:false};
      fs.writeFileSync(cfg.result,JSON.stringify({...saved,ok:saved.ok===true&&focusEmulation.reset,focus_emulation:focusEmulation},null,2),'utf8');
      if(!focusEmulation.reset)throw Error('Focus emulation cleanup failed: '+focusEmulation.reset_error);
    }
  } finally {if(ws)ws.close();}
}
'''
_CDP_SCRIPT = _CDP_SCRIPT.replace('__PNX_FRAME_VISUAL__', json.dumps(_PRESENTATION_VISUAL_EXPRESSION))


class _PrivateJob:
    """KILL_ON_JOB_CLOSE applies only to this newly spawned EXE and descendants."""
    def __init__(self, process):
        class Basic(ctypes.Structure):
            _fields_ = [("process_time", ctypes.c_longlong), ("job_time", ctypes.c_longlong),
                        ("flags", wintypes.DWORD), ("min_working", ctypes.c_size_t),
                        ("max_working", ctypes.c_size_t), ("active", wintypes.DWORD),
                        ("affinity", ctypes.c_size_t), ("priority", wintypes.DWORD),
                        ("scheduling", wintypes.DWORD)]
        class Extended(ctypes.Structure):
            _fields_ = [("basic", Basic), ("io", ctypes.c_ulonglong * 6),
                        ("process_memory", ctypes.c_size_t), ("job_memory", ctypes.c_size_t),
                        ("peak_process_memory", ctypes.c_size_t), ("peak_job_memory", ctypes.c_size_t)]
        self.kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        self.kernel.CreateJobObjectW.argtypes = [ctypes.c_void_p, wintypes.LPCWSTR]
        self.kernel.CreateJobObjectW.restype = wintypes.HANDLE
        self.kernel.SetInformationJobObject.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p, wintypes.DWORD]
        self.kernel.AssignProcessToJobObject.argtypes = [wintypes.HANDLE, wintypes.HANDLE]
        self.kernel.QueryInformationJobObject.argtypes = [wintypes.HANDLE, ctypes.c_int, ctypes.c_void_p,
                                                        wintypes.DWORD, ctypes.c_void_p]
        self.kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        self.handle = self.kernel.CreateJobObjectW(None, None)
        if not self.handle:
            raise ctypes.WinError(ctypes.get_last_error())
        limits = Extended()
        limits.basic.flags = 0x2000  # JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
        try:
            if not self.kernel.SetInformationJobObject(self.handle, 9, ctypes.byref(limits), ctypes.sizeof(limits)):
                raise ctypes.WinError(ctypes.get_last_error())
            if not self.kernel.AssignProcessToJobObject(self.handle, wintypes.HANDLE(process._handle)):
                raise ctypes.WinError(ctypes.get_last_error())
        except Exception:
            self.close()
            raise

    def close(self):
        if self.handle:
            self.kernel.CloseHandle(self.handle)
            self.handle = None

    def pids(self):
        buffer = ctypes.create_string_buffer(65536)
        if not self.kernel.QueryInformationJobObject(self.handle, 3, buffer, len(buffer), None):
            return []
        count = wintypes.DWORD.from_buffer(buffer, 4).value
        if count > (len(buffer) - 8) // ctypes.sizeof(ctypes.c_size_t):
            return []
        return list((ctypes.c_size_t * count).from_buffer(buffer, 8))


def _cache_root() -> Path:
    settings = Path.home() / ".ai-workspace" / "paths.md"
    match = re.search(r"^cache_root:\s*(.+)$", settings.read_text(encoding="utf-8-sig"), re.M)
    if not match:
        raise RuntimeError("cache_root is missing from local paths.md")
    return Path(match.group(1).strip()) / "260725 PDF 추출"


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _runtime_assets(executable: Path) -> dict:
    web = executable.resolve().parent / "_internal" / "web"
    names = ("presentation-frame.html", "presentation-frame.mjs", "pptx-renderer/renderer.mjs", "pptx-renderer/manifest.json")
    return {name: {"bytes": (web / name).stat().st_size, "sha256": _digest(web / name)} for name in names}


def _office_fixtures(root: Path, pages: int = 3, require_fast: bool = False) -> dict:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    inputs, cache = root / "inputs", root / "cache"
    inputs.mkdir()
    cache.mkdir()
    pptx, csv = inputs / "01-slides.pptx", inputs / "02-table.csv"
    marker, csv_marker = "FROZEN_PPTX_NATIVE_9624", "FROZEN_CSV_CELL_8426"
    presentation = Presentation()
    import io
    import random
    from PIL import Image
    for index in range(pages):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        text = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1.5)).text_frame
        text.text = marker + " slide " + str(index + 1)
        text.paragraphs[0].font.size = Pt(28)
        image = io.BytesIO()
        Image.frombytes("RGB", (800, 600), random.Random(20261024 + index).randbytes(800 * 600 * 3)).save(image, format="JPEG", quality=65)
        image.seek(0)
        slide.shapes.add_picture(image, Inches(.5), Inches(2.5), width=Inches(5.5), height=Inches(4.1))
        body = slide.shapes.add_textbox(Inches(6.2), Inches(3), Inches(3), Inches(2)).text_frame
        body.text = "Synthetic local presentation body.\nRead-only original and portable Markdown."
        for paragraph in body.paragraphs:
            paragraph.font.size = Pt(20)
    presentation.save(pptx)
    csv.write_text(f"Name,Value\n{csv_marker},42\nSecond,7\n", encoding="utf-8", newline="")
    (cache / "ui.json").write_text(json.dumps({"folder": str(inputs)}), encoding="utf-8")
    (root / "config.yaml").write_text(json.dumps({"project": {"input_dir": str(inputs), "work_dir": str(root / "work")},
        "grobid": {"url": ""}, "identify": {"resolve_missing_doi": False}, "metadata": {"request_delay_sec": 0}}), encoding="utf-8")
    for source in (pptx, csv):
        source.chmod(stat.S_IREAD)
    return {"inputs": str(inputs), "cache": str(cache), "names": [pptx.name, csv.name],
            "output_dir": str(inputs / "Markdown"), "pdfMarker": marker, "csvMarker": csv_marker,
            "office_mode": True, "office_pages": pages, "require_presentation": require_fast, "scan_mode": False, "source_digests_before": {str(path): _digest(path) for path in (pptx, csv)}}


def _fixtures(root: Path, *, scan=False) -> dict:
    import fitz
    from pubnexus import general_pdf, store
    inputs = root / "inputs"
    inputs.mkdir()
    pdf, csv = inputs / "01-native.pdf", inputs / "02-table.csv"
    pdf_marker, csv_marker = "FROZEN_PDF_NATIVE_7391", "FROZEN_CSV_CELL_8426"
    with fitz.open() as document:
        page = document.new_page()
        if scan:
            from PIL import Image, ImageDraw, ImageFont
            image = Image.new("RGB", (1800, 450), "white")
            drawing = ImageDraw.Draw(image)
            font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 76)
            drawing.text((75, 170), "Frozen OCR marker 7391", font=font, fill="black")
            data = io.BytesIO()
            image.save(data, format="PNG")
            page.insert_image(page.rect, stream=data.getvalue())
            pdf_marker = "7391"
        else:
            for index in range(12):
                page.insert_text((55, 60 + index * 25), pdf_marker + f" synthetic original line {index + 1}.")
        document.save(pdf)
    csv.write_text(f"Name,Value\n{csv_marker},42\nSecond,7\n", encoding="utf-8", newline="")
    cache = root / "cache"
    previous = os.environ.get("PUBNEXUS_STORE")
    os.environ["PUBNEXUS_STORE"] = str(cache)
    try:
        sha = store.file_sha1(pdf)
        if scan:
            with fitz.open(pdf) as original:
                if original[0].get_text().strip():
                    raise AssertionError("Scan fixture must not contain a native text layer")
            doc = {"sha1": sha, "source": "pdf_general", "meta": {"title": "Frozen scan PDF"},
                   "body_text": [], "tables": [], "figures": [], "references": []}
        else:
            doc, report = general_pdf.ensure_fulltext({"sha1": sha, "source": "pdf_general",
                                                     "meta": {"title": "Frozen native PDF"}}, pdf)
            if report["status"] == "failed" or pdf_marker not in json.dumps(doc):
                raise AssertionError("Native PDF cache fixture extraction failed")
        store.save(sha, doc, pdf)
        (cache / "ui.json").write_text(json.dumps({"folder": str(inputs)}), encoding="utf-8")
    finally:
        if previous is None:
            os.environ.pop("PUBNEXUS_STORE", None)
        else:
            os.environ["PUBNEXUS_STORE"] = previous
    configuration = {"project": {"input_dir": str(inputs), "work_dir": str(root / "work")},
                     "grobid": {"url": ""}, "identify": {"resolve_missing_doi": False},
                     "metadata": {"request_delay_sec": 0}}
    (root / "config.yaml").write_text(json.dumps(configuration), encoding="utf-8")
    for source in (pdf, csv):
        source.chmod(stat.S_IREAD)
    return {"inputs": str(inputs), "cache": str(cache), "names": [pdf.name, csv.name],
            "output_dir": str(inputs / "Markdown"),
            "pdfMarker": pdf_marker, "csvMarker": csv_marker,
            "scan_mode": scan, "native_text_empty": scan,
            "source_digests_before": {str(p): _digest(p) for p in (pdf, csv)}}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--exe", type=Path)
    parser.add_argument("--out-dir", type=Path)
    parser.add_argument("--node", default=shutil.which("node"))
    parser.add_argument("--timeout", type=float, default=90)
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--office", action="store_true", help="Actual synthetic three-slide PPTX original/zoom/viewer-extraction/lazy MD")
    parser.add_argument("--office-pages", type=int, choices=(3, 71), default=3, help="Synthetic distinct JPEG slides; 71 exercises on-demand rendering")
    parser.add_argument("--office-fast", action="store_true", help="Require sandboxed presentation and first decoded ready within 2000ms")
    parser.add_argument("--scan", action="store_true", help="Use a raster-only PDF and an empty canonical cache")
    parser.add_argument("--check-update", action="store_true", help="Explicitly query the real release manifest; never install")
    parser.add_argument("--foreground-test", action="store_true", help="Diagnostic hidden WebView: disable Chromium background scheduling, never production")
    parser.add_argument("--focus-emulation-test", action="store_true", help="Test-only CDP active/focused page emulation; require actual focus and visibility")
    parser.add_argument("--background-book", type=Path, help="Copy a real PDF into private TEMP and verify browsing/search/cancel during its extraction")
    args = parser.parse_args()
    if not args.office and (args.office_pages != 3 or args.office_fast):
        parser.error("--office-pages/--office-fast require --office")
    if args.office and (args.scan or args.background_book):
        parser.error("--office cannot be combined with --scan or --background-book")
    if sys.platform != "win32":
        parser.error("This harness requires Windows WebView2")
    if args.timeout <= 0 or not args.node:
        parser.error("A positive timeout and Node.js 22+ are required")
    if args.foreground_test and args.focus_emulation_test:
        parser.error("Choose one diagnostic condition at a time")
    if not args.prepare_only and (not args.exe or not args.exe.is_file()):
        parser.error("Pass the freshly built EXE through --exe")
    cache_root = _cache_root().resolve()
    output = (args.out_dir or cache_root / "frozen-smoke" /
              (datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:8])).resolve()
    if not output.is_relative_to(cache_root):
        parser.error("Receipts must stay inside the configured project's DevCache")
    output.mkdir(parents=True, exist_ok=False)
    receipt = {"ok": False, "mode": "prepare-only" if args.prepare_only else "frozen-exe",
               "started_at": datetime.now(timezone.utc).isoformat(), "output_dir": str(output),
               "harness_elevated": bool(ctypes.windll.shell32.IsUserAnAdmin())}
    process = job = root = None
    try:
        probe = subprocess.run([args.node, "--input-type=module", "-e", "console.log(typeof WebSocket)"],
                               capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=10,
                               creationflags=subprocess.CREATE_NO_WINDOW)
        if probe.returncode or (probe.stdout or '').strip() != "function":
            raise RuntimeError("Node.js has no built-in WebSocket; no packages were installed")
        with tempfile.TemporaryDirectory(prefix="pubnexus-frozen-") as temporary:
            root = Path(temporary)
            fixtures = _office_fixtures(root, pages=args.office_pages, require_fast=args.office_fast) if args.office else _fixtures(root, scan=args.scan)
            if args.background_book:
                original = args.background_book.resolve(strict=True)
                before = _digest(original)
                book = Path(fixtures['inputs']) / '03-book.pdf'
                shutil.copy2(original, book)
                if _digest(book) != before:
                    raise AssertionError('Private book copy failed integrity verification')
                book.chmod(stat.S_IREAD)
                fixtures['names'].append(book.name)
                fixtures['source_digests_before'][str(book)] = before
                fixtures['background_book'] = True
                receipt['book_original'] = {'path': str(original), 'sha256_before': before}
            sources = [Path(path) for path in fixtures["source_digests_before"]]
            try:
                receipt.update(fixtures=fixtures, isolated_root=str(root), source_readonly=True)
                from pubnexus.version import VERSION
                config = {**fixtures, "timeout": args.timeout, "result": str(root / "cdp-result.json"),
                          "expected_version": VERSION, "check_update": args.check_update,
                          "focus_emulation_test": args.focus_emulation_test,
                          "listShot": str(output / "list.png"), "compareShot": str(output / "comparison.png"),
                          "bookActive": str(output / "book-running.json"),
                          "probe": str(output / "cdp-probe.json")}
                if args.prepare_only:
                    receipt.update(ok=True, fixtures_validated=True, exe_launched=False)
                else:
                    receipt["runtime_assets_before"] = _runtime_assets(args.exe)
                    source_web = Path(__file__).resolve().parents[1] / "web"
                    receipt["runtime_assets_expected"] = {name: {"bytes": (source_web / name).stat().st_size,
                        "sha256": _digest(source_web / name)} for name in ("presentation-frame.html", "presentation-frame.mjs")}
                    if any(receipt["runtime_assets_before"][name] != expected for name, expected in receipt["runtime_assets_expected"].items()):
                        raise AssertionError("Frozen frame assets do not match the current approved source")
                    with socket.socket() as reserve:
                        reserve.bind(("127.0.0.1", 0))
                        config["port"] = reserve.getsockname()[1]
                    config_path = root / "cdp-config.json"
                    config_path.write_text(json.dumps(config), encoding="utf-8")
                    environment = os.environ.copy()
                    foreground_flags = ['--disable-backgrounding-occluded-windows',
                                        '--disable-renderer-backgrounding', '--disable-background-timer-throttling'] if args.foreground_test else []
                    environment.update(PUBNEXUS_STORE=fixtures["cache"], PUBNEXUS_ROOT=str(root),
                                       PUBNEXUS_OUTPUT=fixtures["output_dir"],
                                       WEBVIEW2_ADDITIONAL_BROWSER_ARGUMENTS=' '.join([f"--remote-debugging-port={config['port']}", *foreground_flags]),
                                       WEBVIEW2_USER_DATA_FOLDER=str(root / "profile"), PYTHONDONTWRITEBYTECODE="1",
                                       PNX_TEST_DEBUG_PORT=str(config["port"]))
                    startup = subprocess.STARTUPINFO()
                    startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
                    startup.wShowWindow = subprocess.SW_HIDE
                    with open(output / "exe.log", "wb") as log:
                        process = subprocess.Popen([str(args.exe.resolve()), "--hidden"], cwd=str(root),
                                                   env=environment, stdout=log, stderr=log, stdin=subprocess.DEVNULL,
                                                   creationflags=subprocess.CREATE_NO_WINDOW, startupinfo=startup)
                        job = _PrivateJob(process)
                        receipt["job_bound"] = bool(job)
                        receipt.update(exe=str(args.exe.resolve()), exe_sha256=_digest(args.exe), pid=process.pid)
                        if foreground_flags:
                            import psutil
                            observed_args = []
                            owner = psutil.Process(process.pid)
                            limit = time.monotonic() + 12
                            while time.monotonic() < limit:
                                if process.poll() is not None or not owner.is_running():
                                    raise RuntimeError('Private frozen app exited before diagnostic flag verification')
                                for child in owner.children(recursive=True):
                                    try:
                                        command_line = child.cmdline()
                                        if all(flag in command_line for flag in foreground_flags):
                                            observed_args.append({'pid': child.pid, 'arguments': command_line,
                                                                  'priority': int(child.nice())})
                                    except psutil.Error:
                                        pass
                                if observed_args:
                                    break
                                time.sleep(.1)
                            receipt['foreground_test'] = {'requested_flags': foreground_flags,
                                                          'observed_owned_processes': observed_args,
                                                          'applied': bool(observed_args), 'production': False}
                            if not observed_args:
                                raise RuntimeError('Diagnostic WebView background flags were not applied; no valid A/B condition')
                        (output / 'running.json').write_text(json.dumps(
                            {'own_pid': process.pid, 'config': config, 'exe_sha256': receipt['exe_sha256']},
                            ensure_ascii=False, indent=2), encoding='utf-8')
                        cdp_script = root / "cdp-smoke.mjs"
                        cdp_script.write_text(_CDP_SCRIPT, encoding="utf-8")
                        cdp = subprocess.run([args.node, str(cdp_script), str(config_path)],
                                             capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=args.timeout + 20,
                                             creationflags=subprocess.CREATE_NO_WINDOW)
                        stdout, stderr = cdp.stdout or '', cdp.stderr or ''
                        (output / "cdp.log").write_text(stdout + stderr, encoding="utf-8")
                        result_path = Path(config['result'])
                        if result_path.is_file():
                            receipt['webview'] = json.loads(result_path.read_text(encoding='utf-8'))
                        if cdp.returncode:
                            raise RuntimeError("Frozen WebView2 bridge smoke failed: " + stderr[-3000:])
                        result = json.loads(Path(config["result"]).read_text(encoding="utf-8"))
                        if not result.get("ok"):
                            raise AssertionError("CDP returned an unsuccessful receipt")
                        markdown = {}
                        expected = {("01-slides.md" if args.office else "01-native.md"): fixtures["pdfMarker"], "02-table.md": fixtures["csvMarker"]}
                        for name, marker in expected.items():
                            saved = Path(fixtures["output_dir"]) / name
                            text = saved.read_text(encoding="utf-8")
                            if marker not in text or (name.endswith("table.md") and "| " + marker + " | 42 |" not in text):
                                raise AssertionError("Saved Markdown body/table mismatch: " + name)
                            markdown[name] = {"sha256": _digest(saved), "bytes": saved.stat().st_size,
                                              "body_checked": True, "table_checked": name.endswith("table.md")}
                        if args.scan:
                            pdf = Path(fixtures["inputs"]) / fixtures["names"][0]
                            sha = hashlib.sha1(pdf.read_bytes()).hexdigest()
                            cached = Path(fixtures["cache"]) / "docs" / sha[:2] / (sha + ".json")
                            doc = json.loads(cached.read_text(encoding="utf-8"))
                            report = doc.get("extraction_report") or {}
                            if doc.get("source") == "pdf_sdk":
                                ocr_pages = {page["page"] for page in report.get("pages", [])
                                             if page.get("ocr_performed") and page.get("ocr_method") == "tesseract"}
                                ocr = [block for block in doc.get("sdk_blocks", [])
                                       if block.get("source", {}).get("page") in ocr_pages and block.get("text", "").strip()]
                            else:
                                ocr = [p for section in doc.get("body_text", []) for p in section.get("paragraphs", [])
                                       if p.get("source", {}).get("method") == "tesseract"]
                            if not report.get("ocr_performed") or not ocr:
                                raise AssertionError("Scan result lacks actual OCR provenance")
                            receipt["ocr"] = {"performed": True, "paragraphs": len(ocr),
                                              "warnings": report.get("warnings", []), "status": report.get("status")}
                        receipt.update(ok=True, exe_launched=True, webview=result, markdown=markdown)
                after = {str(path): _digest(path) for path in sources}
                receipt.update(source_digests_after=after, sources_unchanged=after == fixtures["source_digests_before"])
                if not receipt["sources_unchanged"]:
                    raise AssertionError("A synthetic original changed")
                if args.background_book:
                    checksum = _digest(args.background_book)
                    receipt['book_original']['sha256_after'] = checksum
                    if checksum != receipt['book_original']['sha256_before']:
                        raise AssertionError('Original book changed during background validation')
            finally:
                try:
                    partial_result = root / 'cdp-result.json'
                    if partial_result.is_file() and 'webview' not in receipt:
                        receipt['webview'] = json.loads(partial_result.read_text(encoding='utf-8'))
                    # Failed diagnostics also preserve source integrity evidence.
                    after = {str(path): _digest(path) for path in sources}
                    receipt.update(source_digests_after=after,
                                   sources_unchanged=after == fixtures['source_digests_before'])
                    if not receipt['sources_unchanged']:
                        raise AssertionError('A private fixture original changed')
                    if args.background_book:
                        receipt['book_original']['sha256_after'] = _digest(args.background_book)
                        if receipt['book_original']['sha256_after'] != receipt['book_original']['sha256_before']:
                            raise AssertionError('Original book changed during diagnostics')
                except Exception as error:
                    receipt.update(ok=False, diagnostic_integrity_error=f'{type(error).__name__}: {error}')
                finally:
                    if process:
                        receipt["exe_returncode_before_cleanup"] = process.poll()
                    if job:
                        receipt["owned_job_pids_before_cleanup"] = job.pids()
                        job.close()
                        job = None
                    if process and process.poll() is None:
                        subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"],
                                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                       timeout=10, creationflags=subprocess.CREATE_NO_WINDOW)
                    if process:
                        process.wait(timeout=10)
                    for source in sources:
                        source.chmod(stat.S_IREAD | stat.S_IWRITE)
        receipt["temporary_inputs_removed"] = not root.exists()
        if not receipt["temporary_inputs_removed"]:
            raise AssertionError("Temporary fixture/profile cleanup did not complete")
    except Exception as error:
        if isinstance(error, subprocess.TimeoutExpired):
            def decoded(value):
                return value.decode('utf-8', errors='replace') if isinstance(value, bytes) else (value or '')
            (output / 'cdp.log').write_text(decoded(error.stdout) + decoded(error.stderr), encoding='utf-8')
            receipt.update(ok=False, error=f'CDP diagnostic timeout: {error.timeout}s', cdp_timeout=True)
        else:
            receipt.update(ok=False, error=f"{type(error).__name__}: {error}")
    finally:
        if "runtime_assets_before" in receipt:
            try:
                receipt["runtime_assets_after"] = _runtime_assets(args.exe)
                receipt["runtime_assets_unchanged"] = receipt["runtime_assets_before"] == receipt["runtime_assets_after"]
                if not receipt["runtime_assets_unchanged"]:
                    receipt.update(ok=False, runtime_assets_error="Runtime assets changed during this private diagnostic")
            except Exception as error:
                receipt.update(ok=False, runtime_assets_error=f"{type(error).__name__}: {error}")
        if root is not None:
            receipt["temporary_inputs_removed"] = not root.exists()
        receipt["finished_at"] = datetime.now(timezone.utc).isoformat()
        (output / "receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": receipt["ok"], "receipt": str(output / "receipt.json"),
                      "error": receipt.get("error")}, ensure_ascii=False), flush=True)
    return 0 if receipt["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
