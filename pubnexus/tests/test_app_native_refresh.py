"""Existing image-only PPTX results stay readable while the new adapter upgrades them."""
import copy
from unittest.mock import patch
import unittest

import test_app_export as support
from pubnexus import documents, store


class NativeRefreshTests(unittest.TestCase):
    setUp = support.AppExportTests.setUp

    def presentation(self):
        from pptx import Presentation
        source = self.root / 'native.pptx'
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = 'Native text remains readable.'
        deck.save(source)
        self.app.pdfs = [source]
        old = documents.extract_document(source)
        old['extraction_report'].pop('image_ocr_version', None)
        store.save(old['sha1'], old, source)
        return source, old

    def test_old_pptx_cache_is_reextracted_and_the_ocr_version_is_persisted(self):
        source, old = self.presentation()
        updated = copy.deepcopy(old)
        updated['extraction_report']['image_ocr_version'] = 1
        with patch.object(documents, 'extract_document', return_value=updated) as extract:
            self.app._work([source])
        extract.assert_called_once()
        self.assertTrue(callable(extract.call_args.kwargs.get('progress')))
        self.assertEqual(store.load(old['sha1'])['extraction_report']['image_ocr_version'], 1)
        self.assertEqual(self.app.exported[support.app_module._norm(source)]['image_ocr_version'], 1)
        self.assertFalse(self.app.file_rows()[0]['needs_refresh'])

    def test_previous_pptx_markdown_can_be_read_before_the_ocr_upgrade(self):
        source, old = self.presentation()
        markdown = self.app.output_dir / 'native.md'
        markdown.parent.mkdir(parents=True)
        markdown.write_text('# Previous result\n\nNative text remains readable.\n', encoding='utf-8')
        stat = source.stat()
        self.app.exported[support.app_module._norm(source)] = {
            'markdown': str(markdown), 'source_stat': (stat.st_size, stat.st_mtime_ns),
            'source_sha1': old['sha1'], 'source_type': 'document_pptx', 'status': 'complete'}
        row = self.app.file_rows()[0]
        self.assertTrue(row['needs_refresh'])
        self.assertTrue(row['md_ready'], 'A saved MD must remain readable')
        self.assertFalse(row['done'], 'The extraction action must offer the OCR upgrade')
        self.assertIn('Previous result', self.app.result_view(0)['doc']['md_text'])

    def test_native_collection_does_not_claim_ocr_completion(self):
        source, document = self.presentation()
        document['extraction_report']['image_ocr_version'] = 1
        key = support.app_module._norm(source)
        phases = []

        def extract(_source, *, progress, **kwargs):
            for count, total, stage in ((71, 71, 'document_native'),
                                        (0, 71, 'document_ocr'),
                                        (1, 71, 'document_ocr'),
                                        (71, 71, 'document_ocr')):
                progress(count, total, stage)
                phases.append(dict(self.app._file_progress[key]))
            return document

        with patch.object(documents, 'extract_document', side_effect=extract):
            self.app._work([source], force=True)
        self.assertEqual([phase['progress_pct'] for phase in phases[:2]], [20, 20])
        self.assertGreater(phases[2]['progress_pct'], 20)
        self.assertLess(phases[2]['progress_pct'], 22)
        self.assertEqual(phases[-1]['progress_pct'], 94)
        self.assertIn('장', phases[-1]['progress_stage'])
        self.assertEqual(self.app._file_progress[key]['progress_pct'], 100)


if __name__ == '__main__':
    unittest.main()
