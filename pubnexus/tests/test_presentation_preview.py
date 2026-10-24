"""Fast PPTX headers, immutable bytes and logical positions without rendering."""
import hashlib
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import zipfile

from pptx import Presentation
from pptx.util import Inches
from pubnexus import presentation_preview as preview
from pubnexus import documents


class PresentationPreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'slides.pptx'
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(10), Inches(5.625)
        for i in range(3):
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = 'Actual slide ' + str(i + 1)
            if i == 1:
                slide._element.set('show', '0')
        deck.save(self.source)

    def test_header_reads_no_slide_xml_or_media_and_preserves_hidden_order(self):
        reads, real = [], zipfile.ZipFile.read
        def read(archive, name, *args, **kwargs):
            reads.append(name)
            return real(archive, name, *args, **kwargs)
        before = self.source.read_bytes()
        with patch.object(zipfile.ZipFile, 'read', read):
            info = preview.describe(self.source, 4)
        self.assertEqual(info['pages'], [[960, 540]] * 3)
        self.assertEqual(info['slide_parts'], [f'ppt/slides/slide{i}.xml' for i in range(1, 4)])
        self.assertEqual(set(reads), {'ppt/presentation.xml', 'ppt/_rels/presentation.xml.rels'})
        self.assertEqual(info['source_revision'], hashlib.sha256(before).hexdigest())
        self.assertIn('?p=4&r=', info['source_data_url'])
        self.assertEqual(self.source.read_bytes(), before)

    def test_raw_data_requires_the_current_exact_revision(self):
        info = preview.describe(self.source, 0)
        self.assertEqual(preview.read_data(self.source, info['source_revision']), self.source.read_bytes())
        self.source.write_bytes(self.source.read_bytes() + b'Changed after metadata')
        with self.assertRaisesRegex(ValueError, '변경'):
            preview.read_data(self.source, info['source_revision'])
        with self.assertRaises(ValueError):
            preview.read_data(self.source, 'anything')

    def test_budget_rejects_before_unbounded_bytes_read(self):
        with patch.object(preview, 'MAX_BYTES', 64), patch.object(zipfile, 'ZipFile', side_effect=AssertionError('Reject size first')):
            with self.assertRaisesRegex(ValueError, '크기'):
                preview.describe(self.source, 0)

    def test_unsafe_package_paths_are_rejected_by_shared_guard(self):
        with zipfile.ZipFile(self.source, 'a') as archive:
            archive.writestr('../outside.xml', '<root/>')
        with self.assertRaisesRegex(ValueError, '경로'):
            preview.describe(self.source, 0)

    def test_native_shapes_remain_logical_until_the_active_renderer_supplies_display_bounds(self):
        origin = {'slide': 2, 'shape': '3', 'shape_bbox_emu': [914400, 1828800, 2743200, 2743200]}
        doc = {'body_text': [{'paragraphs': [{'id': 'normal', 'text': 'text', 'source': origin},
                                             {'id': 'group', 'text': 'group child', 'source': dict(origin, shape='3/1')}]}]}
        mapping = preview.document_map(doc, preview.describe(self.source, 0))
        self.assertEqual(mapping['blocks'][0]['locations'], [])
        self.assertEqual(mapping['blocks'][0]['slide'], 1)
        self.assertEqual(mapping['blocks'][0]['mapping'], 'unmapped')
        self.assertEqual(mapping['blocks'][1]['locations'], [])
        self.assertEqual(mapping['blocks'][1]['mapping'], 'unmapped')
        self.assertEqual(mapping['source_kind'], 'presentation')

    def picture_fixture(self):
        from PIL import Image
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(.5)).text = 'Native group text'
        buffer = io.BytesIO()
        Image.new('RGB', (30, 20), 'white').save(buffer, format='PNG')
        picture = group.shapes.add_picture(io.BytesIO(buffer.getvalue()), Inches(1), Inches(2))
        picture.crop_left = .1
        picture.rotation = 15
        group.rotation = 10
        path = self.root / 'group-picture.pptx'
        deck.save(path)
        with patch('pubnexus.local_ocr.discover', return_value=None):
            doc = documents.extract_document(path, asset_dir=self.root / 'assets')
        return path, doc, str(picture.shape_id), str(group.shape_id)

    def test_picture_source_carries_native_shape_identity_without_inventing_display_bounds(self):
        path, doc, picture_id, group_id = self.picture_fixture()
        figure = doc['figures'][0]
        self.assertEqual(figure['source'].get('source_shape_id'), picture_id)
        self.assertEqual(figure['source'].get('group_shape_ids'), [group_id])
        mapping = preview.document_map(doc, preview.describe(path, 0))
        block = next(item for item in mapping['blocks'] if item['id'] == figure['id'])
        self.assertEqual(block['logical_source']['node_id'], picture_id)
        self.assertEqual(block['logical_source']['group_node_ids'], [group_id])
        self.assertEqual(block['logical_source']['actual_source_image_id'], figure['id'])
        self.assertEqual(block['logical_source']['shape'], '1/2')
        self.assertEqual(block['logical_source']['shape_bbox_emu'], figure['source']['shape_bbox_emu'])
        self.assertEqual(block['locations'], [])
        self.assertEqual(block['mapping'], 'unmapped')

    def test_legacy_picture_and_ocr_shape_paths_resolve_only_from_matching_original_package(self):
        path, doc, picture_id, group_id = self.picture_fixture()
        figure = doc['figures'][0]
        doc['body_text'][0]['paragraphs'].append({'id': 'old-ocr', 'text': 'Only exists inside the picture',
            'source': dict(figure['source'], image_id=figure['id'], method='tesseract',
                source_image_size_px=[30, 20], image_bbox_px=[1, 2, 12, 14],
                raw_ocr_bbox_px=[1, 2, 12, 14], coordinate_space='ocr_image_pixels')})
        for item, *_ in preview.comparison._items(doc):
            item['source'].pop('source_shape_id', None)
            item['source'].pop('group_shape_ids', None)
        figure['source'].pop('image_sha256', None)
        before = copy.deepcopy(doc)
        header = preview.describe(path, 0)
        mapping = preview.document_map(doc, header, path)
        for identity in (figure['id'], 'old-ocr'):
            logical = next(item for item in mapping['blocks'] if item['id'] == identity)['logical_source']
            self.assertEqual(logical['node_id'], picture_id)
            self.assertEqual(logical['group_node_ids'], [group_id])
            self.assertEqual(logical['actual_source_image_id'], figure['id'])
            self.assertEqual(logical.get('image_sha256'), Path(figure['image']).stem)
        ocr_source = next(item for item in mapping['blocks'] if item['id'] == 'old-ocr')['logical_source']
        for name in ('source_image_size_px', 'image_bbox_px', 'raw_ocr_bbox_px', 'coordinate_space'):
            self.assertEqual(ocr_source.get(name), doc['body_text'][0]['paragraphs'][-1]['source'][name])
        self.assertEqual(doc, before, 'Legacy mapping must not rewrite the user canonical store')
        changed = copy.deepcopy(doc)
        changed['sha1'] = '0' * 40
        mapping = preview.document_map(changed, header, path)
        self.assertTrue(all(not item['logical_source'].get('node_id') for item in mapping['blocks']))

    def test_legacy_shape_backfill_rejects_package_changed_after_source_header(self):
        path, doc, _, _ = self.picture_fixture()
        for item, *_ in preview.comparison._items(doc):
            item['source'].pop('source_shape_id', None)
            item['source'].pop('group_shape_ids', None)
        header = preview.describe(path, 0)
        path.write_bytes(path.read_bytes() + b'Changed after header')
        with self.assertRaisesRegex(ValueError, '변경'):
            preview.document_map(doc, header, path)

    def test_legacy_shape_backfill_rejects_dtd_and_malformed_slide_as_typed_read_errors(self):
        for replacement in (b'<!DOCTYPE p:sld [<!ENTITY unsafe "not expanded">]><p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>',
                            b'<p:sld'):
            with self.subTest(xml=replacement[:24]):
                path, doc, _, _ = self.picture_fixture()
                for item, *_ in preview.comparison._items(doc):
                    item['source'].pop('source_shape_id', None)
                    item['source'].pop('group_shape_ids', None)
                buffer = io.BytesIO()
                with zipfile.ZipFile(path) as original, zipfile.ZipFile(buffer, 'w') as modified:
                    for entry in original.infolist():
                        modified.writestr(entry, replacement if entry.filename == 'ppt/slides/slide1.xml'
                                         else original.read(entry.filename))
                path.write_bytes(buffer.getvalue())
                doc['sha1'] = hashlib.sha1(path.read_bytes()).hexdigest()
                before_bytes, before_doc = path.read_bytes(), copy.deepcopy(doc)
                header = preview.describe(path, 0)
                with self.assertRaisesRegex(ValueError, '슬라이드'):
                    preview.document_map(doc, header, path)
                self.assertEqual(path.read_bytes(), before_bytes)
                self.assertEqual(doc, before_doc)

    def test_asset_paths_are_allowlisted_and_vendor_hash_must_match(self):
        web, vendor = self.root / 'web', self.root / 'vendor'
        web.mkdir(); vendor.mkdir()
        for name in ('presentation-frame.html', 'presentation-frame.mjs'):
            (web / name).write_text('fixture', encoding='utf-8')
        bundle = vendor / 'renderer.mjs'
        bundle.write_bytes(b'Pinned original bundle')
        (vendor / 'manifest.json').write_text(json.dumps({'version': preview.VERSION,
            'sha256': hashlib.sha256(bundle.read_bytes()).hexdigest()}), encoding='utf-8')
        with patch.object(preview, '_web_root', return_value=web), patch.object(preview, '_vendor_root', return_value=vendor):
            self.assertTrue(preview.available())
            self.assertEqual(preview.asset('vendor'), b'Pinned original bundle')
            with self.assertRaises(ValueError):
                preview.asset('../settings.json')
            bundle.write_bytes(b'Tampered bundle')
            self.assertFalse(preview.available())


if __name__ == '__main__':
    unittest.main()
