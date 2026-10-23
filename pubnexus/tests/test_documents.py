"""Native 문서 어댑터의 공개 경계. 모든 문서·그림 fixture는 TEMP에서 만든다."""
from __future__ import annotations

import os
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from pubnexus.documents import DocumentExtractionError, extract_document


class NativeDocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="pubnexus-documents-")
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        env = patch.dict(os.environ, {"PUBNEXUS_STORE": str(self.base / "cache")})
        env.start()
        self.addCleanup(env.stop)

    def test_csv_preserves_quoted_cells_and_reports_logical_source_locations(self):
        path = self.base / "한글 표.csv"
        path.write_text('이름,내용\n홍길동,"첫 줄\n둘째 줄"\n수식,=1+2\n', encoding="utf-8-sig", newline="")
        before = path.read_bytes()

        doc = extract_document(path)

        table = doc["tables"][0]
        self.assertEqual((table["rows"], table["cols"]), (3, 2))
        self.assertEqual(table["cells"][3]["text"], "첫 줄\n둘째 줄")
        self.assertEqual(table["cells"][5]["text"], "=1+2")
        self.assertEqual(doc["source_map"]["coordinate_space"], "logical_document")
        self.assertNotIn("page", table["source"])
        self.assertEqual(doc["extraction_report"]["status"], "complete")
        self.assertEqual(path.read_bytes(), before)

    def test_ragged_csv_obeys_rectangular_cell_budget(self):
        path = self.base / "wide.csv"
        path.write_text("a,b,c,d,e\nf\n", encoding="utf-8")

        with patch("pubnexus.documents.MAX_CELLS", 8):
            with self.assertRaises(DocumentExtractionError) as raised:
                extract_document(path)

        self.assertEqual(raised.exception.code, "limit")

    def test_empty_csv_records_are_also_bounded(self):
        path = self.base / "empty-rows.csv"
        path.write_bytes(b"\n" * 9)
        with patch("pubnexus.documents.MAX_CELLS", 8):
            with self.assertRaises(DocumentExtractionError) as raised:
                extract_document(path)
        self.assertEqual(raised.exception.code, "limit")

    def test_docx_preserves_block_order_merged_cells_and_original_image(self):
        from docx import Document
        from PIL import Image
        path = self.base / "보고서.docx"
        image = self.base / "sample.png"
        Image.new("RGB", (3, 2), "red").save(image)
        original = image.read_bytes()
        source = Document()
        source.add_heading("한글 보고서", level=1)
        source.add_paragraph("DOI: 원문의 첫 문장을 그대로 보존")
        table = source.add_table(rows=3, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "병합 제목"
        table.cell(1, 0).merge(table.cell(2, 0)).text = "세로 병합"
        table.cell(1, 1).text = "A"
        table.cell(2, 1).text = "B"
        source.add_paragraph("표 다음 문장")
        source.add_picture(str(image))
        source.save(path)

        doc = extract_document(path)

        cells = doc["tables"][0]["cells"]
        self.assertTrue(any(c["text"] == "병합 제목" and c["col_span"] == 2 for c in cells))
        self.assertTrue(any(c["text"] == "세로 병합" and c["row_span"] == 2 for c in cells))
        self.assertIn('colspan="2"', doc["tables"][0]["html"])
        paragraphs = [p for s in doc["body_text"] for p in s["paragraphs"]]
        first = next(i for i, p in enumerate(paragraphs) if p["text"].startswith("DOI:"))
        anchor = next(i for i, p in enumerate(paragraphs) if p["refs_table"])
        after = next(i for i, p in enumerate(paragraphs) if p["text"] == "표 다음 문장")
        self.assertLess(first, anchor)
        self.assertLess(anchor, after)
        asset = Path(doc["extraction_report"]["asset_base"]) / doc["figures"][0]["image"]
        self.assertEqual(asset.read_bytes(), original)
        self.assertNotIn("page", paragraphs[first]["source"])

    def test_pptx_preserves_slide_text_merged_table_image_and_speaker_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        from PIL import Image
        path = self.base / "발표.pptx"
        image = self.base / "slide.png"
        Image.new("RGB", (4, 3), "blue").save(image)
        native = Presentation()
        slide = native.slides.add_slide(native.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "한글 발표 내용"
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(5), Inches(2)).table
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 0).text = "병합 머리글"
        table.cell(1, 0).text = "A"
        table.cell(1, 1).text = "B"
        slide.shapes.add_picture(str(image), Inches(1), Inches(4))
        slide.notes_slide.notes_text_frame.text = "발표자만 보는 노트"
        native.save(path)

        doc = extract_document(path)

        paragraphs = [p for s in doc["body_text"] for p in s["paragraphs"]]
        text = next(p for p in paragraphs if p["text"] == "한글 발표 내용")
        notes = next(p for p in paragraphs if p["text"] == "발표자만 보는 노트")
        self.assertEqual(text["source"]["slide"], 1)
        self.assertEqual(notes["source"]["slide"], 1)
        self.assertEqual(doc["tables"][0]["cells"][0]["col_span"], 2)
        asset = Path(doc["extraction_report"]["asset_base"]) / doc["figures"][0]["image"]
        self.assertEqual(asset.read_bytes(), image.read_bytes())

    def test_xlsx_preserves_merges_formula_and_original_embedded_image(self):
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as WorkbookImage
        from PIL import Image
        image = self.base / "sheet.png"
        Image.new("RGB", (5, 4), "green").save(image)
        path = self.base / "통합 문서.xlsx"
        native = Workbook()
        sheet = native.active
        sheet.title = "한글 시트"
        sheet.merge_cells("A1:B1")
        sheet["A1"] = "합친 제목"
        sheet["A2"] = 12
        sheet["B2"] = "=SUM(1,2)"
        sheet.add_image(WorkbookImage(str(image)), "D1")
        native.save(path)
        native.close()

        doc = extract_document(path)

        table = doc["tables"][0]
        self.assertEqual(table["source"]["sheet"], "한글 시트")
        self.assertTrue(any(c["text"] == "합친 제목" and c["col_span"] == 2 for c in table["cells"]))
        formula = next(c for c in table["cells"] if c.get("formula"))
        self.assertEqual(formula["formula"], "=SUM(1,2)")
        self.assertIsNone(formula["cached_value"])
        self.assertEqual(formula["source"]["cell"], "B2")
        self.assertTrue(doc["extraction_report"]["warnings"])
        figure = doc["figures"][0]
        asset = Path(doc["extraction_report"]["asset_base"]) / figure["image"]
        self.assertEqual(asset.read_bytes(), image.read_bytes())

    def test_sparse_xlsx_keeps_empty_cell_coordinates(self):
        from openpyxl import Workbook
        path = self.base / "sparse.xlsx"
        native = Workbook()
        native.active["B2"] = "only one value"
        native.save(path)
        native.close()

        doc = extract_document(path)

        self.assertEqual(doc["tables"][0]["cells"][0]["source"]["cell"], "A1")
        self.assertEqual(doc["tables"][0]["cells"][-1]["text"], "only one value")

    def test_xlsx_date_formula_cache_is_json_serializable(self):
        from openpyxl import Workbook
        path = self.base / "date-cache.xlsx"
        native = Workbook()
        native.active["A1"] = "=TODAY()"
        native.active["A1"].number_format = "yyyy-mm-dd"
        native.save(path)
        native.close()
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(b"<v></v>", b"<v>45292</v>")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)

        doc = extract_document(path)

        self.assertEqual(doc["tables"][0]["cells"][0]["cached_value"], "2024-01-01T00:00:00")
        json.dumps(doc, ensure_ascii=False)

    def test_xlsx_does_not_trust_an_underreported_worksheet_dimension(self):
        from openpyxl import Workbook
        path = self.base / "wrong-dimension.xlsx"
        native = Workbook()
        native.active["B2"] = "실제 뒤쪽 셀"
        native.save(path)
        native.close()
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries["xl/worksheets/sheet1.xml"] = re.sub(
            br'<dimension ref="[^"]+"', b'<dimension ref="A1"', entries["xl/worksheets/sheet1.xml"])
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in entries.items():
                archive.writestr(name, data)

        doc = extract_document(path)

        self.assertEqual((doc["tables"][0]["rows"], doc["tables"][0]["cols"]), (2, 2))
        self.assertEqual(doc["tables"][0]["cells"][-1]["text"], "실제 뒤쪽 셀")

    def test_xlsx_array_formula_is_preserved_without_library_objects(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.formula import ArrayFormula
        path = self.base / "array-formula.xlsx"
        native = Workbook()
        native.active["A1"] = ArrayFormula(ref="A1:A2", text="=SEQUENCE(2)")
        native.save(path)
        native.close()

        doc = extract_document(path)

        cell = doc["tables"][0]["cells"][0]
        self.assertEqual(cell["formula"], "=SEQUENCE(2)")
        self.assertEqual(cell["formula_attributes"]["ref"], "A1:A2")
        json.dumps(doc, ensure_ascii=False)

    def test_cp949_text_preserves_literal_content_and_line_hints(self):
        path = self.base / "text.txt"
        text = "DOI: 이 문장은 지우지 않습니다.\n\n  DEAR EDITOR, 원문 들여쓰기\n"
        path.write_bytes(text.encode("cp949"))

        doc = extract_document(path)

        paragraph = doc["body_text"][0]["paragraphs"][0]
        self.assertEqual(paragraph["text"], text)
        self.assertEqual(paragraph["source"]["line_start"], 1)
        self.assertNotIn("page", paragraph["source"])

    def test_html_preserves_merged_table_and_local_image_without_active_content(self):
        from PIL import Image
        image = self.base / "web.png"
        Image.new("RGB", (2, 2), "yellow").save(image)
        path = self.base / "page.html"
        path.write_text('''<html><head><meta charset="utf-8"><title>문서 제목</title></head>
<body><h1>첫 절</h1><p>앞 <strong>강조</strong> 뒤</p><script>RUN_EVIL_SCRIPT</script>
<table><tr><th colspan="2">합친 헤더</th></tr><tr><td rowspan="2">A</td><td>B</td></tr>
<tr><td>C</td></tr></table><img src="web.png" alt="로컬 그림">
<img src="https://example.invalid/remote.png"><p>마지막 문장</p></body></html>''', encoding="utf-8")

        doc = extract_document(path)

        self.assertEqual(doc["meta"]["title"], "문서 제목")
        texts = [p["text"] for s in doc["body_text"] for p in s["paragraphs"]]
        self.assertIn("앞 강조 뒤", texts)
        self.assertNotIn("RUN_EVIL_SCRIPT", "".join(texts))
        self.assertTrue(any(c["row_span"] == 2 for c in doc["tables"][0]["cells"]))
        self.assertEqual(len(doc["figures"]), 1)
        asset = Path(doc["extraction_report"]["asset_base"]) / doc["figures"][0]["image"]
        self.assertEqual(asset.read_bytes(), image.read_bytes())
        self.assertEqual(doc["extraction_report"]["status"], "partial")

    def test_html_table_line_breaks_do_not_concatenate_numbers(self):
        path = self.base / "numbers.html"
        path.write_text("<table><tr><td>10<br>20</td><td><p>30</p><p>40</p></td></tr></table>", encoding="utf-8")
        doc = extract_document(path)
        self.assertEqual(doc["tables"][0]["cells"][0]["text"], "10\n20")
        self.assertEqual(doc["tables"][0]["cells"][1]["text"], "30\n40")

    def test_hwp_missing_configured_engine_has_an_actionable_dependency_error(self):
        path = self.base / "test.hwp"
        path.write_bytes(b"synthetic engine dependency fixture")

        with patch.dict(os.environ, {"PUBNEXUS_RHWP": str(self.base / "missing-rhwp.exe")}):
            with self.assertRaises(DocumentExtractionError) as raised:
                extract_document(path)

        self.assertEqual(raised.exception.code, "dependency")
        self.assertIn("PUBNEXUS_RHWP", str(raised.exception))

    def test_actual_rhwp_extracts_only_synthetic_hwpx_and_hwp5(self):
        engine = os.environ.get("PUBNEXUS_RHWP") or shutil.which("rhwp")
        if not engine:
            try:
                from pubnexus.preview import find_engine
                engine = find_engine("rhwp")
            except ImportError:
                pass
        if not engine:
            self.skipTest("rHWP가 없는 환경; mock/의존성 오류 시험은 별도 실행")
        spec = self.base / "spec.json"
        spec.write_text(json.dumps({"version": "1", "title": "합성 한글 문서", "blocks": [
            {"type": "heading", "level": 1, "text": "첫 절"},
            {"type": "paragraph", "text": "한글 본문을 원문 그대로 보존합니다."},
            {"type": "table", "rows": [["열 하나", "열 둘"], ["값 A", "값 B"]]},
        ]}, ensure_ascii=False), encoding="utf-8")
        hwpx = self.base / "synthetic.hwpx"
        result = subprocess.run([str(engine), "scaffold", str(spec), "-o", str(hwpx), "--json"],
                                capture_output=True, timeout=20, shell=False)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", "replace"))
        hwp = self.base / "synthetic.hwp"
        result = subprocess.run([str(engine), "convert", str(hwpx), str(hwp)],
                                capture_output=True, timeout=20, shell=False)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", "replace"))
        for path in (hwpx, hwp):
            with self.subTest(suffix=path.suffix):
                before = path.read_bytes()
                with patch.dict(os.environ, {"PUBNEXUS_RHWP": str(engine)}):
                    doc = extract_document(path, cancel=lambda: False)
                text = "\n".join(p["text"] for s in doc["body_text"] for p in s["paragraphs"])
                self.assertIn("한글 본문을 원문 그대로 보존합니다.", text)
                self.assertEqual((doc["tables"][0]["rows"], doc["tables"][0]["cols"]), (2, 2))
                self.assertEqual(doc["tables"][0]["cells"][-1]["text"], "값 B")
                self.assertEqual(doc["source_map"]["engine"], "rHWP DocLang 0.6")
                self.assertEqual(path.read_bytes(), before)
        merged = self.base / "merged.hwpx"
        result = subprocess.run([str(engine), "edit", "merge-cells", str(hwpx), "--table", "0",
                                 "--row", "0", "--col", "0", "--end-row", "1", "--end-col", "1",
                                 "-o", str(merged), "--json"], capture_output=True, timeout=20, shell=False)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8", "replace"))
        with patch.dict(os.environ, {"PUBNEXUS_RHWP": str(engine)}):
            merged_doc = extract_document(merged)
        cell = merged_doc["tables"][0]["cells"][0]
        self.assertEqual((cell["row_span"], cell["col_span"]), (2, 2))
        self.assertIn('rowspan="2" colspan="2"', merged_doc["tables"][0]["html"])

    def test_rhwp_timeout_is_an_explicit_failure(self):
        path = self.base / "timeout.hwp"
        path.write_bytes(b"synthetic subprocess boundary fixture")
        engine = self.base / "fake-rhwp.exe"
        engine.write_bytes(b"not executed; subprocess mocked")
        with patch.dict(os.environ, {"PUBNEXUS_RHWP": str(engine)}):
            with patch("subprocess.run", side_effect=subprocess.TimeoutExpired(["rhwp"], 0.01)):
                with self.assertRaises(DocumentExtractionError) as raised:
                    extract_document(path, timeout=1)
        self.assertEqual(raised.exception.code, "timeout")

    def test_cancel_during_native_rows_raises_the_shared_cancel_signal(self):
        from pubnexus.general_pdf import ExtractionCancelled
        path = self.base / "cancel.csv"
        path.write_bytes(b"header,value\na,1\nb,2\nc,3\n")
        calls = 0

        def cancel():
            nonlocal calls
            calls += 1
            return calls >= 5

        with self.assertRaises(ExtractionCancelled):
            extract_document(path, cancel=cancel)
        self.assertEqual(path.read_bytes(), b"header,value\na,1\nb,2\nc,3\n")

    def test_cancel_kills_only_the_private_rhwp_process_before_returning(self):
        from pubnexus.general_pdf import ExtractionCancelled
        path = self.base / "cancel.hwp"
        path.write_bytes(b"synthetic process cancellation fixture")
        engine = self.base / "fake-rhwp.exe"
        engine.write_bytes(b"never executed")
        state = {"started": False, "killed": False}

        class Process:
            returncode = None

            def poll(self):
                return self.returncode

            def kill(self):
                state["killed"] = True
                self.returncode = -9

            def wait(self, timeout=None):
                return self.returncode

        def start(*args, **kwargs):
            state["started"] = True
            return Process()

        with patch.dict(os.environ, {"PUBNEXUS_RHWP": str(engine)}):
            with patch("subprocess.Popen", side_effect=start):
                with patch("subprocess.run", side_effect=AssertionError("cancellable process must be polled")):
                    with self.assertRaises(ExtractionCancelled):
                        extract_document(path, cancel=lambda: state["started"])
        self.assertTrue(state["killed"])

    def test_empty_csv_is_not_reported_as_a_successful_extraction(self):
        path = self.base / "blank.csv"
        path.write_bytes(b"\n\n")
        with self.assertRaises(DocumentExtractionError) as raised:
            extract_document(path)
        self.assertEqual(raised.exception.code, "empty")

    def test_unsupported_encrypted_and_corrupt_documents_have_explicit_error_codes(self):
        cases = [("macro.docm", b"never executed", "unsupported"),
                 ("encrypted.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1fixture", "encrypted"),
                 ("corrupt.xlsx", b"not a ZIP archive", "malformed")]
        for name, data, expected in cases:
            with self.subTest(name=name):
                path = self.base / name
                path.write_bytes(data)
                with self.assertRaises(DocumentExtractionError) as raised:
                    extract_document(path)
                self.assertEqual(raised.exception.code, expected)
                self.assertEqual(path.read_bytes(), data)

    def test_epub_uses_spine_order_and_preserves_internal_image_bytes(self):
        from PIL import Image
        image = self.base / "epub.png"
        Image.new("RGB", (3, 3), "purple").save(image)
        path = self.base / "book.epub"
        container = '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="OEBPS/book.opf"/></rootfiles></container>'
        opf = '''<package xmlns="http://www.idpf.org/2007/opf" version="3.0"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>책 제목</dc:title></metadata>
<manifest><item id="later" href="a.xhtml" media-type="application/xhtml+xml"/><item id="first" href="z.xhtml" media-type="application/xhtml+xml"/></manifest>
<spine><itemref idref="first"/><itemref idref="later"/></spine></package>'''
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("mimetype", "application/epub+zip")
            archive.writestr("META-INF/container.xml", container)
            archive.writestr("OEBPS/book.opf", opf)
            archive.writestr("OEBPS/z.xhtml", '<html><body><p>먼저 읽을 본문</p><img src="images/p.png"/></body></html>')
            archive.writestr("OEBPS/a.xhtml", '<html><body><p>나중 본문</p></body></html>')
            archive.writestr("OEBPS/images/p.png", image.read_bytes())

        doc = extract_document(path)

        self.assertEqual(doc["meta"]["title"], "책 제목")
        paragraphs = [p for s in doc["body_text"] for p in s["paragraphs"] if p["text"]]
        self.assertEqual([p["text"] for p in paragraphs], ["먼저 읽을 본문", "나중 본문"])
        self.assertEqual(paragraphs[0]["source"]["spine_index"], 1)
        asset = Path(doc["extraction_report"]["asset_base"]) / doc["figures"][0]["image"]
        self.assertEqual(asset.read_bytes(), image.read_bytes())

    def test_odt_preserves_spans_explicit_spaces_notes_and_image(self):
        from PIL import Image
        image = self.base / "odf.png"
        Image.new("RGB", (2, 3), "orange").save(image)
        path = self.base / "document.odt"
        content = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:xlink="http://www.w3.org/1999/xlink">
<office:body><office:text><text:h text:outline-level="1">첫 제목</text:h>
<text:p>원문<text:s text:c="2"/>공백<text:note text:id="n1" text:note-class="footnote"><text:note-citation>1</text:note-citation><text:note-body><text:p>각주 본문</text:p></text:note-body></text:note></text:p>
<table:table table:name="표"><table:table-row><table:table-cell table:number-columns-spanned="2"><text:p>머리글</text:p></table:table-cell><table:covered-table-cell/></table:table-row>
<table:table-row><table:table-cell table:number-rows-spanned="2"><text:p>세로 병합</text:p></table:table-cell><table:table-cell><text:p>A</text:p></table:table-cell></table:table-row>
<table:table-row><table:covered-table-cell/><table:table-cell><text:p>B</text:p></table:table-cell></table:table-row></table:table>
<text:p><draw:frame><draw:image xlink:href="Pictures/p.png"/></draw:frame></text:p>
</office:text></office:body></office:document-content>'''
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("mimetype", "application/vnd.oasis.opendocument.text")
            archive.writestr("content.xml", content)
            archive.writestr("Pictures/p.png", image.read_bytes())

        doc = extract_document(path)

        texts = [p["text"] for s in doc["body_text"] for p in s["paragraphs"]]
        self.assertIn("원문  공백[1]", texts)
        self.assertIn("각주 본문", texts)
        self.assertTrue(any(u["kind"] == "footnote" for u in doc["source_map"]["units"]))
        self.assertTrue(any(c["row_span"] == 2 for c in doc["tables"][0]["cells"]))
        self.assertTrue(any(c["col_span"] == 2 for c in doc["tables"][0]["cells"]))
        asset = Path(doc["extraction_report"]["asset_base"]) / doc["figures"][0]["image"]
        self.assertEqual(asset.read_bytes(), image.read_bytes())

    def test_ods_preserves_data_repeats_and_formula_but_does_not_expand_blank_capacity(self):
        path = self.base / "data.ods"
        content = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
<office:body><office:spreadsheet><table:table table:name="데이터">
<table:table-row><table:table-cell office:value-type="float" office:value="10"/><table:table-cell table:formula="of:=SUM([.A1])" office:value-type="float" office:value="10"/><table:table-cell table:number-columns-repeated="16382"/></table:table-row>
<table:table-row table:number-rows-repeated="2"><table:table-cell table:number-columns-repeated="2" office:value-type="string" office:string-value="반복"/></table:table-row>
<table:table-row table:number-rows-repeated="1048573"><table:table-cell table:number-columns-repeated="16384"/></table:table-row>
</table:table></office:spreadsheet></office:body></office:document-content>'''
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            archive.writestr("content.xml", content)

        doc = extract_document(path)

        table = doc["tables"][0]
        self.assertEqual((table["rows"], table["cols"]), (3, 2))
        self.assertEqual(table["source"]["sheet"], "데이터")
        self.assertEqual(table["cells"][1]["formula"], "of:=SUM([.A1])")
        self.assertEqual(table["cells"][1]["cached_value"], "10")
        self.assertEqual([c["text"] for c in table["cells"][-4:]], ["반복"] * 4)
        self.assertEqual(table["source"]["omitted_blank_capacity"]["rows"], 1048573)

    def test_odp_preserves_slide_order_and_speaker_notes_as_separate_units(self):
        path = self.base / "slides.odp"
        content = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"><office:body><office:presentation>
<draw:page draw:name="첫 장"><draw:frame><draw:text-box><text:p>첫 슬라이드</text:p></draw:text-box></draw:frame>
<presentation:notes><draw:frame><draw:text-box><text:p>발표자 노트</text:p></draw:text-box></draw:frame></presentation:notes></draw:page>
<draw:page draw:name="둘째 장"><draw:frame><draw:text-box><text:p>둘째 슬라이드</text:p></draw:text-box></draw:frame></draw:page>
</office:presentation></office:body></office:document-content>'''
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("mimetype", "application/vnd.oasis.opendocument.presentation")
            archive.writestr("content.xml", content)

        doc = extract_document(path)

        units = doc["source_map"]["units"]
        self.assertEqual([u["label"] for u in units if u["kind"] == "slide"], ["첫 장", "둘째 장"])
        note = next(u for u in units if u["kind"] == "speaker_notes")
        self.assertEqual(note["slide"], 1)
        texts = [p["text"] for s in doc["body_text"] for p in s["paragraphs"]]
        self.assertEqual(texts, ["첫 슬라이드", "발표자 노트", "둘째 슬라이드"])

    def test_eml_uses_html_alternative_and_cid_images_but_does_not_execute_attachments(self):
        from email.message import EmailMessage
        from PIL import Image
        image = self.base / "mail.png"
        Image.new("RGB", (2, 4), "cyan").save(image)
        message = EmailMessage()
        message["Subject"] = "한글 메일 제목"
        message["From"] = "sender@example.test"
        message["To"] = "reader@example.test"
        message["Message-ID"] = "<message-1@example.test>"
        message.set_content("PLAIN_ALTERNATIVE_SHOULD_NOT_DUPLICATE")
        message.add_alternative('<html><body><p>메일 본문</p><table><tr><td colspan="2">병합 표</td></tr></table><img src="cid:pic1"/></body></html>', subtype="html")
        html_part = message.get_payload()[1]
        html_part.add_related(image.read_bytes(), maintype="image", subtype="png", cid="<pic1>",
                              filename="inline.png", disposition="inline")
        message.add_attachment(b"not opened or converted", maintype="application", subtype="pdf", filename="../attachment.pdf")
        path = self.base / "message.eml"
        path.write_bytes(message.as_bytes())
        before = path.read_bytes()

        doc = extract_document(path)

        self.assertEqual(doc["meta"]["title"], "한글 메일 제목")
        texts = [p["text"] for s in doc["body_text"] for p in s["paragraphs"]]
        self.assertIn("메일 본문", texts)
        self.assertNotIn("PLAIN_ALTERNATIVE_SHOULD_NOT_DUPLICATE", "".join(texts))
        self.assertEqual(doc["tables"][0]["cells"][0]["col_span"], 2)
        self.assertEqual(doc["mail_messages"][0]["message_id"], "<message-1@example.test>")
        self.assertTrue(any(a["filename"] == "../attachment.pdf" for a in doc["mail_messages"][0]["attachments"]))
        asset = Path(doc["extraction_report"]["asset_base"]) / doc["figures"][0]["image"]
        self.assertEqual(asset.read_bytes(), image.read_bytes())
        self.assertFalse((self.base / "attachment.pdf").exists())
        self.assertEqual(path.read_bytes(), before)


    def test_deeply_nested_mail_is_a_structured_limit_error_and_keeps_original(self):
        pieces = [b"MIME-Version: 1.0\r\n"]
        for index in range(1100):
            pieces.append(f'Content-Type: multipart/mixed; boundary="b{index}"\r\n\r\n--b{index}\r\n'.encode())
        pieces.append(b"Content-Type: text/plain; charset=utf-8\r\n\r\nbody")
        for index in reversed(range(1100)):
            pieces.append(f"\r\n--b{index}--\r\n".encode())
        payload = b"".join(pieces)
        for suffix in (".eml", ".mbox"):
            with self.subTest(suffix=suffix):
                path = self.base / ("nested" + suffix)
                original = (b"From sender@example.test Fri Oct 23 09:00:00 2026\n" + payload
                            if suffix == ".mbox" else payload)
                path.write_bytes(original)
                with self.assertRaises(DocumentExtractionError) as raised:
                    extract_document(path, timeout=5)
                self.assertEqual(raised.exception.code, "limit")
                self.assertEqual(path.read_bytes(), original)

    def test_odf_repeated_spaces_obey_a_cumulative_text_budget(self):
        path = self.base / "spaces.odt"
        content = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>
<text:p>a<text:s text:c="16"/><text:s text:c="16"/><text:s text:c="16"/>b</text:p>
</office:text></office:body></office:document-content>'''
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("content.xml", content)
        with patch("pubnexus.documents.MAX_TEXT_BYTES", 32):
            with self.assertRaises(DocumentExtractionError) as raised:
                extract_document(path)
        self.assertEqual(raised.exception.code, "limit")

    def test_markdown_collects_only_confined_local_images_and_exports_portable_links(self):
        from PIL import Image
        from pubnexus.export import export_document
        source_dir = self.base / "source"
        (source_dir / "images").mkdir(parents=True)
        image = source_dir / "images" / "한 글.png"
        Image.new("RGB", (3, 2), "purple").save(image)
        outside = self.base / "private.png"
        Image.new("RGB", (1, 1), "red").save(outside)
        path = source_dir / "notes.md"
        text = ('# 원래 제목\n\n**굵은 글씨** [링크](https://example.test)\n\n'
                '![로컬 그림][local]\n\n[local]: images/한%20글.png\n\n'
                '```text\n![그림 아님](missing.png)\n```\n\n'
                '![외부](https://example.test/not-fetched.png)\n\n'
                '![폴더 밖](../private.png)\n')
        path.write_text(text, encoding="utf-8", newline="")

        doc = extract_document(path)
        result = export_document(doc, path)

        self.assertEqual(doc["native_markdown"], text)
        self.assertEqual(path.read_text(encoding="utf-8"), text)
        self.assertEqual(len(doc["figures"]), 1)
        self.assertTrue(all(Path(doc["extraction_report"]["asset_base"], f["image"]).read_bytes()
                            == image.read_bytes() for f in doc["figures"]))
        output = result.markdown_path.read_text(encoding="utf-8")
        self.assertIn('**굵은 글씨** [링크](https://example.test)', output)
        self.assertIn('```text\n![그림 아님](missing.png)\n```', output)
        self.assertRegex(output, r'!\[로컬 그림\]\(assets/[a-f0-9]+/[a-f0-9]+\.png\)')
        self.assertNotIn('![외부]', output)
        self.assertNotIn('![폴더 밖]', output)
        self.assertNotIn(str(source_dir), output)
        self.assertEqual(doc["extraction_report"]["status"], "partial")
        self.assertTrue(any("외부" in w for w in doc["extraction_report"]["warnings"]))

    def test_mbox_preserves_message_order_and_unescapes_only_mboxrd_body_lines(self):
        path = self.base / "messages.mbox"
        content = (b"From sender@example.test Fri Oct 23 09:00:00 2026\n"
                   b"Subject: First\nMessage-ID: <one@example.test>\n"
                   b"Content-Type: text/plain; charset=utf-8\n\n"
                   b"first body\n>From escaped body\n>>From nested quote\n\n"
                   b"From sender@example.test Fri Oct 23 09:01:00 2026\n"
                   b"Subject: Second\nMessage-ID: <two@example.test>\n"
                   b"Content-Type: text/plain; charset=utf-8\n\nsecond body\n")
        path.write_bytes(content)

        doc = extract_document(path)

        self.assertEqual([m["subject"] for m in doc["mail_messages"]], ["First", "Second"])
        bodies = [p["text"] for s in doc["body_text"] if s["source"].get("message_index")
                  for p in s["paragraphs"]]
        self.assertIn("first body\nFrom escaped body\n>From nested quote\n\n", bodies)
        self.assertIn("second body\n", bodies)
        self.assertEqual([u["message_index"] for u in doc["source_map"]["units"]
                          if u["kind"] == "email_message"], [1, 2])
        self.assertEqual(path.read_bytes(), content)
        with patch("pubnexus.documents.MAX_MAIL_MESSAGES", 1):
            with self.assertRaises(DocumentExtractionError) as raised:
                extract_document(path)
        self.assertEqual(raised.exception.code, "limit")


if __name__ == "__main__":
    unittest.main()
