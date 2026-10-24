"""Fast source windows and images must be tied to the selected original."""
import io
import urllib.parse
import unittest
from unittest.mock import Mock, patch

from docx import Document
from openpyxl import Workbook
from PIL import Image

import test_app_export as support


class NativeSourceApiTests(unittest.TestCase):
    def setUp(self):
        support.AppExportTests.setUp(self)
        self.addCleanup(self.app.renderer.close_all)
        conversion = patch('pubnexus.preview.preview_path', side_effect=AssertionError('no Office/PDF conversion'))
        conversion.start()
        self.addCleanup(conversion.stop)

    def docx(self):
        path = self.root / 'word.docx'
        doc = Document()
        doc.add_paragraph('Native Word original 7192')
        stream = io.BytesIO()
        Image.new('RGB', (10, 8), 'red').save(stream, format='PNG')
        doc.add_picture(io.BytesIO(stream.getvalue()))
        doc.save(path)
        self.app.pdfs = [path]
        return path, stream.getvalue()

    def handler(self, route):
        handler = object.__new__(support.app_module.Handler)
        handler.app = self.app
        handler.path = route
        handler._json = Mock()
        handler._send = Mock()
        handler.send_error = Mock()
        return handler

    def test_word_original_starts_without_pdf_or_markdown(self):
        path, _ = self.docx()
        with patch('pubnexus.preview.preview_path', side_effect=AssertionError('no Office/PDF conversion')), \
             patch.object(self.app, 'doc_view', side_effect=AssertionError('no extraction')):
            result = self.app.source_view(0, expected_path=str(path))
        self.assertEqual(result.get('source_kind'), 'content', result)
        self.assertIn('Native Word', result['blocks'][0]['text'])
        self.assertFalse(self.app.exported)

    def test_native_word_search_needs_no_pdf(self):
        path, _ = self.docx()
        result = support.app_module.Api(self.app).search_source(0, '7192', str(path))
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['results'][0]['locations'][0]['kind'], 'content')

    def test_excel_window_and_image_routes_are_revision_bound(self):
        path, png = self.docx()
        view = self.app.source_view(0)
        image = next(block for block in view['blocks'] if block['kind'] == 'image')
        params = {'p': 0, 'expected_path': str(path), 'r': view['source_revision'], 'id': image['asset_id']}
        handler = self.handler('/native-asset?' + urllib.parse.urlencode(params))
        handler.do_GET()
        self.assertEqual(handler._send.call_args.args[:2], (png, 'image/png'))
        book = Workbook()
        book.active.cell(1, 1, 'First window')
        book.active.cell(121, 1, 'Later window')
        sheet_path = self.root / 'sheet.xlsx'
        book.save(sheet_path)
        book.close()
        self.app.pdfs = [sheet_path]
        view = self.app.source_view(0)
        self.assertEqual(view['source_kind'], 'sheet')
        params = {'p': 0, 'expected_path': str(sheet_path), 'r': view['source_revision'], 'start_row': 121, 'row_count': 1}
        handler = self.handler('/native-preview?' + urllib.parse.urlencode(params))
        handler.do_GET()
        result = handler._json.call_args.args[0]
        self.assertEqual(result['rows'][0][0], 'Later window')
        self.assertEqual(result['path'], str(sheet_path))
        params['expected_path'] = str(path)
        handler = self.handler('/native-preview?' + urllib.parse.urlencode(params))
        handler.do_GET()
        self.assertEqual(handler._json.call_args.args[0]['code'], 'stale_selection')


if __name__ == '__main__':
    unittest.main()
