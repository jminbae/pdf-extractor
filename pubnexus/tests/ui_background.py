"""Offline frontend-only seam for background extraction interaction tests.

The real UI JavaScript is served unchanged. Native responses are deterministic
fixtures; no App instance, document extractor, user file or settings is opened.
"""
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import argparse
import io
import json
from pathlib import Path
import sys
import tempfile
import threading

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from pubnexus.ui import INDEX_HTML


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--presentation-vendor", type=Path)
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="pnx-ui-background-") as temporary:
        root = Path(temporary)
        presentation_bytes = None
        if args.presentation_vendor:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from PIL import Image, ImageDraw
            picture = Image.new("RGB", (800, 450), "#1687b8")
            ImageDraw.Draw(picture).rectangle((400, 0, 799, 449), fill="#ebaa35")
            image_bytes = io.BytesIO()
            picture.save(image_bytes, format="JPEG", quality=92)
            deck = Presentation()
            for number, marker in enumerate(("FIRST", "SECOND", "THIRD"), 1):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                text = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(9), Inches(1)).text_frame
                text.text = f"FAST_FRAME_{marker}_7391"
                text.paragraphs[0].font.size = Pt(28)
                slide.shapes.add_picture(io.BytesIO(image_bytes.getvalue()), Inches(1), Inches(2), width=Inches(7))
            deck.save(root / "frame-deck.pptx")
            presentation_bytes = (root / "frame-deck.pptx").read_bytes()
        files = [{"name": "A.pdf", "path": str(root / "A.pdf"), "done": False, "md_ready": False, "status": "ready", "pages": 30},
                 {"name": "B.pdf", "path": str(root / "B.pdf"), "done": True, "md_ready": True, "status": "complete", "pages": 30,
                  "markdown": str(root / "Markdown" / "B.md")},
                 {"name": "C.pdf", "path": str(root / "C.pdf"), "done": False, "md_ready": False, "status": "ready", "pages": 30}]
        state = {"files": files, "path": str(root), "output_dir": str(root / "Markdown"),
                 "grobid": "off", "pdf_engine_available": True,
                 "settings": {"ok": True, "theme": "light", "output_dir": str(root / "Markdown"),
                              "auto_update": False, "current_version": "1.2.4", "update": {"status": "idle"}}}
        bootstrap = """<script>window.__calls=[];window.__fixture=STATE;window.__revision='saved-B-v1';
window.__fixture.settings.theme=localStorage.getItem('__fixture-theme')||'light';
window.pywebview={api:new Proxy({},{get:(_,name)=>(...args)=>{window.__calls.push({name,args});
if(name==='state')return Promise.resolve(window.__fixture);
if(name==='settings')return Promise.resolve(window.__fixture.settings);
if(name==='set_theme'){
if(window.__themeRejectNext){window.__themeRejectNext=false;return Promise.resolve({ok:false,why:'테마 저장 실패 fixture'});}
window.__fixture.settings={...window.__fixture.settings,theme:args[0]};localStorage.setItem('__fixture-theme',args[0]);return Promise.resolve(window.__fixture.settings);}
if(name==='extract_selected')return Promise.resolve(window.__rejectNext?{ok:true,n:0,queued:true,accepted_paths:[]}:{ok:true,n:1,queued:!!window.__queueMode});
if(name==='search'&&args[2]!==window.__revision)return Promise.resolve({ok:false,code:'stale_revision',why:'새 Markdown이 저장되었습니다. MD를 다시 열어 검색하세요.'});
if(name==='search'||name==='search_source')return Promise.resolve({results:[{text:'needle fixture',side:name==='search_source'?'source':'result',page:2,locations:[{page:2,bbox:[20,30,100,60],height:500}]}]});
return Promise.resolve({ok:true,n:1});}})};</script>""".replace("STATE", json.dumps(state))

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *_):
                pass

            def do_GET(self):
                from urllib.parse import urlparse, parse_qs
                url = urlparse(self.path)
                if url.path == "/":
                    data, mime = INDEX_HTML.replace("</head>", bootstrap + "</head>"), "text/html; charset=utf-8"
                elif url.path == "/files":
                    data, mime = json.dumps(state), "application/json"
                elif url.path == "/open":
                    index = int(parse_qs(url.query).get("p", [0])[0])
                    payload = {"key": 1, "revision": "saved-B-v1", "name": files[index]["name"], "pages": [[400,500]] * 30,
                        "source_ready": True, "source_unit_label": "쪽", "source_label": "원본 PDF",
                        "doc": {"extracted": True, "html": "<h1>Stored B</h1><p data-block-id='b-old'>Old snapshot mapped paragraph</p>" + "".join(f"<p>Stored paragraph {i} needle</p>" for i in range(80)),
                                "md_text": "# Stored B\n\nneedle"}}
                    if parse_qs(url.query).get("mode", [""])[0] == "result":
                        payload = {"name": payload["name"], "doc": payload["doc"], "revision": payload["revision"],
                                   "source_ready": False, "pages": []}
                    data, mime = json.dumps(payload), "application/json"
                elif url.path == "/document-map":
                    data, mime = json.dumps({"revision": "saved-B-v1", "pages": [{"width":400,"height":500}] * 30,
                        "blocks": [{"id":"b-old", "mapping":"exact", "locations":[{"page":0,"bbox":[20,30,100,60],"width":400,"height":500}]}]}), "application/json"
                elif url.path == "/page":
                    data, mime = '<svg xmlns="http://www.w3.org/2000/svg" width="400" height="500"><rect width="400" height="500" fill="white"/></svg>', "image/svg+xml"
                elif url.path == "/regression":
                    data, mime = Path(__file__).with_suffix(".js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/queue-regression":
                    data, mime = Path(__file__).with_name("ui_queue.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/resources-regression":
                    data, mime = Path(__file__).with_name("ui_resources.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/minimal-regression":
                    data, mime = Path(__file__).with_name("ui_minimal.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/viewer-regression":
                    data, mime = Path(__file__).with_name("ui_viewer.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/presentation-regression":
                    data, mime = Path(__file__).with_name("ui_presentation.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/presentation-host-regression":
                    data, mime = Path(__file__).with_name("ui_presentation_host.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/presentation-fallback-regression":
                    data, mime = Path(__file__).with_name("ui_presentation_fallback.js").read_text(encoding="utf-8"), "text/javascript; charset=utf-8"
                elif url.path == "/presentation-test-host":
                    data, mime = '<!doctype html><html><body style="margin:0"><div id="host" style="width:1000px;height:700px"></div></body></html>', "text/html; charset=utf-8"
                elif url.path == "/presentation-data" and presentation_bytes is not None:
                    if parse_qs(url.query).get("r", [""])[0] != "fixture-source-v1":
                        self.send_error(409)
                        return
                    data, mime = presentation_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                elif url.path in ("/presentation-frame", "/presentation-frame.mjs"):
                    filename = "presentation-frame.html" if url.path == "/presentation-frame" else "presentation-frame.mjs"
                    asset = Path(__file__).resolve().parents[1] / "web" / filename
                    if not asset.is_file():
                        self.send_error(404)
                        return
                    data, mime = asset.read_text(encoding="utf-8"), "text/html; charset=utf-8" if filename.endswith("html") else "text/javascript; charset=utf-8"
                elif url.path == "/presentation-vendor.js" and args.presentation_vendor:
                    data, mime = args.presentation_vendor.read_bytes(), "text/javascript; charset=utf-8"
                else:
                    self.send_error(404)
                    return
                payload = data if isinstance(data, bytes) else data.encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", mime)
                self.send_header("Content-Length", str(len(payload)))
                if url.path in ("/presentation-frame.mjs", "/presentation-vendor.js"):
                    self.send_header("Access-Control-Allow-Origin", "*")
                if url.path == "/presentation-frame":
                    origin = f"http://127.0.0.1:{server.server_port}"
                    self.send_header("Content-Security-Policy", f"default-src 'none'; script-src {origin}; style-src 'unsafe-inline'; img-src data: blob:; font-src data: blob:; connect-src 'none'; worker-src 'none'; base-uri 'none'; form-action 'none'")
                self.end_headers()
                self.wfile.write(payload)

            def do_POST(self):
                if self.path != "/shutdown":
                    self.send_error(404)
                    return
                self.send_response(200)
                self.end_headers()
                threading.Thread(target=server.shutdown, daemon=True).start()

        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        print(json.dumps({"url": f"http://127.0.0.1:{server.server_port}/", "frontend_fixture": True}), flush=True)
        try:
            server.serve_forever()
        finally:
            server.server_close()


if __name__ == "__main__":
    main()
