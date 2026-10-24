"""Read-only native source windows; no PDF/MD/extraction prerequisite."""
import hashlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import zipfile

from docx import Document
from openpyxl import Workbook
from PIL import Image

from pubnexus import native_preview
from pubnexus.preview import PreviewUnavailable


class NativePreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='native-preview-test-')
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)

    def docx(self):
        path = self.base / 'source.docx'
        doc = Document()
        doc.add_paragraph('<script>literal 한글 😀</script>')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'A1'
        table.cell(1, 1).text = 'B2'
        image = io.BytesIO()
        Image.new('RGB', (12, 8), 'white').save(image, format='PNG')
        doc.add_picture(io.BytesIO(image.getvalue()))
        doc.add_paragraph('Last paragraph')
        doc.save(path)
        return path, image.getvalue()

    def xlsx(self):
        path = self.base / 'source.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.title = 'First 한글'
        for row in range(1, 301):
            sheet.cell(row, 1, row)
            sheet.cell(row, 2, '<img onerror=alert(1)>')
        sheet.cell(2, 3, '=SUM(A1:A2)')
        book.create_sheet('Second').cell(1, 1, 'Second value')
        book.create_sheet('Hidden').sheet_state = 'hidden'
        book.save(path)
        book.close()
        return path

    def test_docx_native_order_window_and_original_image_are_read_without_extraction(self):
        path, png = self.docx()
        before = path.read_bytes()
        with (patch('pubnexus.documents.extract_document', side_effect=AssertionError('no extraction')),
              patch('pubnexus.preview.preview_path', side_effect=AssertionError('no conversion'))):
            first = native_preview.view(path, count=1)
            self.assertEqual(first['source_kind'], 'content')
            self.assertEqual(first['blocks'][0]['text'], '<script>literal 한글 😀</script>')
            self.assertEqual(first['next_start'], 1)
            self.assertTrue(first['truncated'])
            table = native_preview.view(path, start=1, count=1, revision=first['source_revision'])
            self.assertEqual(table['blocks'][0]['rows'], [['A1', ''], ['', 'B2']])
            image = native_preview.view(path, start=2, count=1)['blocks'][-1]
            self.assertEqual(image['kind'], 'image')
            self.assertNotIn('path', image)
            self.assertNotIn('member', image)
            raw, mime = native_preview.read_asset(path, image['asset_id'], first['source_revision'])
            self.assertEqual(raw, png)
            self.assertEqual(mime, 'image/png')
            self.assertEqual(image['image_sha256'], hashlib.sha256(png).hexdigest())
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(first['source_revision'], hashlib.sha256(before).hexdigest())
        self.assertEqual(sorted(p.name for p in self.base.iterdir()), [path.name])

    def test_xlsx_uses_read_only_visible_window_and_preserves_formula_and_literal_text(self):
        path = self.xlsx()
        from openpyxl.reader import excel
        calls = []
        real = excel.ExcelReader
        def load(*args, **kwargs):
            calls.append(kwargs)
            return real(*args, **kwargs)
        with patch.object(excel, 'ExcelReader', side_effect=load):
            first = native_preview.view(path, row_count=2, col_count=3)
        self.assertEqual(first['source_kind'], 'sheet')
        self.assertEqual([s['name'] for s in first['sheets']], ['First 한글', 'Second', 'Hidden'])
        self.assertEqual(first['sheets'][-1]['state'], 'hidden')
        self.assertEqual(first['rows'][1], ['2', '<img onerror=alert(1)>', '=SUM(A1:A2)'])
        self.assertEqual(len(first['rows']), 2)
        self.assertTrue(first['has_more'])
        self.assertTrue(all(c['read_only'] and not c['data_only'] and not c['keep_links'] for c in calls))
        second = native_preview.view(path, sheet_index=1, start_row=1, row_count=1, col_count=1)
        self.assertEqual(second['rows'], [['Second value']])
        later = native_preview.view(path, start_row=201, row_count=3, col_count=1)
        self.assertEqual(later['rows'], [['201'], ['202'], ['203']])

    def test_changed_revision_cancel_and_bad_windows_stop_before_parser(self):
        path, _ = self.docx()
        current = native_preview.view(path)['source_revision']
        path.write_bytes(path.read_bytes() + b'changed')
        with patch('docx.Document', side_effect=AssertionError('reject before parse')):
            with self.assertRaisesRegex(PreviewUnavailable, '변경'):
                native_preview.view(path, revision=current)
            with self.assertRaisesRegex(PreviewUnavailable, '취소'):
                native_preview.view(path, cancel=lambda: True)
            for params in ({'count': 0}, {'count': 201}, {'start': -1}, {'row_count': 1000}, {'col_count': 1000}):
                with self.subTest(params=params), self.assertRaises(PreviewUnavailable):
                    native_preview.view(path, **params)
        self.assertIsNone(native_preview.view(self.base / 'unsupported.pdf'))

    def test_zip_path_asset_identifier_and_byte_budgets_are_guarded(self):
        path, _ = self.docx()
        revision = native_preview.view(path)['source_revision']
        with self.assertRaises(PreviewUnavailable):
            native_preview.read_asset(path, '../word/document.xml', revision)
        with patch.object(native_preview, 'MAX_SOURCE_BYTES', 64), patch('docx.Document', side_effect=AssertionError()):
            with self.assertRaisesRegex(PreviewUnavailable, '크기'):
                native_preview.view(path)
        with zipfile.ZipFile(path, 'a') as archive:
            archive.writestr('../outside', 'bad')
        with self.assertRaisesRegex(PreviewUnavailable, '경로'):
            native_preview.view(path)

    def test_asset_reuses_revision_bound_bytes_without_reparsing_or_rehashing_each_picture(self):
        path, png = self.docx()
        view = native_preview.view(path)
        image = next(block for block in view['blocks'] if block['kind'] == 'image')
        with (patch('docx.Document', side_effect=AssertionError('cached image must not reparse')),
              patch.object(native_preview, '_Read', side_effect=AssertionError('cached image must not rehash'))):
            for _ in range(3):
                self.assertEqual(native_preview.read_asset(path, image['asset_id'], view['source_revision'])[0], png)
        path.write_bytes(path.read_bytes() + b'changed')
        with self.assertRaisesRegex(PreviewUnavailable, '변경'):
            native_preview.read_asset(path, image['asset_id'], view['source_revision'])

    def test_docx_and_xlsx_search_return_exact_native_addresses_without_pdf(self):
        path, _ = self.docx()
        result = native_preview.search(path, '한글')
        self.assertEqual(result['results'][0]['locations'][0],
                         {'kind': 'content', 'block_id': 'docx-b000001', 'offset': 16, 'length': 2})
        sheet = self.xlsx()
        result = native_preview.search(sheet, 'Second value')
        loc = result['results'][0]['locations'][0]
        self.assertEqual({k: loc[k] for k in ('kind', 'sheet_index', 'sheet', 'cell', 'row', 'col')},
                         {'kind': 'sheet', 'sheet_index': 1, 'sheet': 'Second', 'cell': 'A1', 'row': 1, 'col': 1})

    def test_formula_only_docx_explicitly_reports_partial_content_and_original_layout_limit(self):
        from docx.oxml import parse_xml
        path = self.base / 'formula.docx'
        native = Document()
        paragraph = native.add_paragraph()
        paragraph._p.append(parse_xml('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x=1</m:t></m:r></m:oMath>'))
        native.save(path)
        result = native_preview.view(path)
        self.assertTrue(result['partial'])
        self.assertTrue(any('수식' in warning for warning in result['warnings']))
        self.assertEqual(result['blocks'][0]['text'], '')

    def test_image_metadata_and_cache_have_independent_byte_and_count_budgets(self):
        path, png = self.docx()
        doc = Document(path)
        doc.add_picture(io.BytesIO(png))
        doc.save(path)
        with patch.object(native_preview, 'MAX_WINDOW_IMAGES', 1):
            result = native_preview.view(path)
        self.assertEqual(sum(b['kind'] == 'image' for b in result['blocks']), 1)
        self.assertTrue(any('그림' in warning and '제한' in warning for warning in result['warnings']))
        with patch.object(native_preview, 'MAX_ASSET_CACHE_BYTES', 1):
            native_preview.view(path)
        with native_preview._assets_lock:
            self.assertEqual(len(native_preview._assets), 0)

    def test_nonfinite_timeout_and_empty_search_are_bounded_before_native_parser(self):
        path, _ = self.docx()
        with patch('docx.Document', side_effect=AssertionError('no native parse')):
            self.assertEqual(native_preview.search(path, '')['results'], [])
            for timeout in (float('nan'), float('inf'), -1):
                with self.assertRaises(PreviewUnavailable):
                    native_preview.search(path, 'query', timeout=timeout)

    def test_xlsx_without_dimension_metadata_still_offers_next_real_row_window(self):
        path = self.base / 'streamed.xlsx'
        book = Workbook(write_only=True)
        sheet = book.create_sheet('Streamed')
        for row in range(5):
            sheet.append([row])
        book.save(path)
        book.close()
        first = native_preview.view(path, row_count=2, col_count=1)
        self.assertTrue(first['has_more'])
        self.assertEqual(first['rows'], [['0'], ['1']])
        final = native_preview.view(path, start_row=5, row_count=2, col_count=1)
        self.assertFalse(final['has_more'])
        self.assertEqual(final['rows'], [['4']])

    def test_xlsx_reuses_standard_stream_parser_without_whole_sheet_dimension_pass(self):
        from openpyxl.worksheet._reader import WorkSheetParser
        path = self.xlsx()
        with patch.object(WorkSheetParser, 'parse_dimensions', side_effect=AssertionError('no whole-sheet prepass')):
            result = native_preview.view(path, row_count=2, col_count=3)
            self.assertEqual(result['rows'][1][-1], '=SUM(A1:A2)')

    def test_xlsx_preserves_raw_formula_cached_value_shared_formula_dates_and_hangul(self):
        from datetime import datetime
        from lxml import etree
        path = self.base / 'formulas.xlsx'
        book = Workbook()
        sheet = book.active
        sheet.append([1, '한글 원문', '=A1*2', datetime(2026, 10, 24)])
        sheet.append([2, '둘째', '=A2*2'])
        book.save(path)
        book.close()
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        buffer = io.BytesIO()
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(buffer, 'w') as modified:
            for entry in original.infolist():
                body = original.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    root = etree.fromstring(body)
                    for row, value in ((1, '2'), (2, '4')):
                        cell = root.find(f'.//{ns}c[@r="C{row}"]')
                        formula = cell.find(ns + 'f')
                        formula.set('t', 'shared')
                        formula.set('si', '0')
                        if row == 1:
                            formula.set('ref', 'C1:C2')
                        else:
                            formula.text = None
                        cell.find(ns + 'v').text = value
                    body = etree.tostring(root)
                modified.writestr(entry, body)
        path.write_bytes(buffer.getvalue())
        result = native_preview.view(path, start_row=2, row_count=1, col_count=4)
        self.assertEqual(result['rows'][0][:3], ['2', '둘째', '=A2*2'])
        self.assertEqual(result['formula_cells'][0]['cached_value'], 4)
        self.assertEqual(result['formula_cells'][0]['formula'], '=A2*2')
        first = native_preview.view(path, row_count=1, col_count=4)
        self.assertEqual(first['rows'][0][-1], '2026-10-24T00:00:00')

    def test_formula_cached_metadata_obeys_the_same_window_byte_budget(self):
        from lxml import etree
        path = self.xlsx()
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        buffer = io.BytesIO()
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(buffer, 'w') as modified:
            for entry in original.infolist():
                body = original.read(entry.filename)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    root = etree.fromstring(body)
                    cell = root.find(f'.//{ns}c[@r="C2"]')
                    cell.set('t', 'str')
                    cell.find(ns + 'v').text = 'x' * 1000
                    body = etree.tostring(root)
                modified.writestr(entry, body)
        path.write_bytes(buffer.getvalue())
        with patch.object(native_preview, 'MAX_WINDOW_BYTES', 100), self.assertRaisesRegex(PreviewUnavailable, '내용 창'):
            native_preview.view(path, start_row=2, row_count=1, col_count=3)

    def test_odt_and_epub_reuse_native_content_parsers_and_cache_only_original_image_bytes(self):
        png = io.BytesIO()
        Image.new('RGB', (12, 8), 'white').save(png, format='PNG')
        odt = self.base / 'source.odt'
        content = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
            xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
            xmlns:xlink="http://www.w3.org/1999/xlink"><office:body><office:text><text:p>한글 첫 문단</text:p>
            <text:p>둘째 문단</text:p><draw:frame><draw:image xlink:href="Pictures/source.png"/></draw:frame>
            </office:text></office:body></office:document-content>'''
        with zipfile.ZipFile(odt, 'w') as archive:
            archive.writestr('content.xml', content)
            archive.writestr('Pictures/source.png', png.getvalue())
        epub = self.base / 'source.epub'
        with zipfile.ZipFile(epub, 'w') as archive:
            archive.writestr('META-INF/container.xml', '<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
            archive.writestr('OPS/book.opf', '<package><manifest><item id="c1" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest><spine><itemref idref="c1"/></spine></package>')
            archive.writestr('OPS/chapter.xhtml', '<html><body><p>한글 첫 문단</p><p>둘째 문단</p><img src="source.png"/><script>never execute</script><img src="https://invalid.test/private.png"/></body></html>')
            archive.writestr('OPS/source.png', png.getvalue())
        for path in (odt, epub):
            with self.subTest(suffix=path.suffix):
                before = path.read_bytes()
                with patch('pubnexus.store.root', side_effect=AssertionError('no user store')):
                    first = native_preview.view(path, count=1)
                    self.assertTrue(first['truncated'])
                    result = native_preview.view(path)
                    self.assertTrue(any('한글 첫 문단' in b.get('text', '') for b in result['blocks']))
                    image = next(b for b in result['blocks'] if b['kind'] == 'image')
                    body, mime = native_preview.read_asset(path, image['asset_id'], result['source_revision'])
                    self.assertEqual((body, mime), (png.getvalue(), 'image/png'))
                self.assertNotIn('never execute', str(result['blocks']))
                self.assertEqual(path.read_bytes(), before)

    def test_hwp_and_hwpx_reuse_existing_read_only_semantic_adapter_without_pdf_or_store(self):
        def rhwp(sink, path):
            unit = sink.unit('Native heading', 'heading')
            sink.paragraph(unit, '한글 본문')
            sink.table(unit, 1, 2, [{'row': 0, 'col': 0, 'text': 'A'}, {'row': 0, 'col': 1, 'text': 'B'}])
        for suffix in ('.hwp', '.hwpx'):
            path = self.base / ('source' + suffix)
            path.write_bytes(b'synthetic engine input')
            with (patch('pubnexus.documents._rhwp', side_effect=rhwp),
                  patch('pubnexus.preview.preview_path', side_effect=AssertionError('no PDF')),
                  patch('pubnexus.store.root', side_effect=AssertionError('no store'))):
                result = native_preview.view(path, count=2)
                found = native_preview.search(path, '한글')
            self.assertEqual([b['text'] for b in result['blocks']], ['Native heading', '한글 본문'])
            self.assertTrue(result['truncated'])
            self.assertEqual(found['results'][0]['locations'][0]['block_id'], suffix[1:] + '-b000002')

    def test_rtf_uses_striprtf_native_plain_reader_without_office_or_pandoc(self):
        path = self.base / 'source.rtf'
        path.write_bytes(br"{\rtf1\ansi\ansicpg949 \'bf\'f8\'b9\'ae \'c7\'d1\'b1\'db\par Second paragraph}")
        before = path.read_bytes()
        with (patch('pubnexus.external_documents.find_engine', side_effect=AssertionError('no office prerequisite')),
              patch('pubnexus.external_documents._run', side_effect=AssertionError('no conversion')),
              patch('pubnexus.store.root', side_effect=AssertionError('no store'))):
            result = native_preview.view(path)
            found = native_preview.search(path, 'Second')
        self.assertEqual([b['text'] for b in result['blocks']], ['원문 한글', 'Second paragraph'])
        self.assertTrue(result['partial'])
        self.assertTrue(any('그림' in warning for warning in result['warnings']))
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(found['results'][0]['locations'][0]['block_id'], 'rtf-b000002')

    def test_ods_repeated_blank_capacity_is_windowed_without_expansion_and_search_uses_real_cell(self):
        path = self.base / 'source.ods'
        content = '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
            xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
            <office:body><office:spreadsheet><table:table table:name="한글 Sheet">
            <table:table-row><table:table-cell><text:p>원문</text:p></table:table-cell>
            <table:table-cell table:formula="of:=1+1" office:value="2"><text:p>2</text:p></table:table-cell></table:table-row>
            <table:table-row table:number-rows-repeated="1000000"><table:table-cell table:number-columns-repeated="16384"/></table:table-row>
            <table:table-row><table:table-cell><text:p>End</text:p></table:table-cell></table:table-row>
            </table:table><table:table table:name="Second"><table:table-row><table:table-cell><text:p>다른 시트</text:p></table:table-cell>
            </table:table-row></table:table></office:spreadsheet></office:body></office:document-content>'''
        with zipfile.ZipFile(path, 'w') as archive:
            archive.writestr('content.xml', content)
        result = native_preview.view(path, row_count=2, col_count=2)
        self.assertEqual(result['source_kind'], 'sheet')
        self.assertEqual(result['rows'], [['원문', 'of:=1+1'], ['', '']])
        self.assertTrue(result['has_more'])
        self.assertEqual(result['formula_cells'][0]['cached_value'], '2')
        last = native_preview.view(path, start_row=1000002, row_count=2, col_count=1)
        self.assertEqual(last['rows'], [['End']])
        self.assertFalse(last['has_more'])
        self.assertEqual(native_preview.view(path, sheet_index=1, col_count=1)['rows'], [['다른 시트']])
        found = native_preview.search(path, 'End')
        self.assertEqual(found['results'][0]['locations'][0]['row'], 1000002)

    def test_explicit_unsupported_rtf_codepage_returns_to_existing_original_renderer(self):
        path = self.base / 'legacy-korean.rtf'
        path.write_bytes(br'{\rtf1\ansi\ansicpg949 legacy}')
        def conversion(sink, *args, **kwargs):
            sink.warn('외부 엔진 진단: [WARNING] Unsupported code page 949. Text will likely be garbled.')
        import sys
        with (patch.dict(sys.modules, {'striprtf': None, 'striprtf.striprtf': None}),
              patch('pubnexus.external_documents.find_engine', return_value=Path('pandoc.exe')),
              patch('pubnexus.external_documents._run', side_effect=conversion)):
            self.assertIsNone(native_preview.view(path))
            self.assertIsNone(native_preview.search(path, 'legacy'))

    def test_rtf_unicode_nested_cp949_raw_and_strict_invalid_bytes_keep_original_content(self):
        cases = [
            (br'{\rtf1\ansi\uc1\u-10916?{\b \u-20992?}\par \u-10179?\u-8704?}', ['한글', '😀']),
            (b'{\\rtf1\\ansi\\ansicpg949 ' + '한글 원문'.encode('cp949') + b'}', ['한글 원문']),
        ]
        for index, (data, expected) in enumerate(cases):
            path = self.base / f'native{index}.rtf'
            path.write_bytes(data)
            result = native_preview.view(path)
            self.assertEqual([block['text'] for block in result['blocks']], expected)
            self.assertEqual(path.read_bytes(), data)
        bad = self.base / 'bad.rtf'
        bad.write_bytes(br"{\rtf1\ansi\ansicpg949 \'c7}")
        with self.assertRaises(PreviewUnavailable):
            native_preview.view(bad)


if __name__ == '__main__':
    unittest.main()
